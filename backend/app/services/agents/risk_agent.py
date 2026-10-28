"""Agente Especialista en Gestión y Asignación de Riesgos Previsibles Contractuales.
Empresa de Servicios Públicos de Yolombó (ESPY S.A. E.S.P.) - Nit: 811.015.112-4.
Metodología: Colombia Compra Eficiente (CCE) y Ley 142 de 1994.
"""
from typing import Optional, List, Dict, Any, Tuple
from pathlib import Path
import re
import os
import json
from datetime import datetime

from sqlalchemy.orm import Session

from app.config import settings
from app.models.contract import Contrato
from app.models.document import Documento
from app.models.contractor import Contratista
from app.schemas.agent import RiskItemSchema, RiskMatrixInput, RiskMatrixResult
from app.services.document_parser import parse_document_text
from app.services.gemini_client import gemini_service
from app.utils.logger import logger
from app.utils.obligaciones_empresa_builder import is_short_term_period

# Import docx styling tools
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# ReportLab imports for institutional PDF generation
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas


def format_cop_standard(amount: Any) -> str:
    """
    Formatea el valor del contrato con la estructura estándar requerida: $4.800.000
    Sin decimales ni comas, con punto como separador de miles.
    """
    if amount is None or amount == "" or amount == 0:
        return "Conforme a presupuesto oficial"
    try:
        val_float = float(amount)
        val_int = int(round(val_float))
        formatted = f"{val_int:,}".replace(",", ".")
        return f"${formatted}"
    except Exception:
        return f"${amount}"


class RiskMatrixAgent:
    """
    Agente Especialista en Matrices de Riesgos Contractuales para ESPY S.A. E.S.P.
    Garantiza consistencia técnica, jurídica y aritmética absoluta bajo lineamientos CCE.
    """

    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc"}

    def __init__(self):
        self.nombre = "Especialista en Gestión y Asignación de Riesgos Previsibles - ESPY S.A. E.S.P."

    def extract_context_from_contract_and_docs(
        self, 
        input_data: RiskMatrixInput, 
        db: Optional[Session] = None
    ) -> Dict[str, Any]:
        """Extrae y consolida información a partir del contrato en BD y/o documentos cargados."""
        context = {
            "objeto": input_data.objeto or "",
            "tipologia": input_data.tipologia or "PRESTACION_SERVICIOS",
            "contratista_nombre": input_data.contratista_nombre or "",
            "valor_estimado": input_data.valor_estimado or 0.0,
            "plazo_texto": input_data.plazo_texto or "",
            "lugar_ejecucion": input_data.lugar_ejecucion or "Municipio de Yolombó, Antioquia",
            "exige_polizas": input_data.exige_polizas,
            "documentos_texto": input_data.estudio_previo_texto or "",
            "archivos_analizados": [],
            "custom_logo_path": input_data.custom_logo_path or ""
        }

        # 1. Si se suministró contrato_id, extraer metadatos del contrato en base de datos
        if input_data.contrato_id and db:
            contrato = db.query(Contrato).filter(Contrato.id == input_data.contrato_id).first()
            if contrato:
                if not context["objeto"] and contrato.objeto:
                    context["objeto"] = contrato.objeto
                if not input_data.tipologia and contrato.tipologia:
                    context["tipologia"] = contrato.tipologia.value
                if not context["valor_estimado"] and contrato.valor_total:
                    context["valor_estimado"] = contrato.valor_total
                if not context["plazo_texto"]:
                    context["plazo_texto"] = contrato.plazo_texto or f"{contrato.plazo_ejecucion} {contrato.unidad_plazo.value.lower()}"
                if not context["lugar_ejecucion"] and contrato.lugar_ejecucion:
                    context["lugar_ejecucion"] = contrato.lugar_ejecucion
                if contrato.contratista and not context["contratista_nombre"]:
                    context["contratista_nombre"] = contrato.contratista.nombre_o_razon_social
                if context["exige_polizas"] is None:
                    context["exige_polizas"] = bool(contrato.garantias_exigidas and len(contrato.garantias_exigidas) > 0)

                # Si no vinieron document_ids explícitos, usar los del contrato
                if not input_data.document_ids and contrato.documentos:
                    input_data.document_ids = [d.id for d in contrato.documentos]

        # 2. Extraer texto de documentos cargados (Estudio Previo, propuesta, etc.)
        if input_data.document_ids and db:
            doc_texts = []
            for doc_id in input_data.document_ids:
                doc = db.query(Documento).filter(Documento.id == doc_id).first()
                if not doc:
                    continue
                ext = Path(doc.nombre_archivo_original).suffix.lower()
                if ext in self.ALLOWED_EXTENSIONS:
                    full_path = settings.BASE_DIR / doc.ruta_relativa
                    if full_path.exists():
                        try:
                            txt = parse_document_text(full_path)
                            if txt:
                                doc_texts.append(f"=== DOCUMENTO: {doc.nombre_archivo_original} ===\n{txt}")
                                context["archivos_analizados"].append(doc.nombre_archivo_original)
                        except Exception as e:
                            logger.warning(f"Error parseando documento {doc.nombre_archivo_original}: {e}")
            if doc_texts:
                combined_docs = "\n\n".join(doc_texts)
                if context["documentos_texto"]:
                    context["documentos_texto"] = f"{context['documentos_texto']}\n\n{combined_docs}"
                else:
                    context["documentos_texto"] = combined_docs

        return context

    def validate_inputs_and_detect_gaps(self, context: Dict[str, Any]) -> Tuple[List[str], List[str]]:
        """
        Paso 1 Obligatorio: Evalúa insumos indispensables y advierte vacíos o inconsistencias del Estudio Previo.
        Retorna (insumos_faltantes_advertencias, observaciones_estudio_previo).
        """
        advertencias = []
        observaciones = []

        texto_completo = (context.get("documentos_texto") or "") + " " + (context.get("objeto") or "")

        # 1. Objeto contractual y alcance técnico
        if not context.get("objeto") and not re.search(r"OBJETO\s*:", texto_completo, re.IGNORECASE):
            advertencias.append("Falta definir con claridad el objeto contractual y el alcance técnico detallado del proceso.")
        
        # 2. Obligaciones y perfil
        if not re.search(r"obligar[aá]|obligaciones|perfil|idoneidad", texto_completo, re.IGNORECASE):
            advertencias.append("No se identificaron las obligaciones específicas detalladas del contratista ni el perfil/idoneidad técnica exigida.")

        # 3. Presupuesto, forma de pago y entregables
        if not context.get("valor_estimado") and not re.search(r"presupuesto|valor|\$|forma\s+de\s+pago", texto_completo, re.IGNORECASE):
            advertencias.append("Se requiere precisar el presupuesto estimado, la forma de pago y los entregables condicionados a desembolso.")

        # 4. Plazo y cronograma
        if not context.get("plazo_texto") and not re.search(r"plazo\s*(?:de\s*ejecuci[oó]n)?|cronograma", texto_completo, re.IGNORECASE):
            advertencias.append("Falta especificar el plazo de ejecución y cronograma de cumplimiento de las actividades.")

        # 5. Ubicación o condiciones territoriales en Yolombó
        if not re.search(r"yolomb[oó]|urbana|rural|plantas|redes", texto_completo, re.IGNORECASE):
            advertencias.append("No se delimitaron las condiciones de ejecución territorial (área urbana, veredas de Yolombó, plantas de tratamiento o redes).")

        # 6. Pólizas y garantías
        polizas_no_exigidas = bool(re.search(r"no\s+se\s+requerir[aá]\s+(?:la\s+constituci[oó]n\s+de\s+)?p[oó]lizas|no\s+se\s+exigen\s+garant[ií]as", texto_completo, re.IGNORECASE))
        if polizas_no_exigidas or context.get("exige_polizas") is False or context.get("exige_polizas") is None:
            if context.get("tipologia") == "SUMINISTRO":
                observaciones.append("Al haber estipulado el Estudio Previo que no se exigirán pólizas contractuales, se traslada el riesgo económico a la ESPY. Esto exige aplicar estrictamente el pago contra entrega física e instalación efectiva de los equipos, previa suscripción del Acta de Entrada a Almacén y recibo a satisfacción por el supervisor.")
            else:
                observaciones.append("Al haber estipulado el Estudio Previo que no se exigirán pólizas contractuales, se traslada el riesgo económico a la ESPY. Esto exige aplicar estrictamente el pago a mes vencido contra informe y recibo a satisfacción.")

        # Detección de errores y vacíos de redacción/plantilla en el texto analizado
        if re.search(r"\bgenti[oó]n\b|\bconplementarias\b", texto_completo, re.IGNORECASE):
            observaciones.append("Se detectaron y corrigieron erratas de digitación presentes en el estudio previo ('gentión' y/o 'conplementarias' normalizados a 'GESTIÓN' y 'COMPLEMENTARIAS').")

        if context.get("tipologia") == "PRESTACION_SERVICIOS" and re.search(r"contrato\s+de\s+suministro|suministro\s+de\s+bienes|compraventa", texto_completo, re.IGNORECASE):
            observaciones.append("El Estudio Previo contiene referencias cruzadas erróneas a 'contrato de suministro' (posible rezago de plantilla). Se tipificó estrictamente como Prestación de Servicios de Apoyo a la Gestión.")

        if re.search(r"(?:prestacion\s+de\s+servicios\s+para\s+){2,}", texto_completo, re.IGNORECASE):
            observaciones.append("Se detectó repetición redundante de frases en el texto del estudio previo, la cual fue depurada para la estructuración del riesgo.")

        if re.search(r"en\s+es\s+a\s+visitar\s+de\s+en", texto_completo, re.IGNORECASE):
            observaciones.append("Se detectaron fragmentos incompletos de plantilla en la delimitación espacial del estudio previo.")

        # Verificación de idoneidad genérica
        if re.search(r"pues\s+a\s+ello\s+dedica\s+su\s+actividad\s+econ[oó]mica", texto_completo, re.IGNORECASE) and not re.search(r"(?:t[ií]tulo|a[ñn]os\s+de\s+experiencia|tecn[oó]logo|profesional|meses)", texto_completo, re.IGNORECASE):
            observaciones.append("La sección de Idoneidad del Estudio Previo señala de manera genérica la capacidad técnica sin estipular un perfil profesional o técnico específico requerido.")

        return advertencias, observaciones

    def normalize_risk_items(
        self, 
        raw_items: List[Dict[str, Any]], 
        tipologia: str,
        objeto_text: str = "",
        plazo_str: Optional[str] = None
    ) -> List[RiskItemSchema]:
        """
        Garantiza consistencia técnica, verificación aritmética absoluta y orden correlativo sin saltos.
        Probabilidad (1-5) + Impacto (1-5) = Valoración.
        Categorías: Bajo (2-4), Medio (5), Alto (6-7), Extremo (8-10).
        Aplica reglas dinámicas de objeto (mantenimiento vs comunitario) y de plazo (<= 30 días).
        """
        normalized: List[RiskItemSchema] = []
        consecutivo = 1

        obj_lower = (objeto_text or "").lower()
        is_mantenimiento = any(k in obj_lower for k in ["mantenimiento", "herramienta", "equipo", "electromec", "reparaci", "taller", "eléctric", "electric"])
        es_corto_plazo = is_short_term_period(plazo_str=plazo_str)

        for raw in raw_items:
            # 1. Asegurar y acotar Probabilidad e Impacto Inicial
            try:
                p_ini = min(max(int(raw.get("probabilidad_inicial", 2)), 1), 5)
            except Exception:
                p_ini = 2
            try:
                i_ini = min(max(int(raw.get("impacto_inicial", 2)), 1), 5)
            except Exception:
                i_ini = 2
            val_ini = p_ini + i_ini
            if val_ini <= 4:
                cat_ini = "Bajo"
            elif val_ini == 5:
                cat_ini = "Medio"
            elif val_ini in (6, 7):
                cat_ini = "Alto"
            else:
                cat_ini = "Extremo"

            # 2. Asegurar y acotar Probabilidad e Impacto Residual
            try:
                p_res = min(max(int(raw.get("probabilidad_residual", 1)), 1), 5)
            except Exception:
                p_res = 1
            try:
                i_res = min(max(int(raw.get("impacto_residual", 2)), 1), 5)
            except Exception:
                i_res = 2
            val_res = p_res + i_res
            if val_res <= 4:
                cat_res = "Bajo"
            elif val_res == 5:
                cat_res = "Medio"
            elif val_res in (6, 7):
                cat_res = "Alto"
            else:
                cat_res = "Extremo"

            # 3. Tipificación y adecuación SST
            tipo = str(raw.get("tipo", "Operacional")).strip()
            desc = str(raw.get("descripcion", "")).strip()
            cons = str(raw.get("consecuencia", "")).strip()
            trat = str(raw.get("tratamiento", "")).strip()
            mec = str(raw.get("monitoreo_mecanismo", "Revisión documental técnica")).strip()
            per = str(raw.get("monitoreo_periodicidad", "Mensual (con presentación de informe y cuenta)")).strip()

            if any(k in desc.lower() or k in tipo.lower() for k in ["sst", "seguridad y salud", "accidente", "incidente laboral", "epp", "arl"]):
                tipo = "Operacional / SST"

            # 4. Prohibición de conceptos ajenos a servicios personales
            if tipologia == "PRESTACION_SERVICIOS" and not is_mantenimiento:
                desc = re.sub(r"suministro\s+de\s+bienes|materiales\s+de\s+obra|subalternos\s+a\s+cargo", "entregables técnicos acordados", desc, flags=re.IGNORECASE)

            # 5. Si el objeto es de Mantenimiento / Herramientas / Equipos:
            # Eliminar tajantemente menciones a Corantioquia/SSPD o renuencia comunitaria/separación en la fuente
            if is_mantenimiento:
                if "corantioquia" in cons.lower() or "sspd" in cons.lower():
                    cons = re.sub(r"ante\s+autoridades\s+ambientales\s+competentes\s*\([^)]*\)", "en la operación de los servicios públicos de la ESPY", cons, flags=re.IGNORECASE)
                    cons = re.sub(r"\(?Corantioquia/SSPD\)?", "la ESPY", cons, flags=re.IGNORECASE)
                if "separación en la fuente" in desc.lower() or "renuencia comunitaria" in desc.lower() or "apatía" in desc.lower():
                    # Reemplazar por riesgo de daño o pérdida de herramientas bajo custodia
                    tipo = "Operacional / Custodia"
                    desc = "Daño, extravío o deterioro de herramientas o equipos propiedad de la ESPY durante su traslado, custodia o intervención en taller."
                    cons = "Afectación al patrimonio público de la ESPY y retraso en la disponibilidad de equipos para atender contingencias operativas."
                    trat = "Levantamiento de inventario y estado inicial al entregar y recibir los equipos mediante acta suscrita por el supervisor. El contratista asume a todo costo la reposición inmediata de cualquier elemento dañado o extraviado."
                    mec = "Inspección técnica detallada y verificación física contra acta de entrega."
                    per = "Al inicio, durante la ejecución y en la liquidación/pago único" if es_corto_plazo else "Al momento de cada entrega de equipos."

            # 6. Adaptación dinámica al plazo de ejecución (Regla de Periodicidad <= 30 días)
            if es_corto_plazo:
                # Reemplazo de frases mensuales en descripción
                desc = re.sub(
                    r"informes\s+mensuales\s+de\s+actividades\s+y\s+evidencias\s+de\s+ejecuci[oó]n[^.]*dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes",
                    "entrega final e informe técnico dentro del plazo pactado",
                    desc,
                    flags=re.IGNORECASE
                )
                desc = re.sub(r"informes\s+mensuales\s+dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes", "entrega final e informe técnico dentro del plazo pactado", desc, flags=re.IGNORECASE)
                desc = re.sub(r"\binformes\s+mensuales\b", "informe final de actividades", desc, flags=re.IGNORECASE)
                desc = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de pago único", desc, flags=re.IGNORECASE)
                desc = re.sub(r"\bpago\s+mensual\b", "pago único", desc, flags=re.IGNORECASE)
                desc = re.sub(r"\bvalor\s+mensualizado\b", "valor total", desc, flags=re.IGNORECASE)

                # Reemplazo en consecuencia
                cons = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de liquidación y pago", cons, flags=re.IGNORECASE)
                cons = re.sub(r"\bpago\s+mensual\b", "pago único", cons, flags=re.IGNORECASE)

                # Reemplazo en tratamiento
                trat = re.sub(r"\ba\s+mes\s+vencido\b", "contra entrega final a satisfacción", trat, flags=re.IGNORECASE)
                trat = re.sub(r"\btr[aá]mite\s+mensual\b", "trámite de pago", trat, flags=re.IGNORECASE)

                # Reemplazo en periodicidad y monitoreo
                if "mensual" in per.lower() or "quincenal" in per.lower():
                    per = "Al inicio, durante la ejecución y en la liquidación/pago único"
                elif per.lower() in ["diario", "diaria", "permanente"]:
                    per = "Al inicio, durante la ejecución y en la liquidación/pago único"
            else:
                if per.lower() in ["diario", "diaria", "todos los dias", "permanente"]:
                    per = "Mensual (con la presentación de cuenta de cobro e informe)"

            item = RiskItemSchema(
                numero=consecutivo,
                clase=raw.get("clase", "Específico"),
                fuente=raw.get("fuente", "Interno"),
                etapa=raw.get("etapa", "Ejecución"),
                tipo=tipo,
                descripcion=desc,
                consecuencia=cons,
                probabilidad_inicial=p_ini,
                impacto_inicial=i_ini,
                valoracion_inicial=val_ini,
                categoria_inicial=cat_ini,
                asignacion=str(raw.get("asignacion", "Contratista (100%)")).strip(),
                tratamiento=trat,
                probabilidad_residual=p_res,
                impacto_residual=i_res,
                valoracion_residual=val_res,
                categoria_residual=cat_res,
                afecta_ejecucion=str(raw.get("afecta_ejecucion", "Sí")).strip(),
                responsable=str(raw.get("responsable", "Supervisor ESPY y Contratista")).strip(),
                monitoreo_mecanismo=mec,
                monitoreo_periodicidad=per
            )
            normalized.append(item)
            consecutivo += 1

        return normalized


    def fallback_risk_matrix(self, context: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Generador determinista normativo de riesgos bajo CCE y Ley 142/1994 si no hay conexión al LLM.
        Se adapta a la tipología contractual y caso específico.
        """
        tipologia = context.get("tipologia", "PRESTACION_SERVICIOS")
        obj_text = (context.get("objeto") or "").lower()
        is_pgirs = "pgirs" in obj_text or "residuos" in obj_text or "reciclaje" in obj_text

        if tipologia == "PRESTACION_SERVICIOS":
            is_mantenimiento = any(k in obj_text for k in ["mantenimiento", "herramienta", "equipo", "electromec", "reparaci", "taller", "eléctric", "electric"])

            if is_mantenimiento:
                return [
                    {
                        "clase": "Específico",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Operacional / Entrega",
                        "descripcion": "Incumplimiento o retraso injustificado en la entrega oportuna de las herramientas o equipos intervenidos y del informe técnico de reparación dentro del plazo pactado.",
                        "consecuencia": "Parálisis o afectación en las cuadrillas operativas de la ESPY por falta de herramientas para atender daños en redes y mantenimiento municipal.",
                        "probabilidad_inicial": 3, "impacto_inicial": 3,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Pactar cronograma estricto de entrega técnica en taller o almacén. Alertas tempranas por parte de la supervisión y condicionamiento total del pago único a la entrega a satisfacción.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Cotejo documental de fecha de entrega material vs. plazo contractual pactado.",
                        "monitoreo_periodicidad": "Al inicio, durante la ejecución y en la liquidación/pago único"
                    },
                    {
                        "clase": "Específico",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Técnico / Calidad",
                        "descripcion": "Deficiente calidad en la reparación, mano de obra no calificada o instalación de repuestos usados o no originales en las herramientas y equipos de la ESPY.",
                        "consecuencia": "Falla prematura de los equipos reparados, reprocesos técnicos, riesgo de daño patrimonial y retrasos en la operación de servicios públicos.",
                        "probabilidad_inicial": 3, "impacto_inicial": 4,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Obligación de suministrar repuestos 100% nuevos y originales bajo modalidad a todo costo. Pruebas conjuntas de funcionamiento previas al recibo y exigencia de garantía técnica mínima sobre repuestos y mano de obra.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Pruebas de encendido, carga operativa y verificación física de repuestos sustituidos.",
                        "monitoreo_periodicidad": "Al momento de las pruebas técnicas y entrega en recibo a satisfacción."
                    },
                    {
                        "clase": "Específico",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Operacional / Custodia",
                        "descripcion": "Daño, deterioro adicional, avería o pérdida de las herramientas y equipos electromecánicos propiedad de la ESPY mientras permanezcan bajo custodia del contratista.",
                        "consecuencia": "Detrimento patrimonial para la ESPY, disminución del inventario operativo y eventuales controversias jurídicas contractuales.",
                        "probabilidad_inicial": 2, "impacto_inicial": 4,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Suscripción obligatoria de Acta de Entrega con inventario detallado del estado inicial. Responsabilidad civil y contractual directa del contratista por la custodia integral de los bienes hasta su devolución formal.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Inspección física contra inventario y acta de entrega técnica suscrita.",
                        "monitoreo_periodicidad": "Al inicio, durante la custodia en taller y en la entrega formal final."
                    },
                    {
                        "clase": "General",
                        "fuente": "Mixto",
                        "etapa": "Ejecución",
                        "tipo": "Operacional / SST",
                        "descripcion": "Riesgo de incidentes o accidentes laborales mecánicos o eléctricos (cortocircuitos, atrapamientos, cortes o quemaduras) durante la manipulación y reparación de herramientas eléctricas en taller o campo.",
                        "consecuencia": "Lesiones físicas al personal técnico, parálisis temporal de labores de mantenimiento y eventuales reclamaciones laborales.",
                        "probabilidad_inicial": 3, "impacto_inicial": 3,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Afiliación obligatoria y cobertura vigente al Sistema General de Riesgos Laborales (ARL), cumplimiento estricto del protocolo SST, uso riguroso de EPP dieléctricos y corte de energía en banco de trabajo.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor / Líder SST ESPY",
                        "monitoreo_mecanismo": "Validación de afiliación a ARL e inspección de normas de seguridad en taller.",
                        "monitoreo_periodicidad": "Al inicio de actividades y durante la ejecución."
                    },
                    {
                        "clase": "General",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Regulatorio",
                        "descripcion": "Omisión o mora en la acreditación del pago de aportes obligatorios al Sistema de Seguridad Social Integral (Salud, Pensión y ARL) conforme a la normatividad legal aplicable.",
                        "consecuencia": "Imposibilidad jurídica para la ESPY de tramitar el certificado de cumplimiento y efectuar el pago pactado conforme a la Ley 789 de 2002.",
                        "probabilidad_inicial": 2, "impacto_inicial": 2,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Control previo estricto por la supervisión: no se radicará cuenta de cobro ni factura sin la planilla PILA o certificado de aportes al día.",
                        "probabilidad_residual": 1, "impacto_residual": 1,
                        "afecta_ejecucion": "No",
                        "responsable": "Supervisor ESPY y Pagaduría",
                        "monitoreo_mecanismo": "Verificación documental de la planilla de pago a seguridad social.",
                        "monitoreo_periodicidad": "Al momento de radicar la cuenta de cobro y liquidación."
                    },
                    {
                        "clase": "Específico",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Financiero",
                        "descripcion": "Riesgo de afectación patrimonial a la entidad ante la no exigencia de pólizas de cumplimiento en el estudio previo del contrato.",
                        "consecuencia": "Eventual desprotección económica en caso de incumplimiento técnico o retiro injustificado sin finalizar las reparaciones.",
                        "probabilidad_inicial": 2, "impacto_inicial": 3,
                        "asignacion": "Entidad (100% ESPY)",
                        "tratamiento": "Mecanismo estricto de pago único condicionado contra recibo a entera satisfacción de las herramientas reparadas y operativas. Cero anticipos de recursos públicos.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Supervisor y Gerencia ESPY",
                        "monitoreo_mecanismo": "Verificación física y técnica de las herramientas entregadas antes de expedir recibido a satisfacción.",
                        "monitoreo_periodicidad": "En la entrega final de herramientas y liquidación contractual."
                    },
                    {
                        "clase": "General",
                        "fuente": "Interno",
                        "etapa": "Ejecución",
                        "tipo": "Jurídico / Regulatorio",
                        "descripcion": "Riesgo de pretendida relación laboral por subordinación continuada en labores de taller o mantenimiento técnico.",
                        "consecuencia": "Exposición a contingencias laborales ordinarias contra la entidad contratante.",
                        "probabilidad_inicial": 2, "impacto_inicial": 4,
                        "asignacion": "Compartido (ESPY 50% / Contratista 50%)",
                        "tratamiento": "Preservación integral de la autonomía técnica y directiva del contratista. Prohibición de imponer horarios de oficina o subordinación laboral.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "No",
                        "responsable": "Supervisor y Gerencia ESPY / Contratista",
                        "monitoreo_mecanismo": "Supervisión técnica enfocada estrictamente en resultados y estado operativo de los equipos.",
                        "monitoreo_periodicidad": "Durante el seguimiento a la ejecución técnica."
                    }
                ]

            entregables_nombre = "actividades del PGIRS municipal" if is_pgirs else "actividades contractuales pactadas"
            return [
                {
                    "clase": "Específico",
                    "fuente": "Interno",
                    "etapa": "Ejecución",
                    "tipo": "Operacional",
                    "descripcion": f"Incumplimiento o retraso injustificado en la entrega oportuna de los informes de actividades y evidencias de ejecución de las {entregables_nombre} dentro del plazo establecido.",
                    "consecuencia": "Retraso en el seguimiento al cronograma institucional, postergación en la expedición del certificado de cumplimiento y demora en el trámite de cuentas.",
                    "probabilidad_inicial": 3, "impacto_inicial": 2,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "La supervisión concertará con la contratista el cronograma detallado de entregables desde el acta de inicio. Emisión de alertas tempranas tres (3) días hábiles antes del vencimiento. Condicionamiento estricto del trámite de pago a la radicación completa y aprobada del informe y evidencias.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista y Supervisor ESPY",
                    "monitoreo_mecanismo": "Cotejo documental de fecha de radicación vs. plazo contractual.",
                    "monitoreo_periodicidad": "Al inicio, durante la ejecución y en la liquidación/pago único"
                },
                {
                    "clase": "Específico",
                    "fuente": "Interno",
                    "etapa": "Ejecución",
                    "tipo": "Técnico",
                    "descripcion": f"Deficiente calidad técnica, inconsistencia o falta de rigor metodológico en los informes, diagnósticos, registros o memorias de apoyo técnico a las {entregables_nombre} elaboradas por la contratista.",
                    "consecuencia": "Rechazo de informes, inconsistencias técnicas en la planeación institucional de la ESPY, retrabajo administrativo y necesidad de requerimientos reiterados.",
                    "probabilidad_inicial": 2, "impacto_inicial": 3,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Verificación del perfil técnico e idoneidad previa a la suscripción. La supervisión técnica evaluará la consistencia de cada entregable, disponiendo de hasta tres (3) días hábiles para exigir subsanaciones antes de otorgar el recibo a satisfacción.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista y Supervisor ESPY",
                    "monitoreo_mecanismo": "Revisión técnica de fondo y cotejo contra obligaciones contractuales.",
                    "monitoreo_periodicidad": "Con la entrega de cada informe o producto."
                },
                {
                    "clase": "General",
                    "fuente": "Mixto",
                    "etapa": "Ejecución",
                    "tipo": "Operacional / SST",
                    "descripcion": "Ocurrencia de accidentes o incidentes laborales, o afectación a la salud e integridad de la contratista durante actividades de campo, visitas o recorridos en el municipio de Yolombó.",
                    "consecuencia": "Incapacidad temporal, interrupción en el desarrollo de actividades en territorio, posibles reclamaciones contra la entidad y retraso en las metas contractuales.",
                    "probabilidad_inicial": 2, "impacto_inicial": 3,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Afiliación previa y obligatoria al Sistema General de Riesgos Laborales (ARL) en el nivel de riesgo aplicable. Cumplimiento estricto del SG-SST de la ESPY (Decreto 1072 de 2015), porte y uso permanente de EPP/EPI reglamentarios en campo, autocuidado y reporte inmediato por escrito de novedades.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista y Supervisor / Líder SST ESPY",
                    "monitoreo_mecanismo": "Validación de cobertura continua en ARL e inspección de uso de EPP.",
                    "monitoreo_periodicidad": "Al inicio de actividades y al presentarse el evento."
                },
                {
                    "clase": "General",
                    "fuente": "Interno",
                    "etapa": "Ejecución",
                    "tipo": "Regulatorio",
                    "descripcion": "Omisión, mora o liquidación errónea de los aportes obligatorios al Sistema de Seguridad Social Integral (Salud, Pensión y ARL) sobre el IBC normativo no menor al 40% del valor del contrato.",
                    "consecuencia": "Paralización del trámite de pago por mandato de Ley (Leyes 100/1993, 789/2002 y Dec. 1273/2018), posibles glosas administrativas y desprotección prestacional.",
                    "probabilidad_inicial": 3, "impacto_inicial": 2,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Control documental estricto por parte del supervisor: la ESPY no tramitará certificado de cumplimiento ni orden de pago sin la verificación aritmética de la planilla PILA con IBC cotizado mínimo legal debidamente cancelada.",
                    "probabilidad_residual": 1, "impacto_residual": 1,
                    "afecta_ejecucion": "No",
                    "responsable": "Supervisor ESPY y Pagaduría",
                    "monitoreo_mecanismo": "Revisión aritmética y temporal de la planilla PILA y certificado de pago.",
                    "monitoreo_periodicidad": "Requisito previo a cuenta de cobro o pago."
                },
                {
                    "clase": "Específico",
                    "fuente": "Interno",
                    "etapa": "Ejecución",
                    "tipo": "Financiero",
                    "descripcion": "Riesgo de afectación patrimonial a los recursos de la entidad generado por la no exigencia de pólizas de garantía de cumplimiento según lo estipulado en los estudios previos del proceso.",
                    "consecuencia": "Eventual desprotección económica frente a un abandono repentino de actividades o incumplimiento no subsanado previo al desembolso.",
                    "probabilidad_inicial": 2, "impacto_inicial": 3,
                    "asignacion": "Entidad (100% ESPY)",
                    "tratamiento": "Mecanismo de pago vencido post-prestación: se cancela estrictamente contra actividades efectivamente ejecutadas y avaladas por el supervisor. Cero anticipos o desembolsos previos. Liquidación proporcional exclusiva de los días de servicio prestados.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Supervisor y Gerencia ESPY",
                    "monitoreo_mecanismo": "Verificación de productos reales recibidos antes de expedir recibido a satisfacción.",
                    "monitoreo_periodicidad": "Al liquidar o aprobar el pago final."
                },
                {
                    "clase": "General",
                    "fuente": "Interno",
                    "etapa": "Ejecución",
                    "tipo": "Jurídico / Regulatorio",
                    "descripcion": "Riesgo de desnaturalización del contrato de prestación de servicios por subordinación laboral continuada o imposición de horarios de planta, generando pretendida relación laboral (contrato realidad).",
                    "consecuencia": "Exposición a demandas ordinarias laborales contra la ESPY, riesgo de reclamación de prestaciones sociales, cesantías y sanciones legales.",
                    "probabilidad_inicial": 2, "impacto_inicial": 4,
                    "asignacion": "Compartido (ESPY 50% / Contratista 50%)",
                    "tratamiento": "Preservación estricta de la autonomía técnica, administrativa y directiva de la contratista (Art. 23 CST y Ley 142/1994). Prohibición absoluta de exigir horarios de oficina fija, control de reloj biométrico o asignación de tareas operativas ordinarias de planta ajenas al objeto contractual.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "No",
                    "responsable": "Supervisor y Gerencia ESPY / Contratista",
                    "monitoreo_mecanismo": "Supervisión orientada exclusivamente a entregables y metas técnicas sin control de horario.",
                    "monitoreo_periodicidad": "Durante el seguimiento a la ejecución contractual."
                },
                {
                    "clase": "General",
                    "fuente": "Interno",
                    "etapa": "Ejecución / Liquidación",
                    "tipo": "Regulatorio",
                    "descripcion": "Incumplimiento del deber de confidencialidad y reserva o indebida utilización y custodia de la información técnica suministrada por la ESPY (Ley 1581 de 2012 de Protección de Datos Personales).",
                    "consecuencia": "Fuga de información confidencial de la ESPY, vulneración del Habeas Data, eventuales quejas o investigaciones administrativas ante la SIC y daño reputacional.",
                    "probabilidad_inicial": 2, "impacto_inicial": 3,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Inclusión de cláusula contractual expresa de confidencialidad, cesión total de derechos patrimoniales sobre los documentos a favor de la ESPY, uso restringido de datos a fines contractuales y restitución formal de archivos al liquidar el contrato.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "No",
                    "responsable": "Contratista y Supervisor ESPY",
                    "monitoreo_mecanismo": "Suscripción de compromiso al inicio y acta de devolución de información al cierre.",
                    "monitoreo_periodicidad": "En acta de inicio y en liquidación final."
                }
            ]

        elif tipologia == "SUMINISTRO":
            is_gps_risk = any(k in obj_text for k in [
                "gps", "telemetr", "adas", "dms", "jc371", "ci05f", "videovigilancia", "cámara vehicular", "camara vehicular", "parque automotor"
            ])
            if is_gps_risk:
                return [
                    {
                        "clase": "Específico", "fuente": "Externo", "etapa": "Ejecución", "tipo": "Operacional",
                        "descripcion": "Retraso o incumplimiento en la entrega física e instalación en sitio de las unidades GPS LC, cámaras vehiculares JC371 ADAS, cámaras de cabina CI05F DMS y accesorios en el parque automotor de la ESPY.",
                        "consecuencia": "Afectación en la planeación y control operativo de rutas del servicio de aseo y aplazamiento en la puesta en marcha de la telemetría institucional.",
                        "probabilidad_inicial": 3, "impacto_inicial": 3,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Coordinación de cronograma estricto de instalación en sitio y verificación de disponibilidad de los vehículos. Pago condicionado al recibo e instalación a entera satisfacción.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Verificación física de instalación por vehículo y firma de Acta de Entrada a Almacén y Recibo a Satisfacción.",
                        "monitoreo_periodicidad": "Al momento de la entrega e instalación en sitio."
                    },
                    {
                        "clase": "Específico", "fuente": "Interno", "etapa": "Ejecución", "tipo": "Técnico",
                        "descripcion": "Entrega o instalación de dispositivos GPS, cámaras ADAS/DMS, tarjetas MicroSD o periféricos de cabina con defectos de fábrica, fallas de cableado automotriz o disconformidad frente a la propuesta técnica.",
                        "consecuencia": "Fallas en la transmisión telemétrica o grabación de video vehicular y retraso en la entrada en operación regular.",
                        "probabilidad_inicial": 2, "impacto_inicial": 4,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Pruebas técnicas de funcionamiento y conectividad de cada equipo instalado. Reemplazo y corrección inmediata a costa del contratista en un plazo máximo de 48 horas bajo garantía de fábrica.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Pruebas de encendido, señal GPS, visualización de cámaras y alertas ADAS/DMS en vivo.",
                        "monitoreo_periodicidad": "Al momento de las pruebas técnicas y entrega."
                    },
                    {
                        "clase": "Específico", "fuente": "Externo", "etapa": "Ejecución", "tipo": "Tecnológico",
                        "descripcion": "Intermitencia, caída del servicio, pérdida de transmisión de datos satelitales o indisponibilidad de la plataforma web y aplicación móvil de monitoreo 24/7.",
                        "consecuencia": "Pérdida temporal de trazabilidad de los vehículos y retraso en el reporte de alertas de seguridad vial e itinerarios.",
                        "probabilidad_inicial": 2, "impacto_inicial": 3,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Acuerdo de nivel de servicio (SLA) con alta disponibilidad (99.5%), canales de soporte técnico permanente y almacenamiento local redundante en tarjetas MicroSD ante fallas de cobertura de red móvil.",
                        "probabilidad_residual": 1, "impacto_residual": 2,
                        "afecta_ejecucion": "Sí",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Supervisión continua de disponibilidad de plataforma web y reportes periódicos de conectividad.",
                        "monitoreo_periodicidad": "Seguimiento permanente durante el periodo contratado."
                    },
                    {
                        "clase": "General", "fuente": "Externo", "etapa": "Ejecución", "tipo": "Económico",
                        "descripcion": "Fluctuación o incremento sobreviniente en precios de equipos tecnológicos, tasas de cambio, costos de instalación o viáticos de desplazamiento al municipio de Yolombó.",
                        "consecuencia": "Riesgo de reclamaciones económicas o solicitudes de reajuste del valor pactado.",
                        "probabilidad_inicial": 2, "impacto_inicial": 3,
                        "asignacion": "Contratista (100%)",
                        "tratamiento": "Pactar valor total fijo y firme a todo costo en pesos colombianos, asumiendo el contratista los costos directos, indirectos, transporte y viáticos sin lugar a reajustes.",
                        "probabilidad_residual": 1, "impacto_residual": 1,
                        "afecta_ejecucion": "No",
                        "responsable": "Contratista y Supervisor ESPY",
                        "monitoreo_mecanismo": "Verificación de facturación electrónica conforme a la propuesta económica fija adjudicada.",
                        "monitoreo_periodicidad": "Al momento de la radicación de la factura para pago."
                    }
                ]

            return [
                {
                    "clase": "Específico", "fuente": "Externo", "etapa": "Ejecución", "tipo": "Operacional",
                    "descripcion": "Retraso o mora en la entrega física de los bienes, materiales o insumos en el almacén de la ESPY en Yolombó.",
                    "consecuencia": "Desabastecimiento operativo y retraso en mantenimientos de acueducto o alcantarillado.",
                    "probabilidad_inicial": 3, "impacto_inicial": 3,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Pactar plazos perentorios de entrega y verificación de existencias antes de la orden. Cláusula penal y multas por mora.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista y Supervisor ESPY",
                    "monitoreo_mecanismo": "Control de guías de despacho y acta de ingreso a bodega.",
                    "monitoreo_periodicidad": "En cada entrega programada."
                },
                {
                    "clase": "Específico", "fuente": "Interno", "etapa": "Ejecución", "tipo": "Técnico",
                    "descripcion": "Entrega de bienes o repuestos defectuosos, averiados o con especificaciones técnicas disconformes con la ficha técnica.",
                    "consecuencia": "Rechazo del producto, fallas en la red y retraso en la reposición.",
                    "probabilidad_inicial": 2, "impacto_inicial": 4,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Inspección técnica detallada al recibo. Obligación de reemplazo en un plazo máximo de 48 horas sin costo adicional.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista y Almacén / Supervisor ESPY",
                    "monitoreo_mecanismo": "Muestreo y verificación contra ficha técnica oficial.",
                    "monitoreo_periodicidad": "Al momento de la entrega física."
                },
                {
                    "clase": "General", "fuente": "Externo", "etapa": "Ejecución", "tipo": "Económico",
                    "descripcion": "Fluctuación o incremento sobreviniente en precios de insumos, fletes de transporte terrestre o tasas de cambio.",
                    "consecuencia": "Riesgo de desequilibrio económico o solicitud de reajuste presupuestal por parte del proveedor.",
                    "probabilidad_inicial": 3, "impacto_inicial": 2,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Pactar precios unitarios fijos y firmes en pesos colombianos hasta la liquidación contractual.",
                    "probabilidad_residual": 1, "impacto_residual": 2,
                    "afecta_ejecucion": "No",
                    "responsable": "Contratista y Supervisor ESPY",
                    "monitoreo_mecanismo": "Verificación de facturación contra cotización aprobada.",
                    "monitoreo_periodicidad": "Mensual (con la presentación de factura)."
                }
            ]
        else: # OBRA
            return [
                {
                    "clase": "Específico", "fuente": "Externo", "etapa": "Ejecución", "tipo": "De la naturaleza",
                    "descripcion": "Condiciones climáticas adversas (lluvias intensas en Yolombó) o geológicas imprevistas que afecten el terreno de excavación.",
                    "consecuencia": "Demoras en cronograma de obra y posible necesidad de prórroga justificada.",
                    "probabilidad_inicial": 3, "impacto_inicial": 3,
                    "asignacion": "Compartido (ESPY 50% / Contratista 50%)",
                    "tratamiento": "Programación técnica considerando régimen pluviométrico local y protocolos de drenaje y protección de zanjas.",
                    "probabilidad_residual": 2, "impacto_residual": 2,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista e Interventor / Supervisor ESPY",
                    "monitoreo_mecanismo": "Seguimiento en bitácora de obra y comités técnicos.",
                    "monitoreo_periodicidad": "Quincenal."
                },
                {
                    "clase": "General", "fuente": "Interno", "etapa": "Ejecución", "tipo": "Operacional / SST",
                    "descripcion": "Accidentes graves o fatales en excavaciones o zanjas por falta de entibados o incumplimiento del SG-SST.",
                    "consecuencia": "Paralización judicial de la obra, sanciones del Ministerio de Trabajo y daño a trabajadores.",
                    "probabilidad_inicial": 2, "impacto_inicial": 5,
                    "asignacion": "Contratista (100%)",
                    "tratamiento": "Exigencia obligatoria de ARL Riesgo V, SISCONTRATA, coordinador de alturas/espacios confinados y entibados normativos.",
                    "probabilidad_residual": 1, "impacto_residual": 3,
                    "afecta_ejecucion": "Sí",
                    "responsable": "Contratista e Inspector SST",
                    "monitoreo_mecanismo": "Inspecciones de campo no anunciadas y firmas de permisos de trabajo.",
                    "monitoreo_periodicidad": "Quincenal y al presentarse el evento."
                }
            ]

    def generate_risk_matrix(
        self, 
        input_data: RiskMatrixInput, 
        db: Optional[Session] = None, 
        override_api_key: Optional[str] = None
    ) -> RiskMatrixResult:
        """
        Orquesta la tipificación, evaluación y asignación rigurosa de riesgos contractuales.
        """
        # 1. Consolidar contexto de BD y documentos
        context = self.extract_context_from_contract_and_docs(input_data, db)
        
        # 2. Paso 1 Obligatorio: Evaluar insumos indispensables y advertir vacíos
        advertencias, observaciones = self.validate_inputs_and_detect_gaps(context)

        # 3. Normalizar objeto y datos básicos
        objeto_final = context.get("objeto") or "PRESTACIÓN DE SERVICIOS PARA EL APOYO A LA GESTIÓN DE LAS ACTIVIDADES COMPLEMENTARIAS DEL PGIRS DEL MUNICIPIO DE YOLOMBÓ"
        # Sanear objeto de erratas comunes
        objeto_final = re.sub(r"\bgenti[oó]n\b", "GESTIÓN", objeto_final, flags=re.IGNORECASE)
        objeto_final = re.sub(r"\bconplementarias\b", "COMPLEMENTARIAS", objeto_final, flags=re.IGNORECASE)
        objeto_final = objeto_final.strip().upper()

        tipologia_final = context.get("tipologia") or "PRESTACION_SERVICIOS"

        # 4. Invocar LLM para generar matriz de riesgos estructurada
        client = gemini_service._get_client(override_api_key)
        raw_risks = []

        if client:
            system_instruction = (
                "Eres el 'Especialista en Gestión y Asignación de Riesgos Previsibles' para la Empresa de Servicios Públicos de Yolombó (ESPY S.A. E.S.P.), Antioquia.\n"
                "Tu misión es estructurar, evaluar y documentar matrices de riesgos contractuales rigurosas, técnicamente sustentadas y sin inconsistencias aritméticas ni conceptuales, conforme a la metodología de Colombia Compra Eficiente (CCE) y el régimen de empresas de servicios públicos (Ley 142 de 1994 y Manual de Contratación ESPY).\n\n"
                "REGLAS CRÍTICAS DE CALIDAD Y CONSISTENCIA:\n"
                "1. Verificación aritmética absoluta: Valoración Inicial = Probabilidad Inicial (1 a 5) + Impacto Inicial (1 a 5). Valoración Residual = Probabilidad Residual (1 a 5) + Impacto Residual (1 a 5).\n"
                "   Categoría Inicial y Residual: Bajo (2, 3 o 4), Medio (5), Alto (6 o 7), Extremo (8, 9 o 10).\n"
                "2. Consecutivo continuo: Del 1 en adelante sin saltos.\n"
                "3. Adecuación estricta al tipo contractual (No copy-paste):\n"
                "   - Si es PRESTACION DE SERVICIOS personales/apoyo a la gestión: Enfocar en calidad técnica de entregables, oportunidad de informes, cumplimiento de perfiles, SST (Operacional / SST), seguridad social y confidencialidad. PROHIBIDO incluir referencias a suministro de bienes, materiales de obra o subalternos a cargo.\n"
                "   - Si es SUMINISTRO: Oportunidad de entrega, garantías de fábrica, transporte, averías e indexación de precios.\n"
                "   - Si es OBRA: Estabilidad, geología y lluvias en Yolombó, manejo ambiental, permisos y contingencias de redes.\n"
                "4. Seguridad y Salud en el Trabajo: Tipificar obligatoriamente como 'Operacional / SST' o 'Laboral'. Jamás como 'Económico'.\n"
                "5. Periodicidad de supervisión realista: 'En etapa precontractual', 'Quincenal', 'Mensual (con cuenta de cobro e informe)', 'Al presentarse el evento'. Evitar 'Diario'.\n"
                "6. Redacción institucional impecable.\n"
                "Devuelve una lista JSON de objetos de riesgo con los campos requeridos."
            )

            prompt_user = (
                f"Estructura la Matriz de Riesgos CCE para el siguiente proceso contractual de la ESPY S.A. E.S.P.:\n"
                f"- Objeto: {objeto_final}\n"
                f"- Tipología: {tipologia_final}\n"
                f"- Contratista: {context.get('contratista_nombre') or 'Por definir'}\n"
                f"- Presupuesto oficial: {format_cop_standard(context.get('valor_estimado', 0))}\n"
                f"- Plazo: {context.get('plazo_texto') or '2 meses'}\n"
                f"- Lugar de ejecución: {context.get('lugar_ejecucion')}\n"
                f"- Exigencia de pólizas: {'No se exigen pólizas contractuales' if not context.get('exige_polizas') else 'Se exigen pólizas de garantía'}\n\n"
                f"Texto o documentos de contexto del proceso:\n{context.get('documentos_texto', '')[:40000]}\n"
            )

            class _RiskListWrapper(docx.Document if False else dict): # dummy for typing
                pass

            try:
                logger.info("Invocando Gemini API para estructuración de matriz de riesgos contractuales...")
                from pydantic import BaseModel
                class RiskListContainer(BaseModel):
                    riesgos: List[RiskItemSchema]

                resp = client.models.generate_content(
                    model=settings.GEMINI_MODEL,
                    contents=prompt_user,
                    config={
                        "system_instruction": system_instruction,
                        "response_mime_type": "application/json",
                        "response_schema": RiskListContainer,
                        "temperature": 0.1,
                    }
                )
                if resp.text:
                    parsed_res = json.loads(resp.text)
                    raw_risks = parsed_res.get("riesgos", [])
            except Exception as e:
                logger.warning(f"Error generando matriz de riesgos con Gemini LLM ({e}). Empleando generador normativo...")
                raw_risks = self.fallback_risk_matrix(context)
        else:
            raw_risks = self.fallback_risk_matrix(context)

        if not raw_risks:
            raw_risks = self.fallback_risk_matrix(context)

        # 5. Normalizar y verificar rigurosamente los riesgos
        normalized_risks = self.normalize_risk_items(
            raw_risks, 
            tipologia_final,
            objeto_text=objeto_final,
            plazo_str=context.get("plazo_texto")
        )

        # 6. Calcular resúmenes estadísticos
        cat_ini_counts = {"Bajo": 0, "Medio": 0, "Alto": 0, "Extremo": 0}
        cat_res_counts = {"Bajo": 0, "Medio": 0, "Alto": 0, "Extremo": 0}

        for r in normalized_risks:
            cat_ini_counts[r.categoria_inicial] = cat_ini_counts.get(r.categoria_inicial, 0) + 1
            cat_res_counts[r.categoria_residual] = cat_res_counts.get(r.categoria_residual, 0) + 1

        result = RiskMatrixResult(
            agente_nombre=self.nombre,
            objeto=objeto_final,
            tipologia=tipologia_final,
            contratista_nombre=context.get("contratista_nombre"),
            valor_estimado=context.get("valor_estimado"),
            plazo_texto=context.get("plazo_texto"),
            lugar_ejecucion=context.get("lugar_ejecucion"),
            exige_polizas=bool(context.get("exige_polizas")),
            insumos_faltantes_advertencias=advertencias,
            observaciones_estudio_previo=observaciones,
            riesgos=normalized_risks,
            total_riesgos=len(normalized_risks),
            resumen_categorias_inicial=cat_ini_counts,
            resumen_categorias_residual=cat_res_counts,
            custom_logo_path=context.get("custom_logo_path")
        )

        return result

    def resolve_logo_paths(self, result: RiskMatrixResult) -> Tuple[Path, Path]:
        """Resuelve la ruta del logo y la marca de agua a utilizar en el documento."""
        root_dir = Path(settings.WORKSPACE_ROOT)
        default_logo = root_dir / "frontend" / "assets" / "logo.png"
        if not default_logo.exists():
            default_logo = root_dir / "scratch" / "custom_image1.png"
        
        default_wm = root_dir / "scratch" / "custom_image2.png"
        if not default_wm.exists():
            default_wm = default_logo

        custom_logo = None
        if result.custom_logo_path:
            p = Path(result.custom_logo_path)
            if not p.is_absolute():
                candidate = root_dir / p
                if candidate.exists():
                    p = candidate
                else:
                    p = settings.BASE_DIR / p
            if p.exists():
                custom_logo = p

        logo_to_use = custom_logo if custom_logo else default_logo
        wm_to_use = custom_logo if custom_logo else default_wm
        return logo_to_use, wm_to_use

    def generate_docx(self, result: RiskMatrixResult, output_path: Path) -> Path:
        """
        Genera el documento oficial en formato DOCX Horizontal (Landscape) con membrete institucional,
        logo corporativo, objeto contractual completo y tablas estilizadas CCE.
        """
        doc = Document()

        # Configurar Orientación Horizontal (Landscape A4/Letter)
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.0)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

        logo_file, wm_file = self.resolve_logo_paths(result)
        TEAL_RGB = RGBColor(11, 60, 93)
        HEADER_BG = "1A365D"

        def set_cell_background(cell, hex_color):
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading)

        def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
            )
            tcPr.append(tcMar)

        def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
            tblPr = table._tbl.tblPr
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'</w:tblBorders>'
            )
            tblPr.append(borders)

        # 1. Encabezado institucional con Logo en el Header de Word
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_head.paragraph_format.space_after = Pt(2)
        if logo_file and logo_file.exists():
            try:
                p_head.add_run().add_picture(str(logo_file), width=Inches(1.5))
            except Exception as e:
                logger.warning(f"Aviso agregando imagen a header docx: {e}")

        # 2. Pie de página institucional repetitivo en todas las páginas
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.paragraph_format.space_before = Pt(4)
        r_f1 = p_foot.add_run("EMPRESA DE SERVICIOS PUBLICOS DE YOLOMBÓ – NIT. 811015112-4\n")
        r_f1.bold = True
        r_f1.font.size = Pt(7.5)
        r_f1.font.color.rgb = RGBColor(15, 23, 42)
        r_f2 = p_foot.add_run("CELULAR: 3206976172\n")
        r_f2.bold = True
        r_f2.font.size = Pt(7.5)
        r_f2.font.color.rgb = RGBColor(15, 23, 42)
        r_f3 = p_foot.add_run("serviciospublicosyol@hotmail.com\n")
        r_f3.bold = True
        r_f3.font.size = Pt(7.5)
        r_f3.font.color.rgb = RGBColor(2, 132, 199)
        r_f4 = p_foot.add_run("Calle Colombia 20 #20 - 74 Yolombó – Antioquia")
        r_f4.font.size = Pt(7.5)
        r_f4.font.color.rgb = RGBColor(71, 85, 105)

        # Barra decorativa superior del cuerpo (Línea Azul y Verde)
        bar_table = doc.add_table(rows=1, cols=2)
        bar_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(bar_table, color="FFFFFF", sz="0", val="none")
        c_blue, c_green = bar_table.rows[0].cells[0], bar_table.rows[0].cells[1]
        c_blue.width = Inches(8.3)
        c_green.width = Inches(1.9)
        set_cell_background(c_blue, "0284C7")
        set_cell_background(c_green, "16A34A")
        set_cell_margins(c_blue, top=15, bottom=15, left=0, right=0)
        set_cell_margins(c_green, top=15, bottom=15, left=0, right=0)
        c_blue.paragraphs[0].paragraph_format.space_after = Pt(0)
        c_green.paragraphs[0].paragraph_format.space_after = Pt(0)

        doc.add_paragraph().paragraph_format.space_after = Pt(3)

        # 3. Caja informativa del proceso (PROCESO CONTRACTUAL / OBJETO)
        info_table = doc.add_table(rows=4, cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(info_table, color="CBD5E1", sz="4")

        valor_txt = format_cop_standard(result.valor_estimado)
        if result.exige_polizas:
            polizas_txt = "Se exige constitución de pólizas de garantía"
        elif result.tipologia == "SUMINISTRO":
            polizas_txt = "No se exige constitución de pólizas (pago contra entrega física e instalación efectiva de los equipos, previa suscripción del Acta de Entrada a Almacén y recibo a satisfacción por el supervisor)"
        else:
            polizas_txt = "No se exige constitución de pólizas (Pago contra prestación efectiva y recibido a satisfacción)"

        info_data = [
            ("PROCESO CONTRACTUAL / OBJETO:", result.objeto),
            ("CONTRATISTA / TIPO DE CONTRATO:", f"{result.contratista_nombre or 'Contratista Seleccionado'}  |  {result.tipologia.replace('_', ' ')}"),
            ("PRESUPUESTO OFICIAL ESTIMADO:", f"{valor_txt}  |  Lugar: {result.lugar_ejecucion or 'Municipio de Yolombó, Antioquia'}"),
            ("PLAZO DE EJECUCIÓN Y PÓLIZAS:", f"Plazo: {result.plazo_texto or 'Conforme a cronograma'}  |  Garantías: {polizas_txt}")
        ]

        col_widths_info = [Inches(2.2), Inches(8.0)]
        for row_idx, (label, val) in enumerate(info_data):
            row = info_table.rows[row_idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = col_widths_info[0]
            c1.width = col_widths_info[1]
            set_cell_background(c0, "F8FAFC")
            set_cell_margins(c0, top=35, bottom=35, left=60, right=60)
            set_cell_margins(c1, top=35, bottom=35, left=60, right=60)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(label)
            r0.bold = True
            r0.font.size = Pt(7.5)
            r0.font.color.rgb = RGBColor(30, 41, 59)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(val)
            r1.font.size = Pt(7.5)
            r1.font.color.rgb = RGBColor(15, 23, 42)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 4. Tabla resumen de criterios CCE
        p_crit_title = doc.add_paragraph()
        p_crit_title.paragraph_format.space_after = Pt(2)
        r_crit = p_crit_title.add_run("CRITERIOS DE VALORACIÓN Y CALIFICACIÓN (METODOLOGÍA COLOMBIA COMPRA EFICIENTE)")
        r_crit.bold = True
        r_crit.font.size = Pt(8.0)
        r_crit.font.color.rgb = RGBColor(15, 23, 42)

        crit_table = doc.add_table(rows=2, cols=3)
        crit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(crit_table, color="CBD5E1", sz="4")

        crit_headers = ["ESCALA DE PROBABILIDAD (1 a 5)", "ESCALA DE IMPACTO (1 a 5)", "VALORACIÓN Y CATEGORIZACIÓN (P + I)"]
        c_row0 = crit_table.rows[0]
        for i, h_text in enumerate(crit_headers):
            cell = c_row0.cells[i]
            set_cell_background(cell, HEADER_BG)
            set_cell_margins(cell, top=30, bottom=30, left=45, right=45)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(7.0)
            run.font.color.rgb = RGBColor(255, 255, 255)

        crit_contents = [
            "1 = Raro (Excepcional)\n2 = Improbable (Ocasional)\n3 = Posible (Cualquier momento)\n4 = Probable (Alta probabilidad)\n5 = Casi Cierto (Mayoría de circunstancias)",
            "1 = Insignificante (< 1% del valor)\n2 = Menor (1% al 5% del valor)\n3 = Moderado (5% al 15% del valor)\n4 = Mayor (15% al 30% del valor)\n5 = Catastrófico (> 30% o incumplimiento total)",
            "• BAJO: Suma de 2, 3 o 4\n• MEDIO: Suma de 5\n• ALTO: Suma de 6 o 7\n• EXTREMO: Suma de 8, 9 o 10\nRegla: Valoración = Probabilidad + Impacto"
        ]
        c_row1 = crit_table.rows[1]
        crit_widths = [Inches(3.4), Inches(3.4), Inches(3.4)]
        for i, text in enumerate(crit_contents):
            cell = c_row1.cells[i]
            cell.width = crit_widths[i]
            set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, top=30, bottom=30, left=45, right=45)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(6.8)
            run.font.color.rgb = RGBColor(30, 41, 59)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 5. Tabla de 14 Columnas de la Matriz de Riesgos
        headers_top = [
            "N°", "Clase", "Fuente", "Etapa", "Tipo",
            "Descripción del Riesgo", "Consecuencia de la Ocurrencia",
            "EVALUACIÓN INICIAL",
            "Asignación",
            "Tratamiento / Medida de Control",
            "EVALUACIÓN RESIDUAL",
            "¿Afecta?", "Responsable", "Monitoreo y Revisión"
        ]

        col_widths = [
            Inches(0.32), Inches(0.48), Inches(0.45), Inches(0.52), Inches(0.65),
            Inches(1.50), Inches(1.25), Inches(0.80), Inches(0.65), Inches(1.48),
            Inches(0.80), Inches(0.38), Inches(0.65), Inches(0.87)
        ]

        table = doc.add_table(rows=1 + len(result.riesgos), cols=14)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="CBD5E1", sz="4")

        # Configurar repetición de encabezado y no división de filas
        hdr_row = table.rows[0]
        hdr_tr = hdr_row._tr.get_or_add_trPr()
        hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        hdr_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for idx, text in enumerate(headers_top):
            cell = hdr_row.cells[idx]
            cell.width = col_widths[idx]
            set_cell_background(cell, HEADER_BG)
            set_cell_margins(cell, top=40, bottom=40, left=30, right=30)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(6.2)
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Detectar condiciones del proceso para blindaje en documento Word
        obj_lower = (result.objeto or "").lower()
        is_mantenimiento = any(k in obj_lower for k in ["mantenimiento", "herramienta", "equipo", "electromec", "reparaci", "taller", "eléctric", "electric"])
        es_corto_plazo = is_short_term_period(plazo_str=result.plazo_texto)

        # Data rows
        for r_idx, r_item in enumerate(result.riesgos, start=1):
            row = table.rows[r_idx]
            row_tr = row._tr.get_or_add_trPr()
            row_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            # Blindaje en DOCX: Sanitizar campos por si venían de caché o LLM previo
            r_desc = r_item.descripcion
            r_cons = r_item.consecuencia
            r_trat = r_item.tratamiento
            r_mec = r_item.monitoreo_mecanismo
            r_per = r_item.monitoreo_periodicidad
            r_tipo = r_item.tipo

            if is_mantenimiento:
                if "corantioquia" in r_cons.lower() or "sspd" in r_cons.lower():
                    r_cons = re.sub(r"ante\s+autoridades\s+ambientales\s+competentes\s*\([^)]*\)", "en la operación de los servicios públicos de la ESPY", r_cons, flags=re.IGNORECASE)
                    r_cons = re.sub(r"\(?Corantioquia/SSPD\)?", "la ESPY", r_cons, flags=re.IGNORECASE)
                if "separación en la fuente" in r_desc.lower() or "renuencia comunitaria" in r_desc.lower() or "apatía" in r_desc.lower():
                    r_tipo = "Operacional / Custodia"
                    r_desc = "Daño, extravío o deterioro de herramientas o equipos propiedad de la ESPY durante su traslado, custodia o intervención en taller."
                    r_cons = "Afectación al patrimonio público de la ESPY y retraso en la disponibilidad de equipos para atender contingencias operativas."
                    r_trat = "Levantamiento de inventario y estado inicial al entregar y recibir los equipos mediante acta suscrita por el supervisor. El contratista asume a todo costo la reposición inmediata de cualquier elemento dañado o extraviado."
                    r_mec = "Inspección técnica detallada y verificación física contra acta de entrega."
                    r_per = "Al inicio, durante la ejecución y en la liquidación/pago único" if es_corto_plazo else "Al momento de cada entrega de equipos."

            if es_corto_plazo:
                r_desc = re.sub(
                    r"informes\s+mensuales\s+de\s+actividades\s+y\s+evidencias\s+de\s+ejecuci[oó]n[^.]*dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes",
                    "entrega final e informe técnico dentro del plazo pactado",
                    r_desc,
                    flags=re.IGNORECASE
                )
                r_desc = re.sub(r"informes\s+mensuales\s+dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes", "entrega final e informe técnico dentro del plazo pactado", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\binformes\s+mensuales\b", "informe final de actividades", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de pago único", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\bpago\s+mensual\b", "pago único", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\bvalor\s+mensualizado\b", "valor total", r_desc, flags=re.IGNORECASE)

                r_cons = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de liquidación y pago", r_cons, flags=re.IGNORECASE)
                r_cons = re.sub(r"\bpago\s+mensual\b", "pago único", r_cons, flags=re.IGNORECASE)

                r_trat = re.sub(r"\ba\s+mes\s+vencido\b", "contra entrega final a satisfacción", r_trat, flags=re.IGNORECASE)
                r_trat = re.sub(r"\btr[aá]mite\s+mensual\b", "trámite de pago", r_trat, flags=re.IGNORECASE)

                if "mensual" in r_per.lower() or "quincenal" in r_per.lower() or "diario" in r_per.lower() or "permanente" in r_per.lower():
                    r_per = "Al inicio, durante la ejecución y en la liquidación/pago único"

            eval_ini_str = f"P:{r_item.probabilidad_inicial}  I:{r_item.impacto_inicial}\nVal:{r_item.valoracion_inicial}\n{r_item.categoria_inicial.upper()}"
            eval_res_str = f"P:{r_item.probabilidad_residual}  I:{r_item.impacto_residual}\nVal:{r_item.valoracion_residual}\n{r_item.categoria_residual.upper()}"
            monitoreo_str = f"{r_mec}\nPeriodicidad: {r_per}"

            row_values = [
                (str(r_item.numero), WD_ALIGN_PARAGRAPH.CENTER, True),
                (r_item.clase, WD_ALIGN_PARAGRAPH.CENTER, False),
                (r_item.fuente, WD_ALIGN_PARAGRAPH.CENTER, False),
                (r_item.etapa, WD_ALIGN_PARAGRAPH.CENTER, False),
                (r_tipo, WD_ALIGN_PARAGRAPH.CENTER, True),
                (r_desc, WD_ALIGN_PARAGRAPH.LEFT, False),
                (r_cons, WD_ALIGN_PARAGRAPH.LEFT, False),
                (eval_ini_str, WD_ALIGN_PARAGRAPH.CENTER, True),
                (r_item.asignacion, WD_ALIGN_PARAGRAPH.CENTER, False),
                (r_trat, WD_ALIGN_PARAGRAPH.LEFT, False),
                (eval_res_str, WD_ALIGN_PARAGRAPH.CENTER, True),
                (r_item.afecta_ejecucion, WD_ALIGN_PARAGRAPH.CENTER, True),
                (r_item.responsable, WD_ALIGN_PARAGRAPH.CENTER, False),
                (monitoreo_str, WD_ALIGN_PARAGRAPH.LEFT, False)
            ]

            for c_idx, (val, align, is_bold) in enumerate(row_values):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                # Color highlighting for categories
                if c_idx == 7:  # Eval Inicial
                    if r_item.categoria_inicial == "Extremo":
                        set_cell_background(cell, "FCA5A5")
                    elif r_item.categoria_inicial == "Alto":
                        set_cell_background(cell, "FEE2E2")
                    elif r_item.categoria_inicial == "Medio":
                        set_cell_background(cell, "FEF3C7")
                    else:
                        set_cell_background(cell, "ECFDF5")
                elif c_idx == 10:  # Eval Residual
                    if r_item.categoria_residual == "Bajo":
                        set_cell_background(cell, "ECFDF5")
                    elif r_item.categoria_residual == "Medio":
                        set_cell_background(cell, "FEF3C7")
                    else:
                        set_cell_background(cell, "FEE2E2")
                else:
                    set_cell_background(cell, "FFFFFF")

                set_cell_margins(cell, top=30, bottom=30, left=30, right=30)
                p = cell.paragraphs[0]
                p.alignment = align
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.size = Pt(5.8)
                run.bold = is_bold
                run.font.color.rgb = RGBColor(15, 23, 42)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 6. Observaciones Técnicas y Recomendaciones
        if result.observaciones_estudio_previo:
            p_obs_title = doc.add_paragraph()
            p_obs_title.paragraph_format.space_after = Pt(2)
            r_obs = p_obs_title.add_run("OBSERVACIONES TÉCNICAS Y CONTROL PREVENTIVO DE RIESGOS (ESTUDIO PREVIO):")
            r_obs.bold = True
            r_obs.font.size = Pt(8.0)
            r_obs.font.color.rgb = RGBColor(15, 23, 42)

            obs_table = doc.add_table(rows=1, cols=1)
            obs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(obs_table, color="CBD5E1", sz="4")
            cell_obs = obs_table.rows[0].cells[0]
            cell_obs.width = Inches(10.2)
            set_cell_background(cell_obs, "F8FAFC")
            set_cell_margins(cell_obs, top=35, bottom=35, left=60, right=60)

            p_o = cell_obs.paragraphs[0]
            p_o.paragraph_format.space_after = Pt(0)
            full_obs_text = "\n".join([f"• {obs}" for obs in result.observaciones_estudio_previo])
            r_o = p_o.add_run(full_obs_text)
            r_o.font.size = Pt(6.8)
            r_o.font.color.rgb = RGBColor(51, 65, 85)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 7. Firma Oficial
        sig_p = doc.add_paragraph()
        sig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sig_p.paragraph_format.space_before = Pt(8)
        sig_p.paragraph_format.space_after = Pt(0)
        r_line = sig_p.add_run("________________________________________\n")
        r_line.font.size = Pt(8.5)
        r_line.font.color.rgb = RGBColor(100, 116, 139)
        r_name = sig_p.add_run("CAROLINA GIRALDO GÓMEZ\n")
        r_name.bold = True
        r_name.font.size = Pt(8.0)
        r_name.font.color.rgb = RGBColor(15, 23, 42)
        r_pos = sig_p.add_run("Gerente General")
        r_pos.font.size = Pt(7.5)
        r_pos.font.color.rgb = RGBColor(71, 85, 105)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"Matriz de riesgos en formato DOCX generada en: {output_path}")
        return output_path

    def generate_pdf(self, result: RiskMatrixResult, output_path: Path) -> Path:
        """
        Genera el documento oficial en formato PDF Horizontal (Landscape) de alta fidelidad visual,
        con membrete corporativo, marcas de agua, barras decorativas, objeto contractual completo
        y tablas reglamentarias Colombia Compra Eficiente (CCE).
        """
        logo_file, wm_file = self.resolve_logo_paths(result)
        page_width, page_height = landscape(letter)

        class NumberedCanvas(canvas.Canvas):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self._saved_page_states = []

            def showPage(self):
                self._saved_page_states.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                num_pages = len(self._saved_page_states)
                for state in self._saved_page_states:
                    self.__dict__.update(state)
                    self.draw_page_decorations(num_pages)
                    super().showPage()
                super().save()

            def draw_page_decorations(self, total_pages):
                self.saveState()
                
                # 1. Marca de agua suave en el centro de la página
                if wm_file and wm_file.exists():
                    try:
                        self.saveState()
                        self.setFillAlpha(0.06)
                        wm_w, wm_h = 320, 320
                        wm_x = (page_width - wm_w) / 2
                        wm_y = (page_height - wm_h) / 2
                        self.drawImage(str(wm_file), wm_x, wm_y, width=wm_w, height=wm_h, mask='auto')
                        self.restoreState()
                    except Exception as e:
                        logger.warning(f"Aviso dibujando marca de agua en PDF: {e}")

                # 2. Encabezado con Logo oficial en la esquina superior izquierda
                if logo_file and logo_file.exists():
                    try:
                        logo_w = 110
                        logo_h = 56
                        logo_x = 28
                        logo_y = page_height - logo_h - 18
                        self.drawImage(str(logo_file), logo_x, logo_y, width=logo_w, height=logo_h, mask='auto')
                    except Exception as e:
                        logger.warning(f"Aviso dibujando logo en PDF: {e}")

                # 3. Línea decorativa dual institucional (Azul ESPY y Verde ESPY)
                line_y = page_height - 78
                self.setLineWidth(2.5)
                self.setStrokeColor(colors.HexColor("#0284C7"))
                self.line(28, line_y, page_width - 160, line_y)
                
                self.setStrokeColor(colors.HexColor("#16A34A"))
                self.line(page_width - 158, line_y, page_width - 28, line_y)

                # 4. Pie de página institucional centrado
                footer_y = 48
                self.setFont("Helvetica-Bold", 7.5)
                self.setFillColor(colors.HexColor("#0F172A"))
                self.drawCentredString(page_width / 2, footer_y, "EMPRESA DE SERVICIOS PUBLICOS DE YOLOMBÓ – NIT. 811015112-4")
                self.drawCentredString(page_width / 2, footer_y - 10, "CELULAR: 3206976172")
                self.setFillColor(colors.HexColor("#0284C7"))
                self.drawCentredString(page_width / 2, footer_y - 20, "serviciospublicosyol@hotmail.com")
                self.setFont("Helvetica", 7.5)
                self.setFillColor(colors.HexColor("#334155"))
                self.drawCentredString(page_width / 2, footer_y - 30, "Calle Colombia 20 #20 - 74 Yolombó – Antioquia")

                self.restoreState()

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=landscape(letter),
            leftMargin=28,
            rightMargin=28,
            topMargin=84,
            bottomMargin=58
        )

        styles = getSampleStyleSheet()
        
        style_label = ParagraphStyle('InfoLabel', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=6.8, leading=8.5, textColor=colors.HexColor("#1E293B"))
        style_val = ParagraphStyle('InfoVal', parent=styles['Normal'], fontName='Helvetica', fontSize=6.8, leading=8.5, textColor=colors.HexColor("#0F172A"))
        style_crit_hdr = ParagraphStyle('CritHdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=6.8, leading=8.5, alignment=1, textColor=colors.white)
        style_crit_body = ParagraphStyle('CritBody', parent=styles['Normal'], fontName='Helvetica', fontSize=6.2, leading=8.0, textColor=colors.HexColor("#1E293B"))
        style_tbl_hdr = ParagraphStyle('TblHdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=5.5, leading=7.0, alignment=1, textColor=colors.white)
        style_tbl_cell = ParagraphStyle('TblCell', parent=styles['Normal'], fontName='Helvetica', fontSize=5.5, leading=7.0, textColor=colors.HexColor("#0F172A"))
        style_tbl_cell_center = ParagraphStyle('TblCellCenter', parent=styles['Normal'], fontName='Helvetica', fontSize=5.5, leading=7.0, alignment=1, textColor=colors.HexColor("#0F172A"))
        style_tbl_cell_bold_center = ParagraphStyle('TblCellBoldCenter', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=5.5, leading=7.0, alignment=1, textColor=colors.HexColor("#0F172A"))
        crit_title_style = ParagraphStyle('CritTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=6.8, leading=8.5, textColor=colors.HexColor("#0F172A"))
        obs_title_style = ParagraphStyle('ObsTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=6.8, leading=8.5, textColor=colors.HexColor("#0F172A"))

        story = []

        # 1. Tabla de metadatos del proceso (PROCESO CONTRACTUAL / OBJETO)
        valor_txt = format_cop_standard(result.valor_estimado)
        polizas_txt = "Se exige constitución de pólizas de garantía" if result.exige_polizas else "No se exige constitución de pólizas (Pago contra prestación efectiva y recibido a satisfacción)"
        contratista_txt = f"{result.contratista_nombre or 'Contratista Seleccionado'}  |  {result.tipologia.replace('_', ' ')}"
        pto_txt = f"{valor_txt}  |  Lugar: {result.lugar_ejecucion or 'Municipio de Yolombó, Antioquia'}"
        plazo_txt = f"Plazo: {result.plazo_texto or 'Conforme a cronograma'}  |  Garantías: {polizas_txt}"

        info_data = [
            [Paragraph("PROCESO CONTRACTUAL / OBJETO:", style_label), Paragraph(result.objeto, style_val)],
            [Paragraph("CONTRATISTA / TIPO DE CONTRATO:", style_label), Paragraph(contratista_txt, style_val)],
            [Paragraph("PRESUPUESTO OFICIAL ESTIMADO:", style_label), Paragraph(pto_txt, style_val)],
            [Paragraph("PLAZO DE EJECUCIÓN Y PÓLIZAS:", style_label), Paragraph(plazo_txt, style_val)],
        ]
        
        col_w_info = [155, page_width - 56 - 155]
        t_info = Table(info_data, colWidths=col_w_info)
        t_info.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#F8FAFC")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING', (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(t_info)
        story.append(Spacer(1, 6))

        # 2. Criterios CCE
        story.append(Paragraph("CRITERIOS DE VALORACIÓN Y CALIFICACIÓN (METODOLOGÍA COLOMBIA COMPRA EFICIENTE)", crit_title_style))
        story.append(Spacer(1, 2))

        crit_headers = [
            Paragraph("ESCALA DE PROBABILIDAD (1 a 5)", style_crit_hdr),
            Paragraph("ESCALA DE IMPACTO (1 a 5)", style_crit_hdr),
            Paragraph("VALORACIÓN Y CATEGORIZACIÓN (P + I)", style_crit_hdr)
        ]
        crit_c1 = (
            "1 = Raro (Excepcional)<br/>"
            "2 = Improbable (Ocasional)<br/>"
            "3 = Posible (Cualquier momento)<br/>"
            "4 = Probable (Alta probabilidad)<br/>"
            "5 = Casi Cierto (Mayoría de circunstancias)"
        )
        crit_c2 = (
            "1 = Insignificante (< 1% del valor)<br/>"
            "2 = Menor (1% al 5% del valor)<br/>"
            "3 = Moderado (5% al 15% del valor)<br/>"
            "4 = Mayor (15% al 30% del valor)<br/>"
            "5 = Catastrófico (> 30% o incumplimiento total)"
        )
        crit_c3 = (
            "• BAJO: Suma de 2, 3 o 4<br/>"
            "• MEDIO: Suma de 5<br/>"
            "• ALTO: Suma de 6 o 7<br/>"
            "• EXTREMO: Suma de 8, 9 o 10<br/>"
            "Regla: Valoración = Probabilidad + Impacto"
        )
        crit_data = [
            crit_headers,
            [Paragraph(crit_c1, style_crit_body), Paragraph(crit_c2, style_crit_body), Paragraph(crit_c3, style_crit_body)]
        ]
        col_w_crit = [(page_width - 56) / 3] * 3
        t_crit = Table(crit_data, colWidths=col_w_crit)
        t_crit.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A365D")),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING', (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(t_crit)
        story.append(Spacer(1, 6))

        # 3. Tabla de 14 Columnas de la Matriz de Riesgos
        h_titles = [
            "N°", "Clase", "Fuente", "Etapa", "Tipo",
            "Descripción del Riesgo", "Consecuencia de la Ocurrencia",
            "EVALUACIÓN INICIAL", "Asignación", "Tratamiento / Medida de Control",
            "EVALUACIÓN RESIDUAL", "¿Afecta?", "Responsable", "Monitoreo y Revisión"
        ]
        tbl_headers = [Paragraph(f"<b>{h}</b>", style_tbl_hdr) for h in h_titles]
        col_w_tbl = [18, 30, 28, 34, 46, 114, 94, 52, 42, 106, 52, 24, 44, 52]

        tbl_rows = [tbl_headers]
        tbl_styles = [
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A365D")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 2.5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
            ('LEFTPADDING', (0,0), (-1,-1), 3),
            ('RIGHTPADDING', (0,0), (-1,-1), 3),
        ]

        # Detectar condiciones del proceso para blindaje en documento PDF
        obj_lower = (result.objeto or "").lower()
        is_mantenimiento = any(k in obj_lower for k in ["mantenimiento", "herramienta", "equipo", "electromec", "reparaci", "taller", "eléctric", "electric"])
        es_corto_plazo = is_short_term_period(plazo_str=result.plazo_texto)

        for idx, r in enumerate(result.riesgos, start=1):
            r_desc = r.descripcion
            r_cons = r.consecuencia
            r_trat = r.tratamiento
            r_mec = r.monitoreo_mecanismo
            r_per = r.monitoreo_periodicidad
            r_tipo = r.tipo

            if is_mantenimiento:
                if "corantioquia" in r_cons.lower() or "sspd" in r_cons.lower():
                    r_cons = re.sub(r"ante\s+autoridades\s+ambientales\s+competentes\s*\([^)]*\)", "en la operación de los servicios públicos de la ESPY", r_cons, flags=re.IGNORECASE)
                    r_cons = re.sub(r"\(?Corantioquia/SSPD\)?", "la ESPY", r_cons, flags=re.IGNORECASE)
                if "separación en la fuente" in r_desc.lower() or "renuencia comunitaria" in r_desc.lower() or "apatía" in r_desc.lower():
                    r_tipo = "Operacional / Custodia"
                    r_desc = "Daño, extravío o deterioro de herramientas o equipos propiedad de la ESPY durante su traslado, custodia o intervención en taller."
                    r_cons = "Afectación al patrimonio público de la ESPY y retraso en la disponibilidad de equipos para atender contingencias operativas."
                    r_trat = "Levantamiento de inventario y estado inicial al entregar y recibir los equipos mediante acta suscrita por el supervisor. El contratista asume a todo costo la reposición inmediata de cualquier elemento dañado o extraviado."
                    r_mec = "Inspección técnica detallada y verificación física contra acta de entrega."
                    r_per = "Al inicio, durante la ejecución y en la liquidación/pago único" if es_corto_plazo else "Al momento de cada entrega de equipos."

            if es_corto_plazo:
                r_desc = re.sub(
                    r"informes\s+mensuales\s+de\s+actividades\s+y\s+evidencias\s+de\s+ejecuci[oó]n[^.]*dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes",
                    "entrega final e informe técnico dentro del plazo pactado",
                    r_desc,
                    flags=re.IGNORECASE
                )
                r_desc = re.sub(r"informes\s+mensuales\s+dentro\s+de\s+los\s+primeros\s+cinco\s*\(\s*5\s*\)\s*d[ií]as\s+h[aá]biles\s+de\s+cada\s+mes", "entrega final e informe técnico dentro del plazo pactado", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\binformes\s+mensuales\b", "informe final de actividades", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de pago único", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\bpago\s+mensual\b", "pago único", r_desc, flags=re.IGNORECASE)
                r_desc = re.sub(r"\bvalor\s+mensualizado\b", "valor total", r_desc, flags=re.IGNORECASE)

                r_cons = re.sub(r"\btr[aá]mite\s+mensual\s+de\s+cuentas\b", "trámite de liquidación y pago", r_cons, flags=re.IGNORECASE)
                r_cons = re.sub(r"\bpago\s+mensual\b", "pago único", r_cons, flags=re.IGNORECASE)

                r_trat = re.sub(r"\ba\s+mes\s+vencido\b", "contra entrega final a satisfacción", r_trat, flags=re.IGNORECASE)
                r_trat = re.sub(r"\btr[aá]mite\s+mensual\b", "trámite de pago", r_trat, flags=re.IGNORECASE)

                if "mensual" in r_per.lower() or "quincenal" in r_per.lower() or "diario" in r_per.lower() or "permanente" in r_per.lower():
                    r_per = "Al inicio, durante la ejecución y en la liquidación/pago único"

            ini_text = f"P:{r.probabilidad_inicial}  I:{r.impacto_inicial}<br/><b>Val:{r.valoracion_inicial}</b><br/><b>{r.categoria_inicial.upper()}</b>"
            res_text = f"P:{r.probabilidad_residual}  I:{r.impacto_residual}<br/><b>Val:{r.valoracion_residual}</b><br/><b>{r.categoria_residual.upper()}</b>"
            mon_html = f"{r_mec}<br/>Periodicidad: {r_per}".replace('\n', '<br/>')

            row = [
                Paragraph(str(r.numero), style_tbl_cell_bold_center),
                Paragraph(r.clase, style_tbl_cell_center),
                Paragraph(r.fuente, style_tbl_cell_center),
                Paragraph(r.etapa, style_tbl_cell_center),
                Paragraph(f"<b>{r_tipo}</b>", style_tbl_cell_center),
                Paragraph(r_desc, style_tbl_cell),
                Paragraph(r_cons, style_tbl_cell),
                Paragraph(ini_text, style_tbl_cell_center),
                Paragraph(r.asignacion, style_tbl_cell_center),
                Paragraph(r_trat, style_tbl_cell),
                Paragraph(res_text, style_tbl_cell_center),
                Paragraph(f"<b>{r.afecta_ejecucion}</b>", style_tbl_cell_center),
                Paragraph(r.responsable, style_tbl_cell_center),
                Paragraph(mon_html, style_tbl_cell)
            ]
            tbl_rows.append(row)

            # Colores de evaluación inicial
            if r.categoria_inicial.upper() == "EXTREMO":
                tbl_styles.append(('BACKGROUND', (7, idx), (7, idx), colors.HexColor("#FCA5A5")))
            elif r.categoria_inicial.upper() == "ALTO":
                tbl_styles.append(('BACKGROUND', (7, idx), (7, idx), colors.HexColor("#FEE2E2")))
            elif r.categoria_inicial.upper() == "MEDIO":
                tbl_styles.append(('BACKGROUND', (7, idx), (7, idx), colors.HexColor("#FEF3C7")))
            else:
                tbl_styles.append(('BACKGROUND', (7, idx), (7, idx), colors.HexColor("#ECFDF5")))

            # Colores de evaluación residual
            if r.categoria_residual.upper() == "BAJO":
                tbl_styles.append(('BACKGROUND', (10, idx), (10, idx), colors.HexColor("#ECFDF5")))
            elif r.categoria_residual.upper() == "MEDIO":
                tbl_styles.append(('BACKGROUND', (10, idx), (10, idx), colors.HexColor("#FEF3C7")))
            else:
                tbl_styles.append(('BACKGROUND', (10, idx), (10, idx), colors.HexColor("#FEE2E2")))

        t_matrix = Table(tbl_rows, colWidths=col_w_tbl, repeatRows=1)
        t_matrix.setStyle(TableStyle(tbl_styles))
        story.append(t_matrix)
        story.append(Spacer(1, 8))

        # 4. Observaciones Técnicas
        if result.observaciones_estudio_previo:
            story.append(Paragraph("OBSERVACIONES TÉCNICAS Y CONTROL PREVENTIVO DE RIESGOS (ESTUDIO PREVIO):", obs_title_style))
            story.append(Spacer(1, 2))
            obs_bullets = "<br/>".join([f"• {obs}" for obs in result.observaciones_estudio_previo])
            t_obs = Table([[Paragraph(obs_bullets, style_crit_body)]], colWidths=[page_width - 56])
            t_obs.setStyle(TableStyle([
                ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('LEFTPADDING', (0,0), (-1,-1), 6),
                ('RIGHTPADDING', (0,0), (-1,-1), 6),
            ]))
            story.append(t_obs)
            story.append(Spacer(1, 12))

        # 5. Firma Oficial
        sig_data = [
            [Paragraph("________________________________________", style_tbl_cell_bold_center)],
            [Paragraph("<b>CAROLINA GIRALDO GÓMEZ</b><br/>Gerente General", style_tbl_cell_bold_center)]
        ]
        t_sig = Table(sig_data, colWidths=[240])
        t_sig.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('TOPPADDING', (0,0), (-1,-1), 2),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ]))
        wrapper_sig = Table([[t_sig]], colWidths=[page_width - 56])
        wrapper_sig.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        story.append(KeepTogether(wrapper_sig))

        doc.build(story, canvasmaker=NumberedCanvas)
        logger.info(f"Matriz de riesgos en formato PDF generada en: {output_path}")
        return output_path


risk_agent = RiskMatrixAgent()

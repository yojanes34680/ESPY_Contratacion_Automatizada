"""Excel Master Matrix synchronization engine using openpyxl. Preserves corporate styling, formats, and cell borders."""
import threading
from pathlib import Path
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.models.contract import Contrato
from app.utils.logger import logger

# Bloqueo global de hilo para evitar concurrencia conflictiva y corrupción en el archivo Excel
excel_lock = threading.Lock()

COLUMNS_CONFIG = [
    {"header": "Vigencia", "width": 12, "align": "center"},
    {"header": "No. Contrato", "width": 22, "align": "center"},
    {"header": "Modalidad Selección", "width": 24, "align": "left"},
    {"header": "Tipología Contractual", "width": 26, "align": "left"},
    {"header": "Identificación Contratista", "width": 22, "align": "center"},
    {"header": "Nombre / Razón Social Contratista", "width": 38, "align": "left"},
    {"header": "Objeto Contractual", "width": 55, "align": "left"},
    {"header": "Valor Total (COP)", "width": 22, "align": "right", "format": "$#,##0"},
    {"header": "Plazo Ejecución", "width": 16, "align": "center"},
    {"header": "Lugar Ejecución", "width": 24, "align": "left"},
    {"header": "CDP / Rubro", "width": 22, "align": "center"},
    {"header": "No. RP", "width": 14, "align": "center"},
    {"header": "Supervisor Designado", "width": 30, "align": "left"},
    {"header": "Cargo Supervisor", "width": 26, "align": "left"},
    {"header": "Documento Supervisor", "width": 22, "align": "center"},
    {"header": "Fecha Firma (Suscripción)", "width": 22, "align": "center"},
    {"header": "Fecha Inicio (Acta)", "width": 20, "align": "center"},
    {"header": "Fecha Terminación", "width": 20, "align": "center"},
    {"header": "Estado Contractual", "width": 18, "align": "center"},
]

def get_or_create_master_workbook() -> openpyxl.Workbook:
    """Abre el archivo Excel maestro o lo inicializa con diseño corporativo si no existe."""
    settings.MASTER_EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    
    wb = None
    if settings.MASTER_EXCEL_PATH.exists():
        try:
            wb = openpyxl.load_workbook(str(settings.MASTER_EXCEL_PATH))
        except Exception as e:
            logger.warning(f"No se pudo cargar Excel existente ({e}), creando nuevo...")

    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Matriz de Contratación"
    else:
        ws = wb.active

    ws.views.sheetView[0].showGridLines = True

    # 1. Título institucional
    ws.merge_cells("A1:S1")
    title_cell = ws["A1"]
    title_cell.value = "EMPRESA DE SERVICIOS PÚBLICOS (ESP) - MATRIZ MAESTRA DE CONTRATACIÓN PÚBLICA"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # 2. Subtítulo normativo
    ws.merge_cells("A2:S2")
    sub_cell = ws["A2"]
    sub_cell.value = "Régimen Legal: Ley 142 de 1994 y Principios de la Función Administrativa y Fiscal (Art. 209 y 267 C.P.)"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    sub_cell.fill = PatternFill(start_color="0A4D8C", end_color="0A4D8C", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 3. Fila de Encabezados (Fila 4)
    ws.row_dimensions[3].height = 10  # Espaciador
    ws.row_dimensions[4].height = 28
    
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    for col_idx, col_cfg in enumerate(COLUMNS_CONFIG, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.value = col_cfg["header"]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
        # Ajustar ancho de columna
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_cfg["width"]

    wb.save(str(settings.MASTER_EXCEL_PATH))
    logger.info(f"Matriz de Contratación Maestra inicializada en: {settings.MASTER_EXCEL_PATH}")
    return wb

def sync_contract_to_master_excel(contrato: Contrato) -> Dict[str, Any]:
    """
    Sincroniza un contrato con el libro maestro Excel.
    Si el contrato ya existe (por número de contrato), actualiza sus datos en la misma fila;
    de lo contrario, anexa una nueva fila con el formato y bordes correspondientes.
    Protegido concurrentemente con excel_lock.
    """
    with excel_lock:
        wb = get_or_create_master_workbook()
        ws = wb.active

        # Extraer datos formateados
        c = contrato.contratista
        doc_id_str = f"{c.tipo_documento.value} {c.numero_documento}"
        if c.digito_verificacion:
            doc_id_str += f"-{c.digito_verificacion}"

        supervisor_doc = f"{contrato.supervisor_tipo_documento or 'CC'} {contrato.supervisor_documento}" if contrato.supervisor_documento else "Por definir"

        row_values = [
            contrato.vigencia,
            contrato.numero_contrato,
            contrato.modalidad_seleccion.value.replace("_", " ").title(),
            contrato.tipologia.value.replace("_", " ").title(),
            doc_id_str,
            c.nombre_o_razon_social,
            contrato.objeto,
            contrato.valor_total or 0.0,
            f"{contrato.plazo_ejecucion} {contrato.unidad_plazo.value.title()}",
            contrato.lugar_ejecucion or "Jurisdicción ESP",
            f"{contrato.cdp_numero or 'S/N'} ({contrato.cdp_rubro or 'General'})",
            contrato.rp_numero or "Pendiente",
            contrato.supervisor_nombre or "Por definir",
            contrato.supervisor_cargo or "Por definir",
            supervisor_doc,
            str(contrato.fecha_suscripcion or "Pendiente"),
            str(contrato.fecha_acta_inicio or "Pendiente"),
            str(contrato.fecha_terminacion or "Pendiente"),
            contrato.estado.value
        ]

        # Buscar si ya existe el número de contrato en la columna B (fila 5 en adelante)
        target_row = None
        max_r = ws.max_row
        for r in range(5, max_r + 1):
            cell_val = ws.cell(row=r, column=2).value
            if cell_val and str(cell_val).strip() == str(contrato.numero_contrato).strip():
                target_row = r
                break

        action = "ACTUALIZADO" if target_row else "CREADO"
        if not target_row:
            # Fila nueva: si max_row es menor que 4, empezar en 5
            target_row = max(5, max_r + 1)

        # Aplicar estilos de fila
        thin_border = Border(
            left=Side(style="thin", color="DDDDDD"),
            right=Side(style="thin", color="DDDDDD"),
            top=Side(style="thin", color="DDDDDD"),
            bottom=Side(style="thin", color="DDDDDD")
        )
        # Alternar color de fila (cebreado sutil)
        bg_color = "F9FBFC" if target_row % 2 == 0 else "FFFFFF"
        data_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        data_font = Font(name="Calibri", size=9.5, color="222222")

        for col_idx, (val, col_cfg) in enumerate(zip(row_values, COLUMNS_CONFIG), start=1):
            cell = ws.cell(row=target_row, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.fill = data_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal=col_cfg["align"],
                vertical="center",
                wrap_text=(col_cfg["header"] == "Objeto Contractual")
            )
            if "format" in col_cfg:
                cell.number_format = col_cfg["format"]

        ws.row_dimensions[target_row].height = 24
        wb.save(str(settings.MASTER_EXCEL_PATH))
        logger.info(f"Contrato {contrato.numero_contrato} {action} en matriz Excel (Fila {target_row}).")

        return {
            "numero_contrato": contrato.numero_contrato,
            "fila_excel": target_row,
            "accion": action,
            "archivo_maestro": str(settings.MASTER_EXCEL_PATH)
        }

def sync_all_contracts_to_master_excel(db_session) -> Dict[str, Any]:
    """Sincroniza todos los contratos registrados en la base de datos hacia el Excel maestro."""
    contratos = db_session.query(Contrato).order_by(Contrato.id.asc()).all()
    results = []
    for c in contratos:
        res = sync_contract_to_master_excel(c)
        results.append(res)
    return {
        "total_sincronizados": len(results),
        "detalles": results,
        "archivo_maestro": str(settings.MASTER_EXCEL_PATH)
    }

def delete_contract_from_master_excel(contract_number: str) -> Dict[str, Any]:
    """
    Elimina la fila correspondiente a un contrato en la Matriz Maestra de Excel (.xlsx).
    Protegido concurrentemente con excel_lock.
    """
    with excel_lock:
        wb = get_or_create_master_workbook()
        ws = wb.active

        target_row = None
        for r in range(5, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=2).value
            if cell_val and str(cell_val).strip() == str(contract_number).strip():
                target_row = r
                break

        if not target_row:
            return {
                "eliminado": False,
                "numero_contrato": contract_number,
                "mensaje": f"El contrato '{contract_number}' no se encontró en la matriz Excel."
            }

        ws.delete_rows(target_row, 1)
        wb.save(str(settings.MASTER_EXCEL_PATH))
        logger.info(f"Contrato {contract_number} eliminado de la matriz Excel (Fila anterior {target_row}).")
        return {
            "eliminado": True,
            "numero_contrato": contract_number,
            "fila_eliminada": target_row,
            "archivo_maestro": str(settings.MASTER_EXCEL_PATH)
        }

def generate_filtered_excel(vigencia: int, db_session) -> Path:
    """
    Genera un archivo Excel con el diseño institucional filtrado para una vigencia fiscal específica.
    """
    export_dir = settings.BASE_DIR / "storage" / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    out_path = export_dir / f"Matriz_Contratacion_Vigencia_{vigencia}_ESP.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Vigencia {vigencia}"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título institucional
    ws.merge_cells("A1:S1")
    title_cell = ws["A1"]
    title_cell.value = f"EMPRESA DE SERVICIOS PÚBLICOS (ESP) - MATRIZ DE CONTRATACIÓN - VIGENCIA FISCAL {vigencia}"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. Subtítulo
    ws.merge_cells("A2:S2")
    sub_cell = ws["A2"]
    sub_cell.value = f"CONSOLIDADO DE CONTRATOS ESTATALES Y PRIVADOS LEY 142/1994 - EJERCICIO FISCAL {vigencia}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="DDDDDD")
    sub_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 10

    # 3. Encabezados
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    for col_idx, col_cfg in enumerate(COLUMNS_CONFIG, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.value = col_cfg["header"]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_cfg["width"]

    ws.row_dimensions[4].height = 28

    # 4. Datos
    contratos = db_session.query(Contrato).filter(Contrato.vigencia == vigencia).order_by(Contrato.numero_contrato.asc()).all()

    data_font = Font(name="Calibri", size=9)
    border_data = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD")
    )

    for idx, contrato in enumerate(contratos, start=5):
        c = contrato.contratista
        doc_id_str = f"{c.tipo_documento.value} {c.numero_documento}" if c else "-"
        if c and c.digito_verificacion:
            doc_id_str += f"-{c.digito_verificacion}"
        c_nombre = c.nombre_o_razon_social if c else "Sin Contratista"

        supervisor_doc = f"{contrato.supervisor_tipo_documento or 'CC'} {contrato.supervisor_documento}" if contrato.supervisor_documento else "Por definir"

        row_values = [
            contrato.vigencia,
            contrato.numero_contrato,
            contrato.modalidad_seleccion.value.replace("_", " ").title(),
            contrato.tipologia.value.replace("_", " ").title(),
            doc_id_str,
            c_nombre,
            contrato.objeto,
            contrato.valor_total or 0.0,
            f"{contrato.plazo_ejecucion} {contrato.unidad_plazo.value.title()}",
            contrato.lugar_ejecucion or "Jurisdicción ESP",
            f"{contrato.cdp_numero or 'S/N'} ({contrato.cdp_rubro or 'General'})",
            contrato.rp_numero or "Pendiente",
            contrato.supervisor_nombre or "Por definir",
            contrato.supervisor_cargo or "Por definir",
            supervisor_doc,
            str(contrato.fecha_suscripcion or "Pendiente"),
            str(contrato.fecha_acta_inicio or "Pendiente"),
            str(contrato.fecha_terminacion or "Pendiente"),
            contrato.estado.value
        ]

        fill_color = "F9FAFB" if idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx, (val, col_cfg) in enumerate(zip(row_values, COLUMNS_CONFIG), start=1):
            cell = ws.cell(row=idx, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.fill = row_fill
            cell.border = border_data
            cell.alignment = Alignment(
                horizontal=col_cfg["align"],
                vertical="center",
                wrap_text=(col_cfg["header"] == "Objeto Contractual")
            )
            if "format" in col_cfg:
                cell.number_format = col_cfg["format"]

        ws.row_dimensions[idx].height = 24

    wb.save(str(out_path))
    logger.info(f"Matriz Excel para vigencia {vigencia} generada con éxito ({len(contratos)} contratos): {out_path}")
    return out_path


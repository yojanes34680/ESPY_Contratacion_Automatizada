import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import pytest
import io
import docx
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def create_sample_estudio_docx(objeto_text: str, justificacion_text: str) -> io.BytesIO:
    doc = docx.Document()
    doc.add_heading("ESTUDIO PREVIO DE CONTRATACIÓN", level=1)
    doc.add_paragraph(f"OBJETO CONTRACTUAL: {objeto_text}")
    doc.add_paragraph("JUSTIFICACIÓN DE LA NECESIDAD:")
    doc.add_paragraph(justificacion_text)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def test_analyze_files_with_only_estudio_file():
    """Valida que el endpoint /unspsc/analyze-files funciona enviando UNICAMENTE el archivo de estudio previo."""
    objeto_str = "SUMINISTRO DE TAPAS Y CAJAS DE MEDIDORES PARA EL SISTEMA DE ACUEDUCTO"
    justif_str = (
        "La ESPY requiere el suministro de tapas y cajillas de medidores para la reposición "
        "y normalización de acometidas domiciliarias de acueducto en el municipio de Yolombó."
    )
    docx_file = create_sample_estudio_docx(objeto_str, justif_str)
    
    response = client.post(
        "/api/unspsc/analyze-files",
        data={
            "objeto": "",
            "estudio_previo_text": "",
            "propuesta_text": "",
        },
        files={
            "estudio_file": ("estudio_previo.docx", docx_file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
    )
    
    assert response.status_code == 200, f"Error: {response.text}"
    data = response.json()
    assert "codigos_seleccionados" in data
    assert "objeto_detectado" in data
    assert data["objeto_detectado"] is not None and len(data["objeto_detectado"]) > 0
    assert "tapas" in data["objeto_detectado"].lower() or "medidores" in data["objeto_detectado"].lower()
    
    # Cada código debe indicar fuente_cita = 'estudio_previo'
    for c in data["codigos_seleccionados"]:
        assert c.get("fuente_cita") == "estudio_previo"

def test_analyze_files_with_estudio_and_propuesta():
    """Valida que cuando se envían estudio previo y propuesta, los ítems se extraen de la propuesta y fuente_cita indica 'propuesta'."""
    objeto_str = "SUMINISTRO DE ELEMENTOS Y HERRAMIENTAS OPERATIVAS"
    justif_str = "La ESPY requiere la adquisición de herramientas para el área operativa."
    estudio_docx = create_sample_estudio_docx(objeto_str, justif_str)

    propuesta_text = """
    ÍTEM 1: Válvulas de compuerta de 2 pulgadas para acueducto
    ÍTEM 2: Tubería a presión RDE 21 de 3 pulgadas
    ÍTEM 3: Cintas de señalización y conos viales de seguridad
    """

    response = client.post(
        "/api/unspsc/analyze-files",
        data={
            "objeto": "",
            "estudio_previo_text": "",
            "propuesta_text": propuesta_text,
        },
        files={
            "estudio_file": ("estudio_previo.docx", estudio_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
    )

    assert response.status_code == 200, f"Error: {response.text}"
    data = response.json()
    assert "codigos_seleccionados" in data
    assert len(data["codigos_seleccionados"]) > 0
    assert data["objeto_detectado"] is not None and len(data["objeto_detectado"]) > 0

    # Cada código debe indicar fuente_cita = 'propuesta'
    for c in data["codigos_seleccionados"]:
        assert c.get("fuente_cita") == "propuesta"

def test_analyze_files_without_any_input_raises_400():
    """Valida que si no se envía ni archivo ni texto del estudio previo, arroja error 400 explicativo."""
    response = client.post(
        "/api/unspsc/analyze-files",
        data={
            "objeto": "",
            "estudio_previo_text": "",
            "propuesta_text": "",
        }
    )
    assert response.status_code == 400
    assert "Debe cargar el estudio previo" in response.json().get("detail", "")

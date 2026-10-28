"""Script de utilidad para limpiar textos de CDP existentes en la base de datos y documentos Word generados.

MODO DE USO:
1. Simulación / Auditoría de lectura (DRY RUN por defecto):
     python tools/limpiar_cdp_existentes.py
   Muestra el reporte de contratos con texto de CDP largo en la base de datos y el reporte
   de documentos Word generados que aún contienen "CCPET" o "clasificadores presupuestales".

2. Aplicar saneamiento a la base de datos:
     python tools/limpiar_cdp_existentes.py --aplicar
   Crea una copia de seguridad segura usando sqlite3 backup y actualiza metadatos_agente.

3. Regenerar documentos Word obsoletos:
     python tools/limpiar_cdp_existentes.py --regenerar
   Guarda primero una copia de respaldo en backups/documentos_<fecha>/ y luego regenera
   con docx_engine los documentos (Estudio Previo, Minuta, Acta de Inicio) afectados.

4. Aplicar base de datos y regenerar documentos en una sola pasada:
     python tools/limpiar_cdp_existentes.py --aplicar --regenerar
"""
import sys
import os
import re
import json
import shutil
import sqlite3
from pathlib import Path
from datetime import datetime

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
BACKEND_DIR = ROOT_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import docx
from app.utils.cdp_cleaner import limpiar_texto_cdp

DB_PATH = BACKEND_DIR / "espy_contratacion.db"
GENERATED_DIR = BACKEND_DIR / "storage" / "generated"

def safe_sqlite_backup(source_db_path: Path, dest_backup_path: Path):
    """Realiza un respaldo consistente de SQLite utilizando la API oficial de backup (seguro para modo WAL)."""
    dest_backup_path.parent.mkdir(parents=True, exist_ok=True)
    src_conn = sqlite3.connect(str(source_db_path))
    dst_conn = sqlite3.connect(str(dest_backup_path))
    try:
        with dst_conn:
            src_conn.backup(dst_conn)
    finally:
        dst_conn.close()
        src_conn.close()

def scan_generated_documents():
    """Escanea la carpeta backend/storage/generated en busca de documentos Word que contengan CCPET o clasificadores."""
    docs_afectados = []
    if not GENERATED_DIR.exists():
        return docs_afectados

    prohibited_patterns = [
        re.compile(r"\bCCPET\b", re.IGNORECASE),
        re.compile(r"clasificadores\s+presupuestales", re.IGNORECASE)
    ]

    for docx_path in GENERATED_DIR.glob("*.docx"):
        if docx_path.name.startswith("~$"):
            continue
        try:
            doc = docx.Document(str(docx_path))
            full_text = " ".join(p.text for p in doc.paragraphs)
            for t in doc.tables:
                for r in t.rows:
                    for c in r.cells:
                        full_text += " " + c.text

            found = []
            for pat in prohibited_patterns:
                if pat.search(full_text):
                    found.append(pat.pattern)

            if found:
                # Deducir contrato de nombre de archivo (ej: Estudio_Previo_Servicios_CPS-050-2026_2026.docx)
                m_num = re.search(r"((?:CPS|CS|CO|CA|CC|CI)-[0-9]{3,4}-[0-9]{4})", docx_path.name, re.IGNORECASE)
                num_contrato = m_num.group(1).upper() if m_num else "DESCONOCIDO"

                tipo_doc = "OTRO"
                name_upper = docx_path.name.upper()
                if "ESTUDIO_PREVIO" in name_upper:
                    tipo_doc = "ESTUDIO_PREVIO"
                elif "MINUTA" in name_upper:
                    tipo_doc = "MINUTA"
                elif "ACTA_INICIO" in name_upper:
                    tipo_doc = "ACTA_INICIO"

                docs_afectados.append({
                    "path": docx_path,
                    "nombre": docx_path.name,
                    "numero_contrato": num_contrato,
                    "tipo_documento": tipo_doc,
                    "coincidencias": found
                })
        except Exception:
            continue

    return docs_afectados

def main():
    aplicar = "--aplicar" in sys.argv
    regenerar = "--regenerar" in sys.argv

    print("=" * 80)
    print("HERRAMIENTA INTEGRAL DE SANEAMIENTO Y LIMPIEZA DE CDP (BD Y DOCUMENTOS WORD)")
    print("=" * 80)
    print(f"Base de datos objetivo: {DB_PATH}")
    print(f"Modo BD: {'APLICAR CAMBIOS (--aplicar)' if aplicar else 'SÓLO LECTURA / SIMULACIÓN (DRY RUN)'}")
    print(f"Modo Documentos: {'REGENERAR WORD OBSOLETOS (--regenerar)' if regenerar else 'REPORTE DE AUDITORÍA'}")
    print("-" * 80)

    if not DB_PATH.exists():
        print(f"Error: No se encontró la base de datos en {DB_PATH}")
        sys.exit(1)

    # 1. Respaldo de seguridad previo de base de datos con SQLite Backup API
    if aplicar:
        backup_dir = ROOT_DIR / "backups"
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_file = backup_dir / f"espy_contratacion_backup_before_cdp_clean_{ts}.db"
        safe_sqlite_backup(DB_PATH, backup_file)
        print(f"[OK] Respaldo consistente de base de datos (sqlite3 backup) creado en:\n     {backup_file}")
        print("-" * 80)

    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()

    cur.execute("SELECT id, numero_contrato, metadatos_agente FROM contratos")
    rows = cur.fetchall()

    contratos_a_modificar = []

    for cid, numero, meta_str in rows:
        if not meta_str:
            continue
        try:
            meta = json.loads(meta_str)
        except Exception:
            continue

        cdp_txt = meta.get("cdp_texto_completo")
        if not cdp_txt:
            continue

        cdp_limpio = limpiar_texto_cdp(cdp_txt)
        if cdp_limpio != cdp_txt:
            meta["cdp_texto_completo"] = cdp_limpio
            contratos_a_modificar.append({
                "id": cid,
                "numero": numero,
                "antes": cdp_txt,
                "despues": cdp_limpio,
                "nuevo_meta": json.dumps(meta, ensure_ascii=False)
            })

    total_bd = len(contratos_a_modificar)
    print(f"1. AUDITORÍA DE BASE DE DATOS:")
    print(f"   Total contratos con texto de CDP largo o residual en BD: {total_bd}")

    if total_bd > 0:
        print("\n   Ejemplos de saneamiento en BD (hasta 3):")
        for item in contratos_a_modificar[:3]:
            print(f"   - Contrato {item['numero']} (ID: {item['id']}):")
            print(f"     [ANTES]   : {item['antes'][:100]}...")
            print(f"     [DESPUÉS] : {item['despues']}\n")

        if aplicar:
            for item in contratos_a_modificar:
                cur.execute(
                    "UPDATE contratos SET metadatos_agente = ? WHERE id = ?",
                    (item["nuevo_meta"], item["id"])
                )
            con.commit()
            print(f"   [OK] ÉXITO: Se actualizaron satisfactoriamente {total_bd} contratos en la base de datos.")
    else:
        print("   [OK] Todos los contratos en la base de datos ya se encuentran limpios.")

    con.close()
    print("-" * 80)

    # 2. Auditoría y Regeneración de Documentos Word Generados
    print("2. AUDITORÍA DE DOCUMENTOS WORD GENERADOS (storage/generated/):")
    docs_afectados = scan_generated_documents()
    print(f"   Total documentos Word encontrados con 'CCPET' o clasificadores: {len(docs_afectados)}")

    if docs_afectados:
        print("\n   Listado de documentos detectados con texto largo:")
        for d in docs_afectados:
            print(f"   - [{d['numero_contrato']}] ({d['tipo_documento']}): {d['nombre']}")

        if regenerar:
            print("\n   INICIANDO REGENERACIÓN SEGURA DE DOCUMENTOS WORD:")
            # Crear directorio de respaldo para archivos viejos
            ts_doc = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_docs_dir = ROOT_DIR / "backups" / f"documentos_{ts_doc}"
            backup_docs_dir.mkdir(parents=True, exist_ok=True)
            print(f"   [OK] Respaldando documentos viejos en: {backup_docs_dir}")

            # Inicializar sesión ORM para usar docx_engine
            from app.database import SessionLocal
            from app.models.contract import Contrato
            from app.services.docx_engine import generate_docx_for_contract

            db = SessionLocal()
            regenerados_ok = 0
            try:
                for d in docs_afectados:
                    # Copiar documento viejo a backup
                    shutil.copy2(d["path"], backup_docs_dir / d["nombre"])

                    num = d["numero_contrato"]
                    tipo = d["tipo_documento"]
                    if num == "DESCONOCIDO" or tipo == "OTRO":
                        print(f"   [SKIP] No se pudo mapear contrato/tipo para: {d['nombre']}")
                        continue

                    contrato_db = db.query(Contrato).filter(Contrato.numero_contrato == num).first()
                    if not contrato_db:
                        print(f"   [WARN] Contrato {num} no existe en BD para regenerar {d['nombre']}")
                        continue

                    # Regenerar documento usando el motor oficial
                    new_path = generate_docx_for_contract(
                        contrato=contrato_db,
                        template_type=tipo,
                        db_session=db
                    )
                    print(f"   [REGENERADO OK] {d['nombre']} -> {new_path.name}")
                    regenerados_ok += 1
            finally:
                db.close()

            print(f"\n   [OK] Total documentos regenerados satisfactoriamente: {regenerados_ok} / {len(docs_afectados)}")
        else:
            print("\n   NOTA: Los documentos Word no se modificaron (Modo reporte de auditoría).")
            print("   Para regenerarlos automáticamente eliminando CCPET, ejecute:")
            print("   python tools/limpiar_cdp_existentes.py --regenerar")
    else:
        print("   [OK] Todos los documentos Word generados ya están limpios y libres de CCPET.")

    print("=" * 80)

if __name__ == "__main__":
    main()

import docx
import re
from pathlib import Path

def patch_custom_plantilla_minuta_suministro():
    p = Path("backend/templates/custom/custom_plantilla_minuta_suministro.docx")
    if not p.exists():
        print(f"File {p} not found")
        return

    doc = docx.Document(str(p))

    # 1. P3: Preámbulo y considerandos
    p3 = doc.paragraphs[3]
    t3 = p3.text
    # Género de la Gerente
    t3 = re.sub(r"\bvecino\s+del\s+Municipio\s+de\s+Yolomb[oó\ufffd]", "vecina del Municipio de Yolombó", t3, flags=re.IGNORECASE)
    t3 = re.sub(r"\bidentificado\s+con\s+la\s+c[eé\ufffd]dula", "identificada con la cédula", t3, flags=re.IGNORECASE)
    t3 = re.sub(r"\bdebidamente\s+facultado\b", "debidamente facultada", t3, flags=re.IGNORECASE)
    # Consideración 2)
    t3 = re.sub(
        r'2\)\s*Que\s+el\s+objeto\s+principal\s+de\s+la\s+Empresa\s+de\s+Servicios\s+P[uú\ufffd]blicos\s+(?:de\s+Yolomb[oó\ufffd]\s+)?E\.?S\.?P\.?\s*(?:es\s+la|[Ll]a)\s+eficiente\s+prestaci[oó\ufffd]n\s+de\s+los\s+servicios\s+p[uú\ufffd]blicos\s+domiciliarios\s+de\s+Acueducto,\s*Alcantarillado\s+y\s+Aseo\s+en\s+el\s+Municipio\s+de\s+Yolomb[oó\ufffd]\s+y\s+para\s+ello\s+debe\s+garantizar\s+las\s+labores\s+tendientes\s+(?:para|a)\s+una\s+eficiente\s+prestaci[oó\ufffd]n\s+de\s+los\s+mismos(?:\s*\"|\s*,|\s*\")*,?',
        '2) Que el objeto principal de la Empresa de Servicios Públicos de Yolombó E.S.P. es la eficiente prestación de los servicios públicos domiciliarios de Acueducto, Alcantarillado y Aseo en el Municipio de Yolombó y para ello debe garantizar las labores tendientes a una eficiente prestación de los mismos, ',
        t3
    )
    p3.text = t3

    # 2. P18: Cláusula Octava (Cesión del Contrato)
    p18 = doc.paragraphs[18]
    p18.text = re.sub(
        r"\s*Que\s+corresponde\s+a\s+la\s+Gerente\s+de\s+LA\s+EMPRESA\s+como\s+ordenadora\s+del\s+gasto[^\n.]*\.",
        "",
        p18.text
    ).strip()

    # 3. P33: Cláusula Décima Sexta (Terminación del Contrato)
    p33 = doc.paragraphs[33]
    p33.text = (
        "DECIMA SEXTA. TERMINACIÓN DEL CONTRATO: El presente contrato podrá terminar por alguno de los siguientes eventos: "
        "a) vencimiento del plazo sin que las partes, por mutuo acuerdo y por escrito, manifiesten su intención de prorrogarlo; "
        "b) por decisión mutua de las partes, lo cual deberá constar en documento escrito por ambas; "
        "c) por incumplimiento de cualquiera de las obligaciones de EL CONTRATISTA, incluida aquella de afiliarse al Sistema General de Seguridad Social o por mora en el pago de las cotizaciones (salud, pensiones, riesgos laborales) por más de un mes o periodo; "
        "d) De acreditarse que el Contratista esté incurso en inhabilidades o incompatibilidades de cualquier naturaleza; "
        "e) Las demás establecidas en la ley."
    )

    # 4. P43: Cláusula Vigésima Primera (Eliminar duplicidad de domicilio)
    p43 = doc.paragraphs[43]
    p43.text = (
        "VIGÉSIMA PRIMERA. JURISDICCIÓN Y RÉGIMEN APLICABLE: Para todos los efectos legales y procesales, "
        "las controversias que no puedan ser resueltas de común acuerdo se someterán a la jurisdicción ordinaria "
        "competente en el Municipio de Yolombó - Antioquia."
    )

    doc.save(str(p))
    print(f"Patched {p} successfully.")

if __name__ == "__main__":
    patch_custom_plantilla_minuta_suministro()

"""
Test script for verifying:
1. Custom Word templates upload, placeholder inspection, generation, and reset.
2. Multi-document upload (up to 8 documents).
3. Multi-document extraction and consolidation.
4. Contractor creation linking all uploaded documents.
"""
import sys
import io
import requests
import docx
from pathlib import Path

BASE_URL = "http://127.0.0.1:8000"

def test_custom_templates():
    print("\n--- TEST 1: Template Status ---")
    r = requests.get(f"{BASE_URL}/api/templates/status")
    assert r.status_code == 200, f"Status failed: {r.text}"
    templates = r.json()
    assert len(templates) == 3, f"Expected 3 templates, got {len(templates)}"
    print(f"Templates retrieved: {[t['tipo'] for t in templates]}")

    print("\n--- TEST 2: Upload Custom Template ---")
    # Create a custom test docx
    custom_doc = docx.Document()
    custom_doc.add_heading("MINUTA PERSONALIZADA DE PRUEBA ESP 2026", level=0)
    custom_doc.add_paragraph("Entre los suscritos, la ESP y {{CONTRATISTA}}, identificado con {{DOCUMENTO}}.")
    custom_doc.add_paragraph("Objeto: {{OBJETO}}")
    custom_doc.add_paragraph("Valor: {{VALOR}} ({{VALOR_LETRAS}})")
    custom_doc.add_paragraph("Plazo: {{PLAZO}}")
    custom_doc.add_paragraph("Obligaciones:\n{{OBLIGACIONES}}")
    custom_doc.add_paragraph("Supervisor: {{SUPERVISOR}}")

    buf = io.BytesIO()
    custom_doc.save(buf)
    buf.seek(0)

    files = {"file": ("mi_minuta_personalizada.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r_upload = requests.post(f"{BASE_URL}/api/templates/upload/MINUTA", files=files)
    assert r_upload.status_code == 200, f"Upload custom template failed: {r_upload.text}"
    detail = r_upload.json()["detalle"]
    assert detail["is_custom"] is True, "Expected is_custom=True"
    print(f"Custom template uploaded: {detail['archivo_nombre']}, is_custom={detail['is_custom']}")
    print(f"Detected placeholders: {detail['placeholders_detectados']}")
    assert "CONTRATISTA" in detail["placeholders_detectados"]
    assert "OBJETO" in detail["placeholders_detectados"]

    print("\n--- TEST 3: Generate Document using Custom Template ---")
    # Call generate on contract 1
    r_gen = requests.post(f"{BASE_URL}/api/templates/generate/1/MINUTA")
    assert r_gen.status_code == 200, f"Generation failed: {r_gen.text}"
    print(f"Document generated with custom template: {r_gen.json()['archivo_nombre']}")

    print("\n--- TEST 4: Reset Custom Template ---")
    r_reset = requests.post(f"{BASE_URL}/api/templates/reset/MINUTA")
    assert r_reset.status_code == 200, f"Reset failed: {r_reset.text}"
    reset_detail = r_reset.json()["detalle"]
    assert reset_detail["is_custom"] is False, "Expected is_custom=False after reset"
    print(f"Template reset: is_custom={reset_detail['is_custom']}, file={reset_detail['archivo_nombre']}")

def test_multidocument_upload_and_extraction():
    print("\n--- TEST 5: Upload Multiple Documents (Batch of 3) ---")
    # File 1: RUT snippet docx
    d1 = docx.Document()
    d1.add_paragraph("DIAN - REGISTRO UNICO TRIBUTARIO\nNIT: 901555444-1\nRAZON SOCIAL: INGENIERIA HIDRAULICA ANDINA S.A.S.\nDIRECCION: CALLE 50 # 40-30\nCIUDAD: MEDELLIN\nTELEFONO: 6042345678\nCORREO: contacto@hidraulicaandina.com")
    buf1 = io.BytesIO()
    d1.save(buf1)
    buf1.seek(0)

    # File 2: Propuesta Económica docx
    d2 = docx.Document()
    d2.add_paragraph("PROPUESTA ECONOMICA Y TECNICA\nVALOR PROPUESTA: $45.000.000 COP\nPLAZO PROPUESTO: 8 MESES\nPERFIL: Ingenieros especialistas en acueductos y tratamiento de agua potable.")
    buf2 = io.BytesIO()
    d2.save(buf2)
    buf2.seek(0)

    # File 3: Certificado docx
    d3 = docx.Document()
    d3.add_paragraph("CERTIFICACION BANCARIA\nBANCO: BANCOLOMBIA\nCUENTA CORRIENTE: 123-456789-01")
    buf3 = io.BytesIO()
    d3.save(buf3)
    buf3.seek(0)

    files = [
        ("files", ("RUT_Hidraulica_Andina.docx", buf1.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("Propuesta_Economica_2026.docx", buf2.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("Certificado_Bancario.docx", buf3.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]

    r_multi = requests.post(f"{BASE_URL}/api/documents/upload-multiple", files=files)
    assert r_multi.status_code == 201, f"Multi upload failed: {r_multi.text}"
    uploaded = r_multi.json()
    assert len(uploaded) == 3, f"Expected 3 uploaded docs, got {len(uploaded)}"
    doc_ids = [doc["id"] for doc in uploaded]
    print(f"3 documents uploaded in batch: IDs {doc_ids}")

    print("\n--- TEST 6: Multi-Document Extraction & Consolidation ---")
    r_ext = requests.post(f"{BASE_URL}/api/agents/extract-multiple-documents", json={"document_ids": doc_ids})
    assert r_ext.status_code == 200, f"Extraction failed: {r_ext.text}"
    ext_data = r_ext.json()
    cons = ext_data["consolidated"]
    print(f"Consolidated Name: {cons.get('nombre_o_razon_social')}")
    print(f"Consolidated Document: {cons.get('numero_documento')}")
    print(f"Consolidated Value: {cons.get('valor_propuesto')}")
    print(f"Consolidated Term: {cons.get('plazo_propuesto')} {cons.get('unidad_plazo_propuesto')}")
    print(f"Consolidated Email: {cons.get('correo_electronico')}")
    print(f"Consolidated Phone: {cons.get('telefono')}")

    assert cons.get("numero_documento") == "901555444", f"Expected 901555444, got {cons.get('numero_documento')}"
    assert cons.get("valor_propuesto") == 45000000.0, f"Expected 45000000, got {cons.get('valor_propuesto')}"
    assert cons.get("plazo_propuesto") == 8, f"Expected 8, got {cons.get('plazo_propuesto')}"

    print("\n--- TEST 7: Create Contractor from Multiple Documents ---")
    r_create = requests.post(f"{BASE_URL}/api/agents/create-contractor-from-docs", json={
        "document_ids": doc_ids,
        "contratista_data": cons
    })
    assert r_create.status_code == 200, f"Create contractor failed: {r_create.text}"
    contractor = r_create.json()
    print(f"Contractor created: ID {contractor['id']} - {contractor['nombre_o_razon_social']}")

    # Verify all 3 docs are linked to contractor
    r_docs = requests.get(f"{BASE_URL}/api/documents?contratista_id={contractor['id']}")
    assert r_docs.status_code == 200
    linked_docs = r_docs.json()
    assert len(linked_docs) >= 3, f"Expected at least 3 linked docs, got {len(linked_docs)}"
    print(f"Verified {len(linked_docs)} documents linked to contractor {contractor['id']}!")

if __name__ == "__main__":
    test_custom_templates()
    test_multidocument_upload_and_extraction()
    print("\n[SUCCESS] ALL NEW CAPABILITIES TESTED AND VERIFIED 100%!")

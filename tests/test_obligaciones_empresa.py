"""
Test Integral para validación de Obligaciones de la Empresa (ESPY)
y Forma de Pago acorde a las nuevas reglas:
1. Generación de 2 a 3 obligaciones técnicas de la ESPY según objeto/tipología.
2. Prohibición de 'honorarios' en mantenimiento, operativos, suministro, obra.
3. Prohibición de 'plataformas o sistemas' en mantenimiento/campo.
4. Regla de periodicidad para plazo <= 30 días (sin 'informes mensuales' ni 'pagos mensuales vencidos').
5. Validación de seguridad social según persona jurídica vs natural.
6. Generación de documentos Word (.docx) Estudio Previo y Minuta.
"""
import os
import sys

# Agregar backend al path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), 'backend')))

from app.utils.obligaciones_empresa_builder import (
    build_obligaciones_empresa,
    build_forma_pago_clausula,
    is_short_term_period
)
from app.services.docx_engine import generate_docx_for_contract

def test_short_term_logic():
    print("--- 1. Testing is_short_term_period ---")
    assert is_short_term_period("15 días", 15, "DIAS") is True
    assert is_short_term_period("30 días", 30, "DIAS") is True
    assert is_short_term_period("1 mes", 1, "MESES") is True
    assert is_short_term_period("45 días", 45, "DIAS") is False
    assert is_short_term_period("2 meses", 2, "MESES") is False
    assert is_short_term_period("6 meses", 6, "MESES") is False
    print("PASS: short term period logic works as expected.")

def test_maintenance_short_term():
    print("\n--- 2. Testing Mantenimiento 15 días (Persona Natural) ---")
    objeto = "MANTENIMIENTO PREVENTIVO Y CORRECTIVO A TODO COSTO DE HERRAMIENTAS Y EQUIPOS MENORES ADSCRITOS AL ÁREA OPERATIVA DE LA EMPRESA DE SERVICIOS PÚBLICOS DE YOLOMBÓ E.S.P."
    tipologia = "PRESTACION_SERVICIOS"
    plazo_str = "15 días"
    plazo_val = 15
    unidad_plazo = "DIAS"
    tipo_persona = "NATURAL"
    valor_total = 4500000

    obligaciones = build_obligaciones_empresa(
        objeto=objeto,
        tipologia=tipologia,
        plazo_str=plazo_str,
        plazo_val=plazo_val,
        unidad_plazo=unidad_plazo,
        tipo_persona=tipo_persona,
        valor_total=valor_total
    )

    full_text = " ".join(obligaciones).lower()
    
    # Verificaciones
    assert "entrega física" in full_text or "entrega material" in full_text, "Falta entrega física de equipos"
    assert "fallas" in full_text or "diagnósticos" in full_text, "Falta reporte de fallas"
    assert "pruebas técnicas" in full_text or "pruebas de funcionamiento" in full_text, "Falta verificación de pruebas técnicas"
    assert "plataforma" not in full_text, "ERROR: Menciona plataformas en mantenimiento de herramientas"
    assert "honorarios" not in full_text, "ERROR: Menciona honorarios en mantenimiento"
    assert "informes mensuales" not in full_text, "ERROR: Menciona informes mensuales con plazo de 15 días"
    assert "pagos mensuales" not in full_text, "ERROR: Menciona pagos mensuales con plazo de 15 días"
    assert "informe final" in full_text, "Debe mencionar informe final para plazo de 15 días"

    forma_pago = build_forma_pago_clausula(
        valor_total=valor_total,
        plazo_str=plazo_str,
        plazo_val=plazo_val,
        unidad_plazo=unidad_plazo,
        tipologia=tipologia,
        tipo_persona=tipo_persona
    )
    fp_lower = forma_pago.lower()
    assert "honorarios" not in fp_lower
    assert "mensual" not in fp_lower
    assert "pago único" in fp_lower
    assert "informe final" in fp_lower

    print("PASS: Mantenimiento 15 días cumplió todas las directrices.")
    print("Muestra de obligaciones:")
    for i, o in enumerate(obligaciones, 1):
        print(f"  {i}. {o}")
    print(f"\nForma de pago:\n  {forma_pago}")

def test_pgirs_juridica():
    print("\n--- 3. Testing PGIRS / Sensibilización 6 meses (Persona Jurídica) ---")
    objeto = "PRESTACIÓN DE SERVICIOS PARA LA SENSIBILIZACIÓN AMBIENTAL Y SOCIALIZACIÓN DE LA RUTA SELECTIVA EN EL MARCO DEL PGIRS DE LA ESPY"
    tipologia = "PRESTACION_SERVICIOS"
    plazo_str = "6 meses"
    plazo_val = 6
    unidad_plazo = "MESES"
    tipo_persona = "JURIDICA"
    valor_total = 30000000

    obligaciones = build_obligaciones_empresa(
        objeto=objeto,
        tipologia=tipologia,
        plazo_str=plazo_str,
        plazo_val=plazo_val,
        unidad_plazo=unidad_plazo,
        tipo_persona=tipo_persona,
        valor_total=valor_total
    )

    full_text = " ".join(obligaciones).lower()
    assert "rutas" in full_text or "programación oficial" in full_text
    assert "pgirs" in full_text or "lineamientos" in full_text
    assert "parafiscales" in full_text or "artículo 50 de la ley 789 de 2002" in full_text
    assert "informe periódico" in full_text or "informes de ejecución" in full_text or "informe" in full_text
    assert "honorarios" not in full_text

    print("PASS: PGIRS Jurídica cumplió las especificaciones.")

def test_supply_and_works():
    print("\n--- 4. Testing Suministro y Obra ---")
    # Suministro
    ob_sum = build_obligaciones_empresa(
        objeto="SUMINISTRO DE REACTIVOS QUÍMICOS Y CLORO PARA LA PLANTA DE TRATAMIENTO DE AGUA POTABLE DE LA ESPY",
        tipologia="SUMINISTRO",
        plazo_str="2 meses",
        plazo_val=2,
        unidad_plazo="MESES",
        tipo_persona="JURIDICA"
    )
    text_sum = " ".join(ob_sum).lower()
    assert "almacén" in text_sum or "ingreso a almacén" in text_sum
    assert "fichas técnicas" in text_sum or "especificaciones técnicas" in text_sum
    assert "honorarios" not in text_sum

    # Obra
    ob_obra = build_obligaciones_empresa(
        objeto="EJECUCIÓN DE OBRAS CIVILES PARA LA OPTIMIZACIÓN DE LA RED DE ALCANTARILLADO DE LA ESPY",
        tipologia="OBRA",
        plazo_str="3 meses",
        plazo_val=3,
        unidad_plazo="MESES",
        tipo_persona="JURIDICA"
    )
    text_obra = " ".join(ob_obra).lower()
    assert "entrega de terreno" in text_obra
    assert "redes" in text_obra
    assert "cantidades de obra" in text_obra
    assert "honorarios" not in text_obra

    print("PASS: Suministro y Obra probados exitosamente.")

def test_docx_generation():
    print("\n--- 5. Testing Generación Real de Documentos DOCX ---")
    # Mock de contrato tipo mantenimiento 15 días
    class MockContractor:
        tipo_persona = "NATURAL"
        nombre = "TALLER Y REPUESTOS JUAN PEREZ"
        razon_social = None
        tipo_documento = "CC"
        numero_documento = "12345678"
        digito_verificacion = None
        representante_legal_nombre = None
        representante_legal_documento = None
        ciudad = "Yolombó"
        direccion = "Calle 20 # 18-10"
        telefono = "3001234567"
        correo = "juan.perez@example.com"
        regimen_tributario = "NO_RESPONSABLE_IVA"

    class MockContract:
        numero_contrato = "CPS-TEST-001-2026"
        vigencia = 2026
        tipologia = "PRESTACION_SERVICIOS"
        modalidad_seleccion = "CONTRATACION_DIRECTA"
        tipo_contrato = "PRESTACIÓN DE SERVICIOS"
        objeto = "MANTENIMIENTO PREVENTIVO Y CORRECTIVO A TODO COSTO DE HERRAMIENTAS ELÉCTRICAS Y EQUIPOS MENORES ADSCRITOS AL ÁREA OPERATIVA DE LA EMPRESA DE SERVICIOS PÚBLICOS DE YOLOMBÓ E.S.P."
        justificacion_necesidad = "La ESPY requiere asegurar la operatividad de sus equipos de campo..."
        marco_juridico = "El presente contrato se rige por la Ley 142 de 1994..."
        objetivo_general = "Garantizar el óptimo funcionamiento técnico de las herramientas y equipos electromecánicos de la ESPY."
        objetivos_especificos = [
            "Diagnosticar el estado técnico y mecánico de las herramientas entregadas.",
            "Ejecutar los mantenimientos preventivos y correctivos a todo costo con repuestos nuevos.",
            "Realizar pruebas técnicas de operatividad y certificar la entrega a satisfacción."
        ]
        obligaciones_estudio_previo = [
            "Realizar diagnóstico de cada herramienta",
            "Suministrar repuestos originales y mano de obra a todo costo",
            "Entregar ficha técnica y registro fotográfico"
        ]
        obligaciones_minuta = [
            "Ejecutar el mantenimiento preventivo y correctivo con altos estándares de calidad",
            "Suministrar e instalar repuestos nuevos y originales a todo costo",
            "Efectuar pruebas de funcionamiento previas a la entrega",
            "Presentar informe final de ejecución con registro fotográfico"
        ]
        # Dejamos obligaciones_empresa como None para probar la generación dinámica
        obligaciones_empresa = None
        valor_total = 3500000
        plazo_ejecucion = 15
        unidad_plazo = "DIAS"
        supervisor_nombre = "ING. CARLOS MARÍN"
        supervisor_cargo = "DIRECTOR OPERATIVO"
        supervisor_dependencia = "DIRECCIÓN OPERATIVA"
        supervisor_tipo_documento = "CC"
        supervisor_documento = "98765432"
        fecha_suscripcion = None
        fecha_acta_inicio = None
        fecha_terminacion = None
        cdp_numero = "CDP-2026-050"
        cdp_fecha = None
        rp_numero = "RP-2026-050"
        rp_fecha = None
        lugar_ejecucion = "Municipio de Yolombó, Antioquia"
        contractor = MockContractor()

    contract = MockContract()
    
    # Generar Minuta
    print("Generando Minuta DOCX...")
    minuta_path = generate_docx_for_contract(contract, "MINUTA")
    assert os.path.exists(minuta_path), f"Archivo Minuta no encontrado: {minuta_path}"
    
    # Generar Estudio Previo
    print("Generando Estudio Previo DOCX...")
    ep_path = generate_docx_for_contract(contract, "ESTUDIO_PREVIO")
    assert os.path.exists(ep_path), f"Archivo Estudio Previo no encontrado: {ep_path}"

    # Inspeccionar contenido de la minuta generada
    import docx
    min_doc = docx.Document(minuta_path)
    min_text = "\n".join([p.text for p in min_doc.paragraphs])
    
    assert "CLÁUSULA SEXTA. - OBLIGACIONES DE LA EMPRESA" in min_text or "SEXTA. OBLIGACIONES DE LA EMPRESA" in min_text or "OBLIGACIONES DE LA EMPRESA" in min_text, "No se encontró encabezado de obligaciones de la empresa en la minuta"
    assert "entrega física" in min_text.lower() or "entrega material" in min_text.lower(), "Falta obligación de entrega física en minuta"
    assert "plataforma" not in min_text.lower(), "ERROR: La minuta contiene mención a 'plataformas'"
    assert "honorarios" not in min_text.lower(), "ERROR: La minuta contiene mención a 'honorarios'"
    assert "informes mensuales" not in min_text.lower(), "ERROR: La minuta contiene 'informes mensuales' en contrato de 15 días"
    
    # Inspeccionar contenido del estudio previo generado
    ep_doc = docx.Document(ep_path)
    ep_text = "\n".join([p.text for p in ep_doc.paragraphs])
    assert "OBLIGACIONES DE LA EMPRESA (ESPY)" in ep_text or "OBLIGACIONES DE LA EMPRESA" in ep_text, "No se encontró sección de obligaciones de la empresa en el estudio previo"
    assert "entrega física" in ep_text.lower() or "entrega material" in ep_text.lower(), "Falta obligación de entrega física en estudio previo"
    assert "plataforma" not in ep_text.lower(), "ERROR: El estudio previo contiene mención a 'plataformas'"
    assert "honorarios" not in ep_text.lower(), "ERROR: El estudio previo contiene 'honorarios'"
    assert "informes mensuales" not in ep_text.lower(), "ERROR: El estudio previo contiene 'informes mensuales' para plazo de 15 días"

    print(f"PASS: Generación de DOCX validada exitosamente.")
    print(f"  Minuta generada en: {minuta_path}")
    print(f"  Estudio Previo generado en: {ep_path}")

if __name__ == '__main__':
    test_short_term_logic()
    test_maintenance_short_term()
    test_pgirs_juridica()
    test_supply_and_works()
    test_docx_generation()
    print("\n==========================================")
    print("TODAS LAS PRUEBAS PASARON EXITOSAMENTE (5/5)")
    print("==========================================")

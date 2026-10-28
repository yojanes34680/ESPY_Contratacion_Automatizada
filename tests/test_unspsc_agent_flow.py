"""Test suite de verificación integral para el Agente Clasificador UNSPSC.
Valida el catálogo oficial (52.845 registros), la selección jerárquica estricta,
la persistencia en contratos y la generación del dictamen Word (.docx).
"""
import sys
from pathlib import Path

# Add backend to path
sys.path.insert(0, str(Path("backend").resolve()))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import docx
from app.database import SessionLocal
from app.models.contract import Contrato
from app.schemas.unspsc import UnspscAnalysisInput
from app.services.agents.unspsc_agent import unspsc_agent

def test_unspsc_agent_full_pipeline():
    print("\n=======================================================")
    print(" INICIANDO TEST DEL AGENTE CLASIFICADOR DE CÓDIGOS UNSPSC")
    print("=======================================================")

    # 1. VERIFICACIÓN DEL CATÁLOGO OFICIAL (52.845 CÓDIGOS)
    print("\n--- 1. Verificación del Catálogo Oficial ---")
    catalog_size = unspsc_agent.get_catalog_size()
    print(f"Total códigos en catálogo: {catalog_size}")
    assert catalog_size == 52845, f"ERROR: Se esperaban 52.845 códigos y se encontraron {catalog_size}."
    print("✓ Confirmado: 52.845 códigos oficiales cargados desde codigos unspsc_contratacion.xlsx.")

    # 2. PRUEBA DE BÚSQUEDA Y MAPEO EN TIEMPO REAL
    print("\n--- 2. Prueba de Búsqueda y Mapeo en el Catálogo ---")
    query_code = "7210"
    results_code = unspsc_agent.search_catalog(query_code, limit=5)
    assert len(results_code) > 0, f"No se encontraron resultados para {query_code}"
    print(f"Búsqueda por prefijo '{query_code}': {len(results_code)} resultados encontrados.")
    for r in results_code[:2]:
        print(f"   -> {r['codigo']} | {r['descripcion']} (Seg: {r['segmento']}, Fam: {r['familia']})")

    query_text = "acueducto"
    results_text = unspsc_agent.search_catalog(query_text, limit=5)
    assert len(results_text) > 0, f"No se encontraron resultados para {query_text}"
    print(f"Búsqueda textual '{query_text}': {len(results_text)} resultados encontrados.")
    for r in results_text[:2]:
        print(f"   -> {r['codigo']} | {r['descripcion']}")
    print("✓ Confirmado: Motor de búsqueda y jerarquización operativo.")

    # 3. ANÁLISIS INTEGRAL DE CONTRATACIÓN (CASO REAL ESPY)
    print("\n--- 3. Análisis Integral del Proceso Contractual ---")
    payload = UnspscAnalysisInput(
        objeto="PRESTACIÓN DE SERVICIOS PARA EL APOYO A LA GESTIÓN EN EL MANTENIMIENTO, MONITOREO Y OPERACIÓN DE LAS PLANTAS DE TRATAMIENTO DE AGUA POTABLE Y RESIDUALES",
        estudio_previo=(
            "La ESP Yolombó requiere garantizar la calidad del agua potable y el saneamiento básico conforme a la Ley 142 de 1994. "
            "Se requiere un perfil técnico para monitoreo físico-químico, dosificación de coagulantes, desinfección, lavado de filtros, "
            "purga de lodos y toma de muestras de agua potable y residual."
        ),
        propuesta_economica=(
            "Actividad 1: Muestreo y caracterización de agua potable ($1.800.000). "
            "Actividad 2: Inspección y mantenimiento preventivo de lechos filtrantes ($2.500.000). "
            "Actividad 3: Operación de sistemas de cloración y bombeo ($3.200.000)."
        )
    )

    result = unspsc_agent.analyze_contract_unspsc(payload)
    print("\n[RESULTADO DEL AGENTE UNSPSC]")
    print(f"Resumen del análisis (máx 4 líneas):\n\"{result.resumen_analisis}\"\n")
    print(f"Total códigos seleccionados: {result.total_codigos}")

    # Validaciones obligatorias de las instrucciones del usuario:
    assert result.resumen_analisis, "El resumen del análisis no debe estar vacío."
    lineas = [l for l in result.resumen_analisis.split("\n") if l.strip()]
    assert len(lineas) <= 4, f"El resumen debe tener máximo 4 líneas (tiene {len(lineas)})."
    assert len(result.codigos_seleccionados) > 0, "Debe seleccionar al menos un código UNSPSC."

    print("\nCódigos UNSPSC Seleccionados:")
    for idx, c in enumerate(result.codigos_seleccionados, 1):
        print(f"{idx}. [{c.codigo_unspsc}] {c.descripcion_codigo}")
        print(f"   Jerarquía: Seg={c.segmento}, Fam={c.familia}, Clase={c.clase}")
        print(f"   Justificación: {c.justificacion_eleccion[:100]}...\n")

        # VALIDACIÓN CRÍTICA: Ningún código puede ser inventado
        assert len(c.codigo_unspsc) == 8, f"El código {c.codigo_unspsc} debe tener 8 dígitos."
        assert c.codigo_unspsc in unspsc_agent._catalog, (
            f"ERROR GRAVE: El código {c.codigo_unspsc} no existe en codigos unspsc_contratacion.xlsx!"
        )
        assert c.descripcion_codigo == unspsc_agent._catalog[c.codigo_unspsc], (
            f"ERROR: La descripción '{c.descripcion_codigo}' no coincide exactamente con el archivo oficial!"
        )
        assert c.justificacion_eleccion, f"El código {c.codigo_unspsc} debe incluir justificación técnica."

    print("✓ Confirmado: Todos los códigos pertenecen 100% estrictamente a la base de datos oficial.")

    # 4. PERSISTENCIA EN CONTRATOS DE LA BASE DE DATOS
    print("\n--- 4. Persistencia en Contrato de la Base de Datos ---")
    db = SessionLocal()
    try:
        contrato = db.query(Contrato).first()
        assert contrato is not None, "Debe existir al menos un contrato en la base de datos."
        print(f"Asignando al contrato ID={contrato.id}, Nro={contrato.numero_contrato}")

        contrato.codigos_unspsc = result.model_dump()
        db.commit()
        db.refresh(contrato)

        assert contrato.codigos_unspsc is not None, "No se guardó codigos_unspsc en el contrato."
        assert len(contrato.codigos_unspsc["codigos_seleccionados"]) == result.total_codigos
        print("✓ Confirmado: Persistencia en base de datos SQLite validada.")

        # 5. GENERACIÓN DEL DICTAMEN OFICIAL WORD (.DOCX)
        print("\n--- 5. Generación de Dictamen Oficial en Word (.docx) ---")
        docx_path = unspsc_agent.generate_unspsc_report_docx(
            result, 
            objeto=payload.objeto, 
            numero_contrato=contrato.numero_contrato
        )
        assert Path(docx_path).exists(), f"El archivo {docx_path} no fue creado."
        print(f"Documento generado: {docx_path}")

        doc = docx.Document(str(docx_path))
        doc_text = "\n".join(p.text for p in doc.paragraphs)
        assert "EMPRESA DE SERVICIOS PÚBLICOS DE YOLOMBÓ" in doc_text
        assert "DICTAMEN Y CLASIFICACIÓN DE CÓDIGOS UNSPSC" in doc_text
        assert len(doc.tables) >= 2, "Debe contener al menos 2 tablas (metadatos y códigos seleccionados)."
        print("✓ Confirmado: Dictamen Word (.docx) generado con membrete institucional y tablas estilizadas.")

    finally:
        db.close()

    # 6. PRUEBA ESPECÍFICA: CONTRATO DE INFRAESTRUCTURA HIDROSANITARIA / ACUEDUCTO (CASO LAS CAMELIAS)
    print("\n--- 6. Prueba Específica: Contrato de Válvula de Sectorización (Las Camelias) ---")
    payload_camelias = UnspscAnalysisInput(
        objeto="INSTALACION DE VALVULA DE SECTORIZACION EN LA RED DE DISTRIBUCION DE ACUEDUCTO DEL SECTOR LAS CAMELIAS DEL MUNICIPIO DE YOLOMBO",
        estudio_previo="Se requiere la instalacion de valvula de sectorizacion, excavacion a cielo abierto, empalme de tuberia a presion RDE 21 y reposicion de pavimento en concreto hidraulico.",
        propuesta_economica="""CONSTRUCTORA E INGENIERIA HIDRAULICA S.A.S.
NIT: 901.234.567-8
1. LOCALIZACION Y REPLANTEO, EXCAVACION MANUAL EN MATERIAL COMUN
2. SUMINISTRO E INSTALACION DE VALVULA DE SECTORIZACION Y TUBERIA A PRESION RDE 21
3. ROTURA Y REPOSICION DE PAVIMENTO EN CONCRETO
"""
    )
    result_camelias = unspsc_agent.analyze_contract_unspsc(payload_camelias)
    cat_id, cat_nombre = unspsc_agent.classify_contract_thematic_category(payload_camelias.objeto, payload_camelias.propuesta_economica)
    assert cat_id == "D", f"ERROR: Se esperaba Categoría D y se obtuvo {cat_id}"
    print(f"✓ Categoría clasificada: {cat_nombre} ({cat_id})")

    selected_codes = [c.codigo_unspsc for c in result_camelias.codigos_seleccionados]
    print(f"Códigos seleccionados: {selected_codes}")

    # Validar presencia de los 3 componentes del APU:
    # 1. Hidrosanitario (72141119, 72141121, 72141124 o 72151100)
    has_hidro = any(c in ["72141119", "72141121", "72141124", "72151100"] for c in selected_codes)
    assert has_hidro, "ERROR: Falta el componente hidrosanitario (acueducto/tubería agua potable)!"

    # 2. Movimiento de tierras (72141511, 72141505 o 72141510)
    has_tierras = any(c in ["72141511", "72141505", "72141510"] for c in selected_codes)
    assert has_tierras, "ERROR: Falta el componente de movimiento de tierras (excavación)!"

    # 3. Pavimentos / Vías (72141103, 72141104 o 72141105)
    has_pav = any(c in ["72141103", "72141104", "72141105"] for c in selected_codes)
    assert has_pav, "ERROR: Falta el componente de pavimentos / vías!"

    # Validar ausencia de códigos incompatibles o de drenaje/alcantarillado
    banned_codes = {"72153103", "72152303", "72151906", "72141120", "72141205"}
    for bc in banned_codes:
        assert bc not in selected_codes, f"ERROR GRAVE: Se seleccionó el código prohibido {bc}!"

    print("✓ Confirmado: Mapeo de 3 a 4 componentes del APU balanceado (Acueducto, Excavación, Pavimento) sin códigos incompatibles.")

    # 7. PRUEBA ESPECÍFICA: REGLA DE TIPOLOGÍA PRIMARIA - CONTRATO DE SUMINISTRO PGIRS (BIENES FÍSICOS)
    print("\n--- 7. Prueba Específica: Contrato de Suministro PGIRS (Bienes Físicos) ---")
    payload_suministro_pgirs = UnspscAnalysisInput(
        objeto="SUMINISTRO DE INSUMOS, MATERIALES, DOTACIONES, EQUIPOS Y LOGISTICA PARA EL PGIRS DE LA EMPRESA DE SERVICIOS PUBLICOS DE YOLOMBO",
        estudio_previo="Se requiere adquirir insumos fisicos para el desarrollo de las actividades del PGIRS, incluyendo canecas para disposicion de residuos, dotaciones de recicladores, vallas informativas, plantulas para reforestacion y equipos de audio.",
        propuesta_economica="""DISTRIBUIDORA Y SUMINISTROS AMBIENTALES S.A.S.
1. Canecas plasticas de 53 galones con tapa y baldes de basura para recoleccion
2. Dotaciones completas para recicladores (overoles de proteccion, guantes de carnaza y botas de seguridad con puntera)
3. Vallas publicitarias de zinc y volantes informativos de separacion en la fuente
4. Plantulas ornamentales y especies vivas para jornadas ambientales
5. Equipos de sonido portatil, bafles y altoparlantes para socializacion
"""
    )
    result_suministro = unspsc_agent.analyze_contract_unspsc(payload_suministro_pgirs)
    cat_id_sum, cat_nombre_sum = unspsc_agent.classify_contract_thematic_category(payload_suministro_pgirs.objeto, payload_suministro_pgirs.propuesta_economica)
    assert cat_id_sum == "C", f"ERROR: Se esperaba Categoría C para Suministro y se obtuvo {cat_id_sum}"
    print(f"✓ Categoría clasificada: {cat_nombre_sum} ({cat_id_sum})")

    selected_sum_codes = [c.codigo_unspsc for c in result_suministro.codigos_seleccionados]
    print(f"Códigos seleccionados: {selected_sum_codes}")

    # Verificar prohibición absoluta de códigos de servicios (>= 70)
    for c in selected_sum_codes:
        seg_num = int(c[:2])
        assert seg_num < 70, f"ERROR GRAVE: En contrato de Suministro se asignó un código de servicio ({c}, Segmento {seg_num})!"

    # Verificar presencia de productos físicos (47 canecas, 46 dotaciones, 55 vallas/señalización, 10 plantas, 52 audio)
    has_canecas = any(c.startswith("47") or c.startswith("2410") for c in selected_sum_codes)
    has_dotacion = any(c.startswith("46") for c in selected_sum_codes)
    has_vallas = any(c.startswith("55") or c.startswith("56") for c in selected_sum_codes)
    has_plantulas = any(c.startswith("10") for c in selected_sum_codes)
    has_audio = any(c.startswith("52") or c.startswith("45") for c in selected_sum_codes)

    assert has_canecas, "ERROR: Falta código de canecas/recipientes de residuos (Segmento 47)!"
    assert has_dotacion, "ERROR: Falta código de dotaciones/ropa de seguridad (Segmento 46)!"
    assert has_vallas, "ERROR: Falta código de vallas/señalización (Segmento 55 o 56)!"
    assert has_plantulas, "ERROR: Falta código de plántulas/vivero (Segmento 10)!"
    assert has_audio, "ERROR: Falta código de equipos de audio/sonido (Segmento 52 o 45)!"

    # Validación de Coincidencia Semántica (Match) en la Justificación
    for item in result_suministro.codigos_seleccionados:
        code_s = item.codigo_unspsc
        j_text = item.justificacion_eleccion.lower()
        if code_s.startswith("47"):
            assert any(w in j_text for w in ["caneca", "balde", "recipiente", "contenedor"]), f"ERROR: Código 47 ({code_s}) no cita caneca/balde en justificación: {item.justificacion_eleccion}"
        elif code_s.startswith("46"):
            assert any(w in j_text for w in ["dotacion", "overol", "guante", "bota", "seguridad"]), f"ERROR: Código 46 ({code_s}) no cita dotación/ropa en justificación: {item.justificacion_eleccion}"
        elif code_s.startswith("55") or code_s.startswith("56"):
            assert any(w in j_text for w in ["valla", "senalamiento", "senal", "volante", "informativo"]), f"ERROR: Código 55/56 ({code_s}) no cita vallas/señalización en justificación: {item.justificacion_eleccion}"
        elif code_s.startswith("10"):
            assert any(w in j_text for w in ["plantula", "planta", "especies vivas", "ornamental", "vivero"]), f"ERROR: Código 10 ({code_s}) no cita plántulas en justificación: {item.justificacion_eleccion}"
        elif code_s.startswith("52") or code_s.startswith("45"):
            assert any(w in j_text for w in ["sonido", "audio", "bafle", "altavoz", "altoparlante"]), f"ERROR: Código 52 ({code_s}) no cita audio/bafles en justificación: {item.justificacion_eleccion}"

    print("✓ Confirmado: Desglose perfecto de 5 familias de bienes manufacturados sin códigos de consultoría ni servicios.")
    print("✓ Confirmado: Emparejamiento semántico 1 a 1 verificado en las justificaciones de cada código (sin desfase de índice).")

    print("\n=======================================================")
    print(" ¡TODAS LAS PRUEBAS DEL AGENTE UNSPSC PASARON AL 100%! ")
    print("=======================================================\n")

if __name__ == "__main__":
    test_unspsc_agent_full_pipeline()

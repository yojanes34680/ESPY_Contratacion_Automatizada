import sys
from pathlib import Path

# Ajuste de path para ejecución dentro de tests/
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import docx
import re
import pytest

@pytest.mark.skip(reason="usa datos reales; se reactiva en el prompt 1")
def test_cs056_documents():

    docs_to_check = [
        Path("backend/storage/generated/Estudio_Previo_Suministro_CS-056-2026_2026.docx"),
        Path("backend/storage/generated/Minuta_Suministro_CS-056-2026_2026.docx"),
        Path("backend/storage/generated/Acta_Inicio_CS-056-2026_2026.docx"),
        Path("backend/storage/generated/Matriz_Riesgos_ESPY_SUMINISTRO__INSTALACI_N_Y_PUES.docx"),
    ]

    prohibited_strings = [
        "R+YPLAN",
        "VAVANCED",
        "SERVICIO GPS PLAN R+",
        "desensamble integral",
        "devanados",
        "rodamientos",
        "aislamiento dieléctrico",
        "aislamiento dielectrico",
        "inexistencia de personal de planta",
        "RL. CARLOS IVAN ESCOBAR\n",
        "RL. CARLOS IVAN ESCOBAR\t",
        "por la otra la MONITOREO",
        ".. Este contrato",
    ]

    for p in docs_to_check:
        assert p.exists(), f"El archivo {p} no existe."
        doc = docx.Document(str(p))
        full_text = " ".join(p_elem.text for p_elem in doc.paragraphs)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    full_text += " " + c.text

        print(f"\n--- Verificando {p.name} ---")
        for bad in prohibited_strings:
            if bad in full_text:
                raise AssertionError(f"ERROR: Encontrado texto prohibido '{bad}' en {p.name}")
        print(f"PASS: {p.name} no contiene ningún texto prohibido.")

    # Verificaciones específicas
    # 1. Estudio Previo
    doc_ep = docx.Document("backend/storage/generated/Estudio_Previo_Suministro_CS-056-2026_2026.docx")
    ep_txt = " ".join(p.text for p in doc_ep.paragraphs)
    assert "dotación tecnológica y fortalecimiento operativo" in ep_txt, "No se encontró el texto de conveniencia corregido en Estudio Previo."
    assert "Suministrar e instalar" in ep_txt or "Suministrar e instalar" in ep_txt, "Objetivo específico mal capitalizado."
    assert "reparación oportuna, eficiente y en términos de economía del acueducto y alcantarillado" not in ep_txt

    # 2. Minuta Suministro
    doc_min = docx.Document("backend/storage/generated/Minuta_Suministro_CS-056-2026_2026.docx")
    min_txt = " ".join(p.text for p in doc_min.paragraphs)
    assert "vecina del Municipio de Yolombó" in min_txt or "vecina del Municipio de Yolomb" in min_txt, "Género de la Gerente no corregido."
    assert "identificada con la cédula" in min_txt, "identificada con la cédula no corregido."
    assert "debidamente facultada" in min_txt, "debidamente facultada no corregido."
    assert "por la otra MONITOREO INTELIGENTE S. A. S." in min_txt, "por la otra MONITOREO no corregido."
    assert "R.L. CARLOS IVAN ESCOBAR GONZALEZ" in min_txt or "CARLOS IVAN ESCOBAR GONZALEZ - R.L." in min_txt or "CARLOS IVAN ESCOBAR GONZALEZ" in min_txt
    assert "Representante Legal Contratista:" in " ".join(c.text for t in doc_min.tables for r in t.rows for c in r.cells), "Representante Legal en tabla no reubicado."

    # 3. Acta de Inicio
    doc_acta = docx.Document("backend/storage/generated/Acta_Inicio_CS-056-2026_2026.docx")
    acta_txt = " ".join(p.text for p in doc_acta.paragraphs)
    assert "CARLOS IVAN ESCOBAR GONZALEZ" in acta_txt, "Nombre de representante legal truncado en Acta de Inicio."
    assert "por la otra MONITOREO INTELIGENTE S. A. S." in acta_txt, "por la otra MONITOREO no corregido en Acta de Inicio."

    # 4. Matriz de Riesgos
    doc_risk = docx.Document("backend/storage/generated/Matriz_Riesgos_ESPY_SUMINISTRO__INSTALACI_N_Y_PUES.docx")
    risk_txt = " ".join(p.text for p in doc_risk.paragraphs)
    for t in doc_risk.tables:
        for r in t.rows:
            for c in r.cells:
                risk_txt += " " + c.text

    assert "unidades GPS LC" in risk_txt or "cámaras vehiculares JC371 ADAS" in risk_txt or "dispositivos GPS" in risk_txt, "Riesgos específicos de GPS no encontrados."
    assert "pago contra entrega física e instalación efectiva de los equipos" in risk_txt, "Observación técnica de suministro en Matriz de Riesgos no corregida."

    print("\n¡TODAS LAS VALIDACIONES DE TEXTO Y REGLAS DE CONTRATO CS-056-2026 PASARON EXITOSAMENTE (100%)!")

if __name__ == "__main__":
    test_cs056_documents()

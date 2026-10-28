"""Test script for Contract Addition Flow (Agente de Adiciones a Contratos)."""
import os
import sys
import docx
from pathlib import Path

# Add backend to path
sys.path.insert(0, os.path.abspath("backend"))

from fastapi.testclient import TestClient
from app.main import app
from app.database import SessionLocal
from app.models.contract import Contrato
from app.models.contractor import Contratista
from app.models.contract_addition import AdicionContrato

client = TestClient(app)

def run_tests():
    print("==================================================")
    print("INICIANDO PRUEBAS DEL AGENTE DE ADICIONES")
    print("==================================================")

    # 1. Test Contracts List endpoint
    print("\n1. Probando /api/additions/contracts-list...")
    resp = client.get("/api/additions/contracts-list")
    assert resp.status_code == 200, f"Error en contracts-list: {resp.status_code} - {resp.text}"
    contracts = resp.json()
    print(f"   -> Contratos disponibles: {len(contracts)}")
    assert len(contracts) > 0, "Debe haber contratos disponibles"

    # Seleccionar primer contrato
    target_c = contracts[0]
    c_id = target_c["id"]
    c_num = target_c["numero_contrato"]
    print(f"   -> Contrato seleccionado: {c_num} (ID: {c_id})")

    # 2. Test Contract Context endpoint
    print(f"\n2. Probando /api/additions/contract-context/{c_id}...")
    resp_ctx = client.get(f"/api/additions/contract-context/{c_id}")
    assert resp_ctx.status_code == 200, f"Error en contract-context: {resp_ctx.text}"
    ctx = resp_ctx.json()
    print(f"   -> Objeto: {ctx['objeto'][:60]}...")
    print(f"   -> Valor Total Actual: ${ctx['valor_total_actual']:,.0f}")
    print(f"   -> Siguiente Adición: {ctx['siguiente_numero_adicion']}")
    print(f"   -> Contratista: {ctx['contratista']['nombre_o_razon_social']}")

    # 3. Test Justification Generation
    print("\n3. Probando generación de justificación con IA...")
    resp_just = client.post("/api/additions/generate-justification", json={
        "contrato_id": c_id,
        "valor_adicion": 5000000,
        "necesidad_usuario": "Mayor volumen de usuarios en el acueducto municipal durante época de sequía."
    })
    assert resp_just.status_code == 200, f"Error generando justificación: {resp_just.text}"
    just_data = resp_just.json()
    justificacion_generada = just_data["justificacion"]
    print(f"   -> Justificación generada ({len(justificacion_generada)} caracteres):")
    print(f"      \"{justificacion_generada[:120]}...\"")

    # 4. Test Creation of Addition with Legal Rep & CDP
    print("\n4. Probando creación de adición y generación de minuta Word...")
    addition_payload = {
        "contrato_id": c_id,
        "numero_adicion": ctx["siguiente_numero_adicion"],
        "valor_adicion": 5000000.0,
        "justificacion": justificacion_generada,
        "cdp_numero": "00217",
        "cdp_fecha": "2026-03-20",
        "cdp_rubro": "2.3.2.02.02.005",
        "cdp_rubro_nombre": "Servicios de Mantenimiento y Operación ESP",
        "cdp_valor": 12000000.0,
        "representante_legal_nombre": "EDWARD MANUEL HERNANDEZ VALENCIA",
        "representante_legal_documento": "CC 1017128843",
        "guardar_en_contratista": True,
        "fecha_adicion": "2026-04-15"
    }

    resp_create = client.post("/api/additions/create", json=addition_payload)
    assert resp_create.status_code == 200, f"Error creando adición: {resp_create.text}"
    res = resp_create.json()
    print(f"   -> Status: {res['status']}")
    print(f"   -> Addition ID: {res['addition_id']}")
    print(f"   -> Nuevo Total Contractual: ${res['valor_nuevo_total']:,.0f}")
    print(f"   -> Download URL: {res['download_url']}")

    addition_id = res["addition_id"]

    # 5. Test Download of generated Word document
    print(f"\n5. Probando descarga de Minuta Word /api/additions/{addition_id}/download-docx...")
    resp_down = client.get(f"/api/additions/{addition_id}/download-docx")
    assert resp_down.status_code == 200, f"Error descargando docx: {resp_down.status_code}"
    docx_bytes = resp_down.content
    print(f"   -> Tamaño de archivo Word: {len(docx_bytes):,} bytes")
    assert len(docx_bytes) > 5000, "El archivo Word debe tener contenido válido"

    # Save to scratch and inspect with python-docx
    test_docx_path = Path("test_adicion_output.docx")
    with open(test_docx_path, "wb") as f:
        f.write(docx_bytes)

    doc_check = docx.Document(str(test_docx_path))
    doc_text = " ".join([p.text for p in doc_check.paragraphs])
    print(f"   -> Párrafos en el documento Word: {len(doc_check.paragraphs)}")
    
    # Assertions on text content
    assert "ADICIÓN NRO." in doc_text, "Debe contener ADICIÓN NRO."
    assert "CAROLINA GIRALDO GÓMEZ" in doc_text, "Debe mencionar a Carolina Giraldo Gómez"
    assert "CONSIDERACIONES" in doc_text, "Debe contener CONSIDERACIONES"
    assert "ACUERDAN" in doc_text, "Debe contener ACUERDAN"
    assert "00217" in doc_text, "Debe contener el número de CDP"
    assert "EDWARD MANUEL HERNANDEZ VALENCIA" in doc_text, "Debe incluir el representante legal"
    assert "1017128843" in doc_text, "Debe incluir la cédula del representante legal"
    
    print("   -> Contenido del documento Word 100% verificado.")

    # Cleanup temp test file
    if test_docx_path.exists():
        test_docx_path.unlink()

    # 6. Verify Database Persistence
    print("\n6. Verificando persistencia en base de datos SQLite...")
    db = SessionLocal()
    try:
        ad_record = db.query(AdicionContrato).filter(AdicionContrato.id == addition_id).first()
        assert ad_record is not None, "Registro de adición debe existir en BD"
        assert ad_record.valor_adicion == 5000000.0
        assert ad_record.cdp_numero == "00217"
        assert ad_record.representante_legal_nombre == "EDWARD MANUEL HERNANDEZ VALENCIA"

        c_updated = db.query(Contrato).filter(Contrato.id == c_id).first()
        assert c_updated.adicion_1 >= 5000000.0, "Contrato adicion_1 debe estar actualizado"
        print(f"   -> Contrato {c_num} adicion_1 actualizado a: ${c_updated.adicion_1:,.0f}")
        print(f"   -> Contrato {c_num} valor_total actualizado a: ${c_updated.valor_total:,.0f}")

        # Check contractor legal rep was updated
        c_contr = db.query(Contratista).filter(Contratista.id == c_updated.contratista_id).first()
        assert c_contr.representante_legal_nombre == "EDWARD MANUEL HERNANDEZ VALENCIA"
        print(f"   -> Contratista actualizado con representante legal: {c_contr.representante_legal_nombre}")
    finally:
        db.close()

    print("\n==================================================")
    print(">>> TODAS LAS PRUEBAS DEL AGENTE PASARON CON EXITO (100%)")
    print("==================================================")

if __name__ == "__main__":
    run_tests()

import os
import sys
from pathlib import Path

# Add backend to path
sys.path.insert(0, str(Path("backend").resolve()))

from app.database import SessionLocal
from app.models.contract import Contrato
from app.services.docx_engine import (
    generate_docx_for_contract,
    get_active_template_path,
    inspect_docx_placeholders
)
import docx

def test_custom_estudio_previo_replacement():
    db = SessionLocal()
    try:
        active_path, is_custom = get_active_template_path("ESTUDIO_PREVIO")
        print(f"Active template for ESTUDIO_PREVIO: {active_path} (is_custom: {is_custom})")
        placeholders = inspect_docx_placeholders(active_path)
        print(f"Detected placeholders in template ({len(placeholders)}): {placeholders}")
        
        # Verify placeholders with spaces/accents are detected properly
        assert any("OBJETIV" in p for p in placeholders), "Should detect OBJETIVOS placeholders"
        
        # Get or pick a contract from DB
        contrato = db.query(Contrato).first()
        if not contrato:
            print("No contracts found in DB to test.")
            return

        # Ensure test contract has specific objectives defined
        contrato.objetivo_general = "Prestar servicios profesionales de asesoría jurídica y representación legal para la Empresa de Servicios Públicos de Yolombó ESP."
        contrato.objetivos_especificos = [
            "Emitir conceptos jurídicos especializados sobre la Ley 142 de 1994 y reglamentación de la CRA.",
            "Revisar y estructurar minutas de contratos de obra, consultoría y prestación de servicios.",
            "Acompañar los procesos de contratación directa y convocatorias públicas de la entidad."
        ]
        db.commit()
        db.refresh(contrato)

        print(f"Testing contract {contrato.numero_contrato} with:")
        print(f"  Objetivo General: {contrato.objetivo_general}")
        print(f"  Objetivos Específicos: {contrato.objetivos_especificos}")

        # Generate DOCX
        output_path = generate_docx_for_contract(contrato, "ESTUDIO_PREVIO")
        print(f"Generated docx at: {output_path}")

        # Open and inspect the generated document
        gen_doc = docx.Document(str(output_path))
        all_text = "\n".join(p.text for p in gen_doc.paragraphs)
        
        print("\n--- Snippet of generated document (first 15 paragraphs) ---")
        for i, p in enumerate(gen_doc.paragraphs[:15]):
            if p.text.strip():
                print(f"P{i+1}: {p.text}")
        print("-----------------------------------------------------------\n")

        # Assertions
        assert "{{OBJETIVOS GENERAL}}" not in all_text, "Placeholder {{OBJETIVOS GENERAL}} was not replaced!"
        assert "{{OBJETIVO GENERAL}}" not in all_text, "Placeholder {{OBJETIVO GENERAL}} was not replaced!"
        assert "{{OBJETIVOS ESPECIFICOS}}" not in all_text, "Placeholder {{OBJETIVOS ESPECIFICOS}} was not replaced!"
        assert "{{OBJETIVOS ESPECÍFICOS}}" not in all_text, "Placeholder {{OBJETIVOS ESPECÍFICOS}} was not replaced!"
        
        # Assert actual values are present
        assert "Prestar servicios profesionales de asesoría jurídica" in all_text, "General objective text not found!"
        assert "Emitir conceptos jurídicos especializados" in all_text, "Specific objective 1 not found!"
        assert "Revisar y estructurar minutas" in all_text, "Specific objective 2 not found!"
        assert "Acompañar los procesos de contratación" in all_text, "Specific objective 3 not found!"

        print("SUCCESS! All general and specific objectives were correctly injected and formatted!")

    finally:
        db.close()

if __name__ == "__main__":
    test_custom_estudio_previo_replacement()

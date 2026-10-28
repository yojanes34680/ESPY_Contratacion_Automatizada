import sys
from pathlib import Path
from datetime import date, datetime
import docx
sys.path.insert(0, str(Path(__file__).parent / "backend"))

from app.services.docx_engine import format_date_spanish, build_contract_context, generate_docx_for_contract
from app.models.contract import Contrato
from app.models.contractor import Contratista, TipoPersona

def run_tests():
    print("--- 1. Testing format_date_spanish function ---")
    assert format_date_spanish("2026-10-15") == "15 de octubre de 2026", f"Failed: {format_date_spanish('2026-10-15')}"
    assert format_date_spanish(date(2026, 10, 15)) == "15 de octubre de 2026"
    assert format_date_spanish(datetime(2026, 10, 15, 12, 0, 0)) == "15 de octubre de 2026"
    assert format_date_spanish("15/10/2026") == "15 de octubre de 2026"
    assert format_date_spanish("2026-03-20") == "20 de marzo de 2026"
    assert format_date_spanish("2027-01-20") == "20 de enero de 2027"
    assert format_date_spanish(None, default="Por definir") == "Por definir"
    assert format_date_spanish("Pendiente de firma") == "Pendiente de firma"
    assert format_date_spanish("15 de octubre de 2026") == "15 de octubre de 2026"
    print("[OK] format_date_spanish unit tests passed!")

    print("\n--- 2. Testing build_contract_context date fields ---")
    mock_contratista = Contratista(
        id=999,
        nombre_o_razon_social="ASOCOMUNAL YOLOMBO",
        tipo_persona=TipoPersona.JURIDICA,
        numero_documento="900123456-7",
        representante_legal_nombre="Representante Legal Prueba",
        representante_legal_documento="12345678"
    )

    mock_contrato = Contrato(
        id=999,
        numero_contrato="ESP-TEST-2026-01",
        vigencia=2026,
        objeto="MANTENIMIENTO DE EQUIPOS Y HERRAMIENTAS",
        tipologia="PRESTACION_SERVICIOS",
        valor_total=15000000.0,
        fecha_estudio_previo="2026-10-15",
        fecha_suscripcion="2026-10-15",
        fecha_acta_inicio="2026-10-15",
        fecha_terminacion="2026-12-15",
        contratista=mock_contratista
    )

    ctx = build_contract_context(mock_contrato)
    assert ctx["FECHA"] == "15 de octubre de 2026", f"ctx['FECHA'] was {ctx['FECHA']}"
    assert ctx["FECHA_ESTUDIO_PREVIO"] == "15 de octubre de 2026"
    assert ctx["FECHA_SUSCRIPCION"] == "15 de octubre de 2026"
    assert ctx["FECHA_ACTA_INICIO"] == "15 de octubre de 2026"
    assert ctx["FECHA_INICIO"] == "15 de octubre de 2026"
    assert ctx["FECHA_TERMINACION"] == "15 de diciembre de 2026"
    assert ctx["FECHA_FINAL"] == "15 de diciembre de 2026"
    print("[OK] build_contract_context date placeholders verified!")

    print("\n--- 3. Testing document generation for ESTUDIO_PREVIO, MINUTA, and ACTA_INICIO ---")
    for t_type in ["ESTUDIO_PREVIO_SERVICIOS", "MINUTA_SERVICIOS", "ACTA_INICIO"]:
        out_path = generate_docx_for_contract(mock_contrato, template_type=t_type)
        assert out_path.exists(), f"File {out_path} not found"
        
        # Read the docx and verify that '15 de octubre de 2026' is present and '2026-10-15' is NOT present as an unformatted date
        d = docx.Document(str(out_path))
        full_text = []
        for p in d.paragraphs:
            full_text.append(p.text)
        for tbl in d.tables:
            for row in tbl.rows:
                for c in row.cells:
                    full_text.append(c.text)
        doc_str = "\n".join(full_text)

        assert "15 de octubre de 2026" in doc_str, f"'15 de octubre de 2026' not found in generated {t_type} docx!"
        assert "2026-10-15" not in doc_str, f"Found unformatted date '2026-10-15' in generated {t_type} docx!"
        print(f"[OK] {t_type} generated docx successfully contains '15 de octubre de 2026' and zero raw '2026-10-15'!")

    print("\n=======================================================")
    print(" ALL SPANISH DATE FORMATTING TESTS PASSED 100%!")
    print("=======================================================")

if __name__ == "__main__":
    run_tests()

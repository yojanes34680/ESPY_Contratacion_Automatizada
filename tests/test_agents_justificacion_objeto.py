"""Test suite for:
1. JustificationAgent: Enforces presence of PDF/DOCX of Propuesta or Estudio de Mercado.
2. ObjectAgent: Formulates contractual object in infinitive, customizable by user.
3. End-to-end integration via /api/agents/generate-justificacion and /api/agents/generate-objeto.
"""
import sys
import io
import docx
from pathlib import Path
from fastapi.testclient import TestClient

# Agregar backend al path para utilidades
sys.path.insert(0, str(Path(__file__).parent / "backend"))

from app.main import app

client = TestClient(app)

def test_justification_requires_proposal_or_market_study():
    print("\n--- TEST 1: Validation - Fail when no proposal/market study exists ---")
    # First, upload only a RUT document
    buf = io.BytesIO(b"%PDF-1.4 Mock RUT DIAN Document text with NIT 901234567-1")
    files = {"file": ("rut_dian_2026.pdf", buf.getvalue(), "application/pdf")}
    r_upload = client.post("/api/documents/upload", files=files, data={"tipo_documento": "RUT"})
    assert r_upload.status_code == 201, f"Upload failed: {r_upload.text}"
    rut_id = r_upload.json()["id"]

    # Try to generate justification using only the RUT
    payload = {
        "document_ids": [rut_id],
        "tipologia": "PRESTACION_SERVICIOS",
        "necesidad_contexto": "Operación de planta potabilizadora"
    }
    r_just = client.post("/api/agents/generate-justificacion", json=payload)
    print("Status code received without proposal:", r_just.status_code)
    assert r_just.status_code == 422, f"Expected 422 Unprocessable Entity, got {r_just.status_code}: {r_just.text}"
    assert "indispensable" in r_just.json()["detail"].lower()
    print("  [OK] Correctly blocked: Indispensable proposal or market study missing.")

def test_justification_with_valid_proposal_document():
    print("\n--- TEST 2: Success - Generate Justification with Proposal Document ---")
    # Create a real docx proposal
    doc = docx.Document()
    doc.add_heading("PROPUESTA TÉCNICA Y ECONÓMICA", level=0)
    doc.add_paragraph("Proponente: SERVICIOS Y MONTAJES HIDRÁULICOS S.A.S. NIT: 900.888.777-1")
    doc.add_paragraph("Objeto de la oferta: Prestación de servicios de mantenimiento electromecánico y control de dosificación química.")
    doc.add_paragraph("Valor total propuesto: $36,000,000 COP M/CTE.")
    doc.add_paragraph("Plazo propuesto: 6 meses.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"file": ("Propuesta_Economica_PTAP_2026.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r_upload = client.post("/api/documents/upload", files=files, data={"tipo_documento": "PROPUESTA"})
    assert r_upload.status_code == 201, f"Upload failed: {r_upload.text}"
    prop_id = r_upload.json()["id"]

    # Call generate-justificacion
    payload = {
        "document_ids": [prop_id],
        "tipologia": "PRESTACION_SERVICIOS",
        "modalidad": "CONTRATACION_DIRECTA",
        "necesidad_contexto": "Mantenimiento especializado de dosificación química en PTAP",
        "contratista_nombre": "SERVICIOS Y MONTAJES HIDRÁULICOS S.A.S.",
        "valor_estimado": 36000000.0
    }
    r_just = client.post("/api/agents/generate-justificacion", json=payload)
    assert r_just.status_code == 200, f"Justification failed: {r_just.text}"
    data = r_just.json()
    print("  Documento analizado:", data["documento_soporte_usado"])
    print("  Tipo detectado:", data["tipo_documento_detectado"])
    print("  Fundamento Ley 142:", data["fundamento_juridico_ley_142"])
    print(f"  Justificación generada:\n{data['justificacion_estudio_previo'][:180]}...")
    assert "Propuesta_Economica_PTAP_2026.docx" in data["documento_soporte_usado"]
    assert "142" in data["justificacion_estudio_previo"]
    print("  [OK] Justification agent elaborated justification with proposal support.")
    return prop_id

def test_object_agent_and_user_customization(prop_id):
    print("\n--- TEST 3: ObjectAgent - Automatic Generation and Customization ---")
    payload = {
        "document_ids": [prop_id],
        "tipologia": "PRESTACION_SERVICIOS",
        "necesidad_descripcion": "mantenimiento electromecanico y control de dosificacion quimica",
        "contratista_nombre": "SERVICIOS Y MONTAJES HIDRÁULICOS S.A.S.",
        "lugar_ejecucion": "la Planta de Tratamiento de Agua Potable de la ESP"
    }
    r_obj = client.post("/api/agents/generate-objeto", json=payload)
    assert r_obj.status_code == 200, f"Generate objeto failed: {r_obj.text}"
    obj_data = r_obj.json()
    print("  Objeto sugerido por el agente:", obj_data["objeto_sugerido"])
    assert obj_data["objeto_sugerido"].lower().startswith("presta")

    # User modifies the object to their liking:
    custom_objeto = "Prestar servicios especializados de mantenimiento electromecánico integral con diagnóstico mensual a satisfacción de la ESP."
    print("  Objeto modificado a gusto del usuario:", custom_objeto)

    import time
    num_cto = f"ESP-TEST-OBJ-{int(time.time())}"
    contract_data = {
        "numero_contrato": num_cto,
        "vigencia": 2026,
        "modalidad_seleccion": "CONTRATACION_DIRECTA",
        "tipologia": "PRESTACION_SERVICIOS",
        "contratista_id": 1,
        "objeto": custom_objeto,
        "justificacion_necesidad": "Justificación fundamentada en propuesta analizada y Ley 142.",
        "valor_total": 36000000.0,
        "plazo_ejecucion": 6,
        "unidad_plazo": "MESES"
    }
    r_create = client.post("/api/contracts", json=contract_data)
    assert r_create.status_code == 201, f"Contract create failed: {r_create.text}"
    contract_id = r_create.json()["id"]
    from app.utils.object_normalizer import normalize_objeto_contractual
    expected_saved_objeto = normalize_objeto_contractual(custom_objeto, tipologia="PRESTACION_SERVICIOS")
    assert r_create.json()["objeto"].strip(".").upper() == expected_saved_objeto.strip(".").upper()
    print(f"  [OK] Contract {contract_id} saved with user customized object normalized: {r_create.json()['objeto']}")

if __name__ == "__main__":
    test_justification_requires_proposal_or_market_study()
    prop_id = test_justification_with_valid_proposal_document()
    test_object_agent_and_user_customization(prop_id)
    print("\n=======================================================")
    print(" ALL TESTS FOR JUSTIFICATION & OBJECT AGENTS PASSED!")
    print("=======================================================")

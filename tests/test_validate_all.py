"""Suite de validación convertida a pruebas pytest aisladas usando auth_client."""
import io
from pathlib import Path
import pytest
import docx
import openpyxl

from app.config import settings
from app.utils.num_to_words import numero_a_letras_cop

def test_num_to_words():
    """Valida la conversión de números a letras en pesos colombianos."""
    v1 = 15450000.0
    l1 = numero_a_letras_cop(v1)
    assert "QUINCE MILLONES" in l1
    assert "PESOS MONEDA CORRIENTE (M/CTE)" in l1

    v2 = 250000000.0
    l2 = numero_a_letras_cop(v2)
    assert "DOSCIENTOS CINCUENTA MILLONES" in l2

def test_api_health_check(auth_client):
    """Verifica el estado del endpoint /api/health."""
    resp = auth_client.get("/api/health")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "healthy"
    assert data["modules_active"]["fase_1_estructura_backend"] is True
    assert data["modules_active"]["fase_2_multi_agent_llm"] is True
    assert data["modules_active"]["fase_3_docx_excel"] is True
    assert data["modules_active"]["fase_4_frontend_ui"] is True

def test_catalogs_typologies_and_modalities(auth_client):
    """Verifica los catálogos normativos para ESP (Ley 142/1994)."""
    resp_typ = auth_client.get("/api/contracts/catalogs/typologies")
    assert resp_typ.status_code == 200
    types = resp_typ.json()
    assert any(t["codigo"] == "PRESTACION_SERVICIOS" for t in types)
    assert any(t["codigo"] == "OBRA" for t in types)
    assert any(t["codigo"] == "SUMINISTRO" for t in types)

    resp_mod = auth_client.get("/api/contracts/catalogs/modalities")
    assert resp_mod.status_code == 200
    mods = resp_mod.json()
    assert any(m["codigo"] == "CONTRATACION_DIRECTA" for m in mods)

def test_contractors_crud(auth_client):
    """Verifica el registro y consulta de contratistas (PN y PJ)."""
    pn_data = {
        "tipo_persona": "NATURAL",
        "tipo_documento": "CC",
        "numero_documento": "1018456789",
        "nombre_o_razon_social": "Carlos Alberto Rodríguez Morales",
        "direccion": "Calle 45 # 13-20",
        "municipio": "Facatativá",
        "departamento": "Cundinamarca",
        "correo_electronico": "carlos.rodriguez@example.com",
        "telefono": "3104567890",
        "regimen_tributario": "NO_RESPONSABLE_IVA",
        "banco_nombre": "Bancolombia",
        "banco_tipo_cuenta": "Ahorros",
        "banco_numero_cuenta": "245-897654-12",
        "perfil_o_experiencia_resumen": "Ingeniero Sanitario con experiencia en PTAP."
    }
    resp_pn = auth_client.post("/api/contractors", json=pn_data)
    assert resp_pn.status_code in (201, 409)

    pj_data = {
        "tipo_persona": "JURIDICA",
        "tipo_documento": "NIT",
        "numero_documento": "901456789",
        "digito_verificacion": "3",
        "nombre_o_razon_social": "Aguas y Obras Civiles de Colombia S.A.S.",
        "nombre_comercial": "Aguas y Obras",
        "representante_legal_nombre": "Martha Helena Gómez Restrepo",
        "representante_legal_documento": "52897456",
        "direccion": "Carrera 7 # 72-15 Of. 402",
        "municipio": "Bogotá D.C.",
        "departamento": "Bogotá D.C.",
        "correo_electronico": "licitaciones@aguasyobras.com.co",
        "telefono": "6017894561",
        "regimen_tributario": "RESPONSABLE_IVA",
        "perfil_o_experiencia_resumen": "Constructora especializada en redes hidráulicas."
    }
    resp_pj = auth_client.post("/api/contractors", json=pj_data)
    assert resp_pj.status_code in (201, 409)

    resp_list = auth_client.get("/api/contractors")
    assert resp_list.status_code == 200
    assert len(resp_list.json()) >= 2

def test_full_contracts_and_documents_flow(auth_client):
    """Flujo completo de contratos, transición de estados, documentos y agentes."""
    # Obtener ID de contratistas existentes
    conts = auth_client.get("/api/contractors").json()
    pn_id = conts[0]["id"]
    pj_id = conts[1]["id"] if len(conts) > 1 else pn_id

    # 1. Contrato Servicios
    c_serv_data = {
        "numero_contrato": "ESP-PS-2026-001",
        "vigencia": 2026,
        "modalidad_seleccion": "CONTRATACION_DIRECTA",
        "tipologia": "PRESTACION_SERVICIOS",
        "contratista_id": pn_id,
        "objeto": "Prestar servicios profesionales especializados como Ingeniero Sanitario para el seguimiento técnico a la operación y calibración dosificadora de la Planta de Tratamiento de Agua Potable (PTAP).",
        "justificacion_necesidad": "La Empresa de Servicios Públicos requiere garantizar la calidad del agua distribuida conforme a la Resolución 2115 de 2007, careciendo en su planta de personal de ingenieros con dedicación técnica en potabilización.",
        "valor_total": 24000000.0,
        "plazo_ejecucion": 6,
        "unidad_plazo": "MESES",
        "estado": "BORRADOR"
    }
    resp_cs = auth_client.post("/api/contracts", json=c_serv_data)
    if resp_cs.status_code == 409:
        c_serv_id = auth_client.get("/api/contracts?query=ESP-PS-2026-001").json()[0]["id"]
    else:
        assert resp_cs.status_code == 201
        c_serv_id = resp_cs.json()["id"]

    # 2. Transición de estado con auditoría
    resp_st = auth_client.patch(
        f"/api/contracts/{c_serv_id}/status",
        json={"nuevo_estado": "EN_REVISION", "observacion": "Estudio previo completado"}
    )
    assert resp_st.status_code == 200
    assert resp_st.json()["estado"] == "EN_REVISION"

    # 3. Subida y descarga de archivo multipart
    contenido_dummy = b"%PDF-1.4 Mock PDF RUT Contratista Carlos Rodriguez para pruebas del sistema ESPY"
    files = {"file": ("RUT_Carlos_Rodriguez_2026.pdf", io.BytesIO(contenido_dummy), "application/pdf")}
    resp_up = auth_client.post(
        "/api/documents/upload",
        data={"contrato_id": str(c_serv_id), "tipo_documento": "RUT", "categoria": "SOPORTE_CONTRATISTA"},
        files=files
    )
    assert resp_up.status_code == 201
    doc_id = resp_up.json()["id"]

    resp_dl = auth_client.get(f"/api/documents/{doc_id}/download")
    assert resp_dl.status_code == 200
    assert resp_dl.content == contenido_dummy

    # 4. Pipeline de agentes
    doc_word = docx.Document()
    doc_word.add_heading("REGISTRO UNICO TRIBUTARIO (RUT) Y PROPUESTA", level=1)
    doc_word.add_paragraph("RAZON SOCIAL: HidroServicios del Valle S.A.S.")
    doc_word.add_paragraph("NIT: 901789654-2")
    doc_word.add_paragraph("REPRESENTANTE LEGAL: Hernando Silva Peña")
    doc_word.add_paragraph("DOCUMENTO REPRESENTANTE: CC 79845123")
    doc_word.add_paragraph("MUNICIPIO: Cali")
    doc_word.add_paragraph("VALOR PROPUESTA: $ 38.500.000 COP")
    doc_word.add_paragraph("PLAZO: 5 MESES")

    word_stream = io.BytesIO()
    doc_word.save(word_stream)
    word_stream.seek(0)

    upload_word = auth_client.post(
        "/api/documents/upload",
        data={"tipo_documento": "PROPUESTA", "categoria": "SOPORTE_CONTRATISTA"},
        files={"file": ("rut_propuesta_hidrovalle.docx", word_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_word.status_code == 201
    doc_word_id = upload_word.json()["id"]

    resp_ext = auth_client.post(f"/api/agents/extract-document/{doc_word_id}")
    assert resp_ext.status_code == 200
    assert resp_ext.json()["numero_documento"] == "901789654"

    resp_autoc = auth_client.post(f"/api/agents/create-contractor-from-doc/{doc_word_id}")
    assert resp_autoc.status_code == 200
    new_cid = resp_autoc.json()["id"]

    # Redacción con agente
    resp_draft = auth_client.post("/api/agents/draft", json={
        "tipologia": "PRESTACION_SERVICIOS",
        "necesidad_descripcion": "Asesoría jurídica y técnica especializada en regulación tarifaria CRA",
        "valor_estimado": 35000000.0,
        "plazo": 7,
        "unidad_plazo": "MESES",
        "contratista_id": new_cid
    })
    assert resp_draft.status_code == 200
    assert resp_draft.json()["tipologia"] == "PRESTACION_SERVICIOS"

    # Generación de documentos Word
    resp_tpls = auth_client.get("/api/templates/available")
    assert resp_tpls.status_code == 200

    resp_gep = auth_client.post(f"/api/templates/generate/{c_serv_id}/ESTUDIO_PREVIO")
    assert resp_gep.status_code == 200

    resp_gmin = auth_client.post(f"/api/templates/generate/{c_serv_id}/MINUTA")
    assert resp_gmin.status_code == 200

    resp_gacta = auth_client.post(f"/api/templates/generate/{c_serv_id}/ACTA_INICIO")
    assert resp_gacta.status_code == 200

    # Sincronización Excel
    resp_sync = auth_client.post(f"/api/excel/sync/{c_serv_id}")
    assert resp_sync.status_code == 200

    resp_sync_all = auth_client.post("/api/excel/sync-all")
    assert resp_sync_all.status_code == 200

    resp_dlex = auth_client.get("/api/excel/download")
    assert resp_dlex.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp_dlex.content))
    assert wb.active["A4"].value == "Vigencia"

    # Cleanup: eliminar contrato de prueba generado para no dejar muestras en base de datos
    auth_client.delete(f"/api/contracts/{c_serv_id}", json={"motivo": "Limpieza de prueba automatizada e2e"})
    auth_client.delete(f"/api/contracts/{c_serv_id}", json={"motivo": "Limpieza física de prueba automatizada e2e"})

def test_spa_endpoints(client):
    """Verifica que la SPA y sus recursos estáticos se sirvan correctamente."""
    resp_html = client.get("/app/", follow_redirects=True)
    assert resp_html.status_code == 200
    assert "ESPY" in resp_html.text

    resp_css = client.get("/app/css/styles.css")
    assert resp_css.status_code == 200

    resp_js = client.get("/app/js/app.js")
    assert resp_js.status_code == 200

    resp_root = client.get("/", follow_redirects=False)
    assert resp_root.status_code in (302, 307)
    assert "/app/" in resp_root.headers.get("location", "")

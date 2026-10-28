"""
Test integral para verificar las correcciones en Minuta del Contrato y Acta de Inicio:
1. Párrafo de comparecencia condicional (Persona Jurídica vs Persona Natural).
2. Estructura de firmas de 3 líneas en Persona Jurídica.
3. Fila de Representante Legal en Tabla Resumen Inicial.
4. Limpieza de OCR en la primera obligación y separación de palabras pegadas.
"""
import sys
import re
from pathlib import Path

# Añadir backend al path
sys.path.insert(0, str(Path(__file__).parent / "backend"))

from app.models.contract import Contrato
from app.models.contractor import Contratista, TipoPersona
from app.services.docx_engine import generate_docx_for_contract
from app.utils.necesidad_builder import extract_items_from_text, clean_ocr_joined_words, clean_signature_footers
from app.utils.obligaciones_builder import build_obligaciones_contratista
import docx


def test_persona_juridica_minuta_and_acta():
    print("\n--- TEST 1 & 2 & 3: PERSONA JURÍDICA (ASOCOMUNAL) EN MINUTA Y ACTA DE INICIO ---")
    contratista = Contratista(
        id=101,
        tipo_persona=TipoPersona.JURIDICA,
        nombre_o_razon_social="ASOCIACION DE JUNTAS DE ACCION COMUNAL Y VIVIENDA COMUNITARIA",
        tipo_documento="NIT",
        numero_documento="800.254.070-8",
        representante_legal_nombre="PAULA ANDREA MONSALVE MEJIA",
        representante_legal_documento="43.593.616"
    )
    contrato = Contrato(
        id=101,
        numero_contrato="CPS-001-2026",
        vigencia=2026,
        objeto="MANTENIMIENTO DE HERRAMIENTAS ELECTRICAS",
        contratista_id=101,
        valor_total=3000000
    )
    contrato.contratista = contratista

    # 1. MINUTA
    out_m = generate_docx_for_contract(contrato, template_type="MINUTA")
    doc_m = docx.Document(str(out_m))

    # Verificar Comparecencia en Minuta
    preamble_m = ""
    for p in doc_m.paragraphs:
        if "por la otra" in p.text.lower():
            preamble_m = p.text
            break
    assert preamble_m, "No se encontró el párrafo de comparecencia en la minuta"
    print("Comparecencia Minuta encontrada:")
    print(" ", preamble_m[preamble_m.find("por la otra"):preamble_m.find("por la otra")+220])
    assert "el señor ASOCIACION" not in preamble_m, "ERROR: Prefijo 'el señor' indebido antes de Persona Jurídica"
    assert "C.C. NIT" not in preamble_m, "ERROR: 'C.C. NIT' indebido en comparecencia"
    assert "por la otra la ASOCIACION DE JUNTAS DE ACCION COMUNAL Y VIVIENDA COMUNITARIA" in preamble_m
    assert "NIT 800.254.070-8" in preamble_m
    assert "representada legalmente por la señora PAULA ANDREA MONSALVE MEJIA" in preamble_m
    assert "identificado(a) con C.C. 43.593.616" in preamble_m
    print("  [OK] Comparecencia en Minuta cumple al 100% la estructura exigida.")

    # Verificar Tabla 0 en Minuta
    t0_m = doc_m.tables[0]
    rep_row_m = None
    for r in t0_m.rows:
        row_txt = [c.text.strip() for c in r.cells]
        if any("representante legal" in c.lower() for c in row_txt):
            rep_row_m = row_txt
            break
    assert rep_row_m, "ERROR: No se encontró la fila de Representante Legal en Tabla 0 de Minuta"
    assert "PAULA ANDREA MONSALVE MEJIA - C.C. 43.593.616" in rep_row_m[1]
    print(f"  [OK] Tabla 0 Minuta incluye Representante Legal: {rep_row_m}")

    # Verificar Bloque de Firmas en Minuta
    sig_paragraphs_m = [p.text for p in doc_m.paragraphs if "ASOCIACION DE JUNTAS" in p.text]
    assert sig_paragraphs_m, "ERROR: No se encontró la razón social en firmas de Minuta"
    all_sig_txt_m = "\n".join(p.text for p in doc_m.paragraphs[-8:])
    assert "PAULA ANDREA MONSALVE MEJIA - R.L." in all_sig_txt_m, "ERROR: Se omitió el nombre del R.L. o se escribió 'RL' suelto en firmas Minuta"
    assert "Contratista" in all_sig_txt_m
    print("  [OK] Firmas en Minuta cumplen 3 líneas para Persona Jurídica.")

    # 2. ACTA DE INICIO
    out_a = generate_docx_for_contract(contrato, template_type="ACTA_INICIO")
    doc_a = docx.Document(str(out_a))

    # Verificar Comparecencia en Acta de Inicio
    preamble_a = ""
    for p in doc_a.paragraphs:
        if "por la otra" in p.text.lower():
            preamble_a = p.text
            break
    assert preamble_a, "No se encontró el párrafo de comparecencia en el Acta de Inicio"
    print("\nComparecencia Acta encontrada:")
    print(" ", preamble_a[preamble_a.find("por la otra"):preamble_a.find("por la otra")+220])
    assert "por la otra la ASOCIACION DE JUNTAS DE ACCION COMUNAL Y VIVIENDA COMUNITARIA" in preamble_a
    assert "NIT 800.254.070-8" in preamble_a
    assert "representada legalmente por la señora PAULA ANDREA MONSALVE MEJIA" in preamble_a
    assert "identificado(a) con C.C. 43.593.616" in preamble_a
    print("  [OK] Comparecencia en Acta de Inicio cumple al 100% la estructura exigida.")

    # Verificar Tabla 0 en Acta de Inicio
    t0_a = doc_a.tables[0]
    rep_row_a = None
    for r in t0_a.rows:
        row_txt = [c.text.strip() for c in r.cells]
        if any("representante legal" in c.lower() for c in row_txt):
            rep_row_a = row_txt
            break
    assert rep_row_a, "ERROR: No se encontró la fila de Representante Legal en Tabla 0 de Acta"
    assert "PAULA ANDREA MONSALVE MEJIA - C.C. 43.593.616" in rep_row_a[1]
    print(f"  [OK] Tabla 0 Acta incluye Representante Legal: {rep_row_a}")

    # Verificar Bloque de Firmas en Acta de Inicio
    all_sig_txt_a = "\n".join(p.text for p in doc_a.paragraphs[-8:])
    assert "ASOCIACION DE JUNTAS DE ACCION COMUNAL Y VIVIENDA COMUNITARIA" in all_sig_txt_a
    assert "PAULA ANDREA MONSALVE MEJIA - R.L." in all_sig_txt_a
    assert "Contratista" in all_sig_txt_a
    print("  [OK] Firmas en Acta de Inicio cumplen 3 líneas para Persona Jurídica.")


def test_persona_natural_comparecencia():
    print("\n--- TEST: PERSONA NATURAL EN MINUTA Y ACTA DE INICIO ---")
    contratista = Contratista(
        id=102,
        tipo_persona=TipoPersona.NATURAL,
        nombre_o_razon_social="CARLOS ALBERTO RESTREPO",
        tipo_documento="CC",
        numero_documento="1.017.123.456"
    )
    contrato = Contrato(
        id=102,
        numero_contrato="CPS-002-2026",
        vigencia=2026,
        objeto="PRESTACION DE SERVICIOS OPERATIVOS",
        contratista_id=102,
        valor_total=2500000
    )
    contrato.contratista = contratista

    out_m = generate_docx_for_contract(contrato, template_type="MINUTA")
    doc_m = docx.Document(str(out_m))
    preamble_m = ""
    for p in doc_m.paragraphs:
        if "por la otra" in p.text.lower():
            preamble_m = p.text
            break
    assert "por la otra el(la) señor(a) CARLOS ALBERTO RESTREPO, identificado(a) con C.C. 1.017.123.456" in preamble_m
    print("  [OK] Persona Natural en Minuta cumple: '...y por la otra el(la) señor(a) [NOMBRE], identificado(a) con C.C. [CÉDULA]...'")

    out_a = generate_docx_for_contract(contrato, template_type="ACTA_INICIO")
    doc_a = docx.Document(str(out_a))
    preamble_a = ""
    for p in doc_a.paragraphs:
        if "por la otra" in p.text.lower():
            preamble_a = p.text
            break
    assert "por la otra el(la) señor(a) CARLOS ALBERTO RESTREPO, identificado(a) con C.C. 1.017.123.456" in preamble_a
    print("  [OK] Persona Natural en Acta de Inicio cumple formato estándar.")


def test_limpieza_primera_obligacion_y_ocr():
    print("\n--- TEST 4: LIMPIEZA DE LA PRIMERA OBLIGACIÓN EN MINUTA (OCR Y PIE DE FIRMA) ---")
    dirty_proposal_text = (
        "Autonamia, compromiso y ASOCOMUNALYOLOMBO participacion NIT800254070-8\n"
        "Asociacion de Juntas de Accion Comunal y Vivienda Comunitaria.\n"
        "Senores:\nEMPRESADESERVICIOSPUBLUICOSDELMUNICIPIODEYOLOMBO\n"
        "ASUNTO:COTIZACIONATODOCOSTO\n"
        "Producto\nVALOR\n"
        "MANTENIMIENTO A TODO.COSTO DE LAS HERRAMIENTAS ELECTRICAS\n"
        "3.000.000\n"
        "DE LA EMPRESA DE SERVICIOS PUBLICOS DELMUNICIPIODE YOLOMBO.\n"
        "PAULAANDREAMONSALVE\n"
        "Presidente ASOcOMUNAL"
    )

    items = extract_items_from_text(dirty_proposal_text)
    print("Ítems extraídos desde propuesta con OCR:")
    for it in items:
        print("  *", it)

    assert len(items) >= 1
    first_item = items[0]
    assert "DELMUNICIPIODE" not in first_item, "ERROR: Palabra pegada DELMUNICIPIODE no fue separada"
    assert "DEL MUNICIPIO DE" in first_item or "DE YOLOMBO" in first_item
    assert "PAULAANDREAMONSALVE" not in first_item, "ERROR: Nombre del firmante concatenado al final de la actividad"
    assert "PRESIDENTE ASOCOMUNAL" not in first_item, "ERROR: Cargo del firmante concatenado a la actividad"

    # Ahora verificar a través de build_obligaciones_contratista (que nutre la Cláusula Quinta de la Minuta)
    oblig_data = build_obligaciones_contratista(
        objeto="MANTENIMIENTO DE HERRAMIENTAS ELECTRICAS",
        tipologia="PRESTACION_SERVICIOS",
        contexto_adicional=dirty_proposal_text
    )
    first_oblig = oblig_data["obligaciones_minuta"][0]
    print(f"Primera obligación generada para Cláusula Quinta:\n  --> {first_oblig}")

    assert "PAULA" not in first_oblig.upper(), "ERROR: Se incluyó el nombre del firmante en la primera obligación"
    assert "PRESIDENTE" not in first_oblig.upper(), "ERROR: Se incluyó el cargo del firmante en la primera obligación"
    assert "ASOCOMUNAL" not in first_oblig.upper(), "ERROR: Se incluyó el nombre de la asociación al final de la obligación técnica"
    assert "DELMUNICIPIODE" not in first_oblig, "ERROR: Falta separar DELMUNICIPIODE"
    print("  [OK] Primera obligación técnica completamente limpia de firmas y palabras pegadas.")


if __name__ == "__main__":
    test_persona_juridica_minuta_and_acta()
    test_persona_natural_comparecencia()
    test_limpieza_primera_obligacion_y_ocr()
    print("\n" + "="*70)
    print("¡TODOS LOS 4 AJUSTES FUERON VERIFICADOS EXITOSAMENTE AL 100%!")
    print("="*70)

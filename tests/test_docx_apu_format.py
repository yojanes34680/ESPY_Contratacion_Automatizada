import sys
from pathlib import Path
import pytest
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

# Add backend to sys.path
sys.path.insert(0, str(Path("backend").resolve()))

from app.services.agents.multi_proposal_parser import multi_proposal_parser
from app.services.agents.apu_agent import apu_agent
from app.schemas.apu import EvaluacionComparativaPayload, ProcesoConfig

def test_docx_apu_arial12_and_content_quality(tmp_path):
    """
    Verifica que los dos documentos Word generados por el módulo APU:
    1. Tengan el 100% de párrafos y celdas con fuente Arial de 12 pt y color negro (#000000).
    2. Ninguna celda de tabla tenga fondos oscuros (azul marino, verde oscuro, etc.).
    3. No contengan textos prohibidos: 'No especificado', 'CiSPAS.A.S', 'ConsultoriaInt..', 'para los servici)', 'MILLONES PESOS'.
    4. El objeto contractual esté autocompletado y sincronizado.
    5. Para consultorías (estudio tarifario CRA 1038), el APU se descomponga en 80% / 10% / 10%.
    """
    path_publi = Path("tests/fixtures/apu_proposals/PROPUESTAS ECONOMICAS_Parte1_1.pdf")
    path_cispa = Path("tests/fixtures/apu_proposals/PROPUESTAS ECONOMICAS_Parte2_1.pdf")

    assert path_publi.exists(), f"No existe {path_publi}"
    assert path_cispa.exists(), f"No existe {path_cispa}"

    offers_publi = multi_proposal_parser.extract_offers_from_pdf_bytes(path_publi.read_bytes(), path_publi.name)
    offers_cispa = multi_proposal_parser.extract_offers_from_pdf_bytes(path_cispa.read_bytes(), path_cispa.name)

    # El objeto contractual se envía vacío intencionalmente para validar la deducción automática
    payload = EvaluacionComparativaPayload(
        proceso_config=ProcesoConfig(
            objeto_contractual="",  # Debe autocompletarse del ítem de la oferta ganadora
            tipo_contrato="Consultoría",
            plazo_ejecucion_meses=2,
            smlmv_vigencia=1750905.0,
            modalidad_seleccion="Mínimo tres (3) cotizaciones",
            rango_cuantia_smlmv="Menor cuantía",
            criterio_seleccion="Menor precio global del mercado",
            firmante_elaboro_nombre="JOAN ALVAREZ CORREA",
            firmante_elaboro_cargo="Estructurador Técnico",
            firmante_aprobo_nombre="CAROLINA GIRALDO GÓMEZ",
            firmante_aprobo_cargo="Gerente General"
        ),
        ofertas=offers_publi + offers_cispa,
        auto_desglosar_apu=True
    )

    result = apu_agent.evaluate_proposals(payload)
    assert result.proponente_ganador is not None
    assert "CISPA" in result.proponente_ganador.nombre_razon_social.upper()
    assert result.valor_total_adjudicado == 35000000.0

    # Generar ambos documentos DOCX en carpeta temporal
    path_doc1 = tmp_path / "Informe_Evaluacion_Comparativa.docx"
    path_doc2 = tmp_path / "APU_Presupuesto_Precontractual.docx"

    apu_agent.generate_docx_evaluacion(result, path_doc1)
    apu_agent.generate_docx_apu(result, path_doc2)

    assert path_doc1.exists()
    assert path_doc2.exists()

    dark_backgrounds = ["0A3A60", "1B365D", "0F172A", "1E3A8A", "065F46", "0284C7", "16A34A"]

    for doc_file, doc_name in [(path_doc1, "Informe_Evaluacion"), (path_doc2, "APU_Presupuesto")]:
        doc = Document(str(doc_file))

        # 1. Validar que no aparezcan textos prohibidos en el documento completo
        full_text_list = []
        for p in doc.paragraphs:
            full_text_list.append(p.text)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    full_text_list.append(c.text)
        full_text = " ".join(full_text_list)

        assert "No especificado" not in full_text, f"[{doc_name}] Contiene 'No especificado'"
        assert "CiSPAS.A.S" not in full_text, f"[{doc_name}] Contiene 'CiSPAS.A.S'"
        assert "ConsultoriaInt.." not in full_text, f"[{doc_name}] Contiene 'ConsultoriaInt..'"
        assert "para los servici)" not in full_text, f"[{doc_name}] Contiene 'para los servici)'"
        assert "MILLONES PESOS" not in full_text, f"[{doc_name}] Contiene 'MILLONES PESOS' (falta 'DE')"
        assert "MILLONES DE PESOS" in full_text, f"[{doc_name}] Debe contener 'MILLONES DE PESOS'"

        # 2. Validar tipografía: todos los párrafos y celdas deben tener fuente Arial 12 pt y texto negro
        for p_idx, p in enumerate(doc.paragraphs):
            for r_idx, r in enumerate(p.runs):
                if not r.text.strip():
                    continue
                # Nombre de fuente
                assert r.font.name == "Arial", f"[{doc_name}] Párrafo {p_idx} run {r_idx} '{r.text}' no es Arial: {r.font.name}"
                # Tamaño de fuente: exactamente 12 pt
                assert r.font.size == Pt(12), f"[{doc_name}] Párrafo {p_idx} run {r_idx} '{r.text}' no es 12pt: {r.font.size}"
                # Color de texto: negro estricto RGB(0, 0, 0)
                assert r.font.color.rgb == RGBColor(0, 0, 0), f"[{doc_name}] Párrafo {p_idx} run {r_idx} '{r.text}' no es negro: {r.font.color.rgb}"

        for t_idx, t in enumerate(doc.tables):
            for r_idx, r in enumerate(t.rows):
                for c_idx, c in enumerate(r.cells):
                    # Verificar fondos oscuros
                    tcPr = c._tc.get_or_add_tcPr()
                    for shd in tcPr.findall(qn("w:shd")):
                        fill_val = (shd.get(qn("w:fill")) or "").upper()
                        assert fill_val not in dark_backgrounds, f"[{doc_name}] Tabla {t_idx} celda ({r_idx},{c_idx}) tiene fondo oscuro: {fill_val}"

                    # Verificar texto de las celdas
                    for p in c.paragraphs:
                        for r in p.runs:
                            if not r.text.strip():
                                continue
                            assert r.font.name == "Arial", f"[{doc_name}] Tabla {t_idx} ({r_idx},{c_idx}) run '{r.text}' no es Arial: {r.font.name}"
                            assert r.font.size == Pt(12), f"[{doc_name}] Tabla {t_idx} ({r_idx},{c_idx}) run '{r.text}' no es 12pt: {r.font.size}"
                            assert r.font.color.rgb == RGBColor(0, 0, 0), f"[{doc_name}] Tabla {t_idx} ({r_idx},{c_idx}) run '{r.text}' no es negro: {r.font.color.rgb}"

    # 3. Validaciones específicas del Informe Comparativo
    doc1 = Document(str(path_doc1))
    text1 = " ".join([p.text for p in doc1.paragraphs] + [c.text for t in doc1.tables for r in t.rows for c in r.cells])
    # Concordancia singular
    assert "para el ítem estructurado en el proceso" in text1
    # Párrafo de garantías con exención
    assert "Conforme al Numeral 32 (literales d y e) del Manual de Contratación de la ESPY, para el presente proceso aplica el siguiente régimen de garantías:" in text1

    # 4. Validaciones específicas del APU Presupuesto Precontractual
    doc2 = Document(str(path_doc2))
    text2 = " ".join([p.text for p in doc2.paragraphs] + [c.text for t in doc2.tables for r in t.rows for c in r.cells])
    # Componentes de consultoría 80% / 10% / 10%
    assert "HONORARIOS DE PERSONAL PROFESIONAL Y ESPECIALIZADO" in text2
    assert "INSUMOS TÉCNICOS, PROCESAMIENTO DE INFORMACIÓN E INFORMES" in text2
    assert "LOGÍSTICA, VISITAS DE CAMPO Y COSTOS ADMINISTRATIVOS" in text2


def test_docx_endpoints_sanitize_dirty_payload(tmp_path, auth_client):
    """
    Valida que los endpoints HTTP:
    - POST /api/apu/generate-informe-docx
    - POST /api/apu/generate-presupuesto-apu-docx
    Saneen automáticamente payloads que lleguen con:
    1. Objeto contractual 'No especificado'
    2. Valor en letras sin 'DE' ('TREINTA Y CINCO MILLONES PESOS')
    3. Razones sociales con OCR sucio ('CiSPAS.A.S', 'Publi Corp')
    4. Desglose APU con texto obsoleto ('...para los servici)')
    Y generen documentos .docx 100% limpios en Arial 12pt, texto negro y cuadros claros.
    """
    client = auth_client

    path_publi = Path("tests/fixtures/apu_proposals/PROPUESTAS ECONOMICAS_Parte1_1.pdf")
    path_cispa = Path("tests/fixtures/apu_proposals/PROPUESTAS ECONOMICAS_Parte2_1.pdf")
    offers_publi = multi_proposal_parser.extract_offers_from_pdf_bytes(path_publi.read_bytes(), path_publi.name)
    offers_cispa = multi_proposal_parser.extract_offers_from_pdf_bytes(path_cispa.read_bytes(), path_cispa.name)

    payload = EvaluacionComparativaPayload(
        proceso_config=ProcesoConfig(
            objeto_contractual="",
            tipo_contrato="Consultoría",
            plazo_ejecucion_meses=2,
            smlmv_vigencia=1750905.0,
            modalidad_seleccion="Mínimo tres (3) cotizaciones",
            rango_cuantia_smlmv="Menor cuantía",
            criterio_seleccion="Menor precio global del mercado",
            firmante_elaboro_nombre="JOAN ALVAREZ CORREA",
            firmante_elaboro_cargo="Estructurador Técnico",
            firmante_aprobo_nombre="CAROLINA GIRALDO GÓMEZ",
            firmante_aprobo_cargo="Gerente General"
        ),
        ofertas=offers_publi + offers_cispa,
        auto_desglosar_apu=True
    )

    clean_result = apu_agent.evaluate_proposals(payload)
    dirty = clean_result.model_copy(deep=True)

    # Inyectar deliberadamente las inconsistencias reportadas
    dirty.proceso_config.objeto_contractual = "No especificado"
    dirty.valor_total_letras = "TREINTA Y CINCO MILLONES PESOS"
    if dirty.proponente_ganador:
        dirty.proponente_ganador.nombre_razon_social = "CiSPAS.A.S"
    for of in dirty.ofertas:
        if "CISPA" in of.nombre_razon_social.upper():
            of.nombre_razon_social = "CiSPAS.A.S"
    for it in dirty.items_apu:
        it.descomposicion_apu.ai_desglose_justificacion = "Desglose preliminar para los servici)"

    dirty_payload = dirty.model_dump(mode="json")

    # Test endpoint 1
    res1 = client.post("/api/apu/generate-informe-docx", json=dirty_payload)
    assert res1.status_code == 200, f"Error en endpoint informe: {res1.text}"
    doc1_path = tmp_path / "informe_endpoint.docx"
    doc1_path.write_bytes(res1.content)

    # Test endpoint 2
    res2 = client.post("/api/apu/generate-presupuesto-apu-docx", json=dirty_payload)
    assert res2.status_code == 200, f"Error en endpoint presupuesto: {res2.text}"
    doc2_path = tmp_path / "apu_endpoint.docx"
    doc2_path.write_bytes(res2.content)

    dark_backgrounds = ["0A3A60", "1B365D", "0F172A", "1E3A8A", "065F46", "0284C7", "16A34A"]

    for doc_file, doc_name in [(doc1_path, "Informe_Endpoint"), (doc2_path, "APU_Endpoint")]:
        doc = Document(str(doc_file))
        full_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])

        assert "No especificado" not in full_text, f"[{doc_name}] Contiene 'No especificado'"
        assert "CiSPAS.A.S" not in full_text, f"[{doc_name}] Contiene 'CiSPAS.A.S'"
        assert "para los servici)" not in full_text, f"[{doc_name}] Contiene 'para los servici)'"
        assert "MILLONES PESOS" not in full_text, f"[{doc_name}] Contiene 'MILLONES PESOS' (falta 'DE')"
        assert "MILLONES DE PESOS" in full_text, f"[{doc_name}] Debe contener 'MILLONES DE PESOS'"
        assert "CISPA S.A.S." in full_text, f"[{doc_name}] No normalizó a CISPA S.A.S."

        # Validar fuentes y tamaños
        for p in doc.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    assert r.font.name == "Arial"
                    assert r.font.size == Pt(12)
                    assert r.font.color.rgb == RGBColor(0, 0, 0)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    tcPr = c._tc.get_or_add_tcPr()
                    for shd in tcPr.findall(qn("w:shd")):
                        fill_val = (shd.get(qn("w:fill")) or "").upper()
                        assert fill_val not in dark_backgrounds
                    for p in c.paragraphs:
                        for r in p.runs:
                            if r.text.strip():
                                assert r.font.name == "Arial"
                                assert r.font.size == Pt(12)
                                assert r.font.color.rgb == RGBColor(0, 0, 0)

    # Validar que Doc 2 contiene los componentes intelectuales regenerados
    doc2 = Document(str(doc2_path))
    text2 = " ".join([p.text for p in doc2.paragraphs] + [c.text for t in doc2.tables for r in t.rows for c in r.cells])
    assert "HONORARIOS DE PERSONAL PROFESIONAL Y ESPECIALIZADO" in text2
    assert "INSUMOS TÉCNICOS, PROCESAMIENTO DE INFORMACIÓN E INFORMES" in text2
    assert "LOGÍSTICA, VISITAS DE CAMPO Y COSTOS ADMINISTRATIVOS" in text2

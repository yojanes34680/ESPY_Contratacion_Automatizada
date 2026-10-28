"""Automated tests for Supervisor details, Contract Dates, Excel Sync/Delete, and Typology Metrics."""
import time
import openpyxl
from pathlib import Path

def test_supervisor_dates_and_excel(auth_client):
    r = auth_client.get("/api/health")
    assert r.status_code == 200, f"Health check failed: {r.text}"

    # Get or create contractor
    r_c = auth_client.get("/api/contractors")
    contractors = r_c.json()
    if not contractors:
        c_res = auth_client.post("/api/contractors", json={
            "tipo_persona": "NATURAL",
            "tipo_documento": "CC",
            "numero_documento": "123456789",
            "nombre_o_razon_social": "Contratista Supervisor Test",
            "correo_electronico": "test@espy.com",
            "telefono": "3000000000"
        })
        cid = c_res.json()["id"]
    else:
        cid = contractors[0]["id"]

    num_contrato = f"ESP-OB-2026-TEST{int(time.time() * 10) % 100000}"
    payload = {
        "numero_contrato": num_contrato,
        "vigencia": 2026,
        "modalidad_seleccion": "INVITACION_PUBLICA",
        "tipologia": "OBRA",
        "contratista_id": cid,
        "objeto": "Construcción de tanque de almacenamiento y desarenador PTAP sector norte",
        "justificacion_necesidad": "Asegurar continuidad y presión en el suministro de agua potable.",
        "valor_total": 120000000.0,
        "plazo_ejecucion": 10,
        "unidad_plazo": "MESES",
        "supervisor_nombre": "Ing. Carlos Alberto Montoya",
        "supervisor_cargo": "Subdirector de Infraestructura y Obras",
        "supervisor_tipo_documento": "CC",
        "supervisor_documento": "71987654",
        "fecha_suscripcion": "2026-03-15",
        "fecha_acta_inicio": "2026-03-20",
        "fecha_terminacion": "2027-01-20",
        "estado": "EN_EJECUCION"
    }
    r_create = auth_client.post("/api/contracts", json=payload)
    assert r_create.status_code == 201, f"Create contract failed: {r_create.text}"
    contract_id = r_create.json()["id"]

    # Verify contract data via GET
    r_get = auth_client.get(f"/api/contracts/{contract_id}")
    assert r_get.status_code == 200
    c_data = r_get.json()
    assert c_data["supervisor_cargo"] == "Subdirector de Infraestructura y Obras"

    # Sync with Excel
    r_sync = auth_client.post(f"/api/excel/sync/{contract_id}")
    assert r_sync.status_code == 200, f"Sync failed: {r_sync.text}"
    sync_res = r_sync.json()
    excel_file_path = Path(sync_res["resultado"]["archivo_maestro"])

    # Check Excel file content
    wb = openpyxl.load_workbook(str(excel_file_path))
    ws = wb.active
    target_row = sync_res["resultado"]["fila_excel"]
    headers = [ws.cell(row=4, column=c).value for c in range(1, 20)]
    assert "Supervisor Designado" in headers

    sup_nom = ws.cell(row=target_row, column=13).value
    sup_cargo = ws.cell(row=target_row, column=14).value
    sup_doc = ws.cell(row=target_row, column=15).value
    f_firma = ws.cell(row=target_row, column=16).value
    f_inicio = ws.cell(row=target_row, column=17).value
    f_fin = ws.cell(row=target_row, column=18).value

    assert sup_nom == "Ing. Carlos Alberto Montoya"
    assert sup_cargo == "Subdirector de Infraestructura y Obras"
    assert "71987654" in str(sup_doc)
    assert str(f_firma) == "2026-03-15"
    assert str(f_inicio) == "2026-03-20"
    assert str(f_fin) == "2027-01-20"
    wb.close()

    # Update Contract and Re-Sync in Excel
    update_payload = {
        "supervisor_cargo": "Gerente de Infraestructura Hidráulica",
        "fecha_terminacion": "2027-03-20"
    }
    r_up = auth_client.put(f"/api/contracts/{contract_id}", json=update_payload)
    assert r_up.status_code == 200, f"Update failed: {r_up.text}"
    
    r_resync = auth_client.post(f"/api/excel/sync/{contract_id}")
    assert r_resync.status_code == 200
    resync_res = r_resync.json()
    assert resync_res["resultado"]["accion"] == "ACTUALIZADO"
    
    wb2 = openpyxl.load_workbook(str(excel_file_path))
    ws2 = wb2.active
    updated_cargo = ws2.cell(row=target_row, column=14).value
    updated_fin = ws2.cell(row=target_row, column=18).value
    wb2.close()
    assert updated_cargo == "Gerente de Infraestructura Hidráulica"
    assert str(updated_fin) == "2027-03-20"

    # Delete Contract from Excel Matrix
    r_del = auth_client.delete(f"/api/excel/contract/{contract_id}")
    assert r_del.status_code == 200, f"Delete from excel failed: {r_del.text}"
    del_res = r_del.json()
    assert del_res["resultado"]["eliminado"] is True

    # Verify that contract is no longer in Excel
    wb3 = openpyxl.load_workbook(str(excel_file_path))
    ws3 = wb3.active
    found = False
    for r in range(5, ws3.max_row + 1):
        if ws3.cell(row=r, column=2).value == num_contrato:
            found = True
            break
    wb3.close()
    assert not found, f"Contract {num_contrato} was still found in Excel!"

    # Verify that contract STILL exists in database
    r_still = auth_client.get(f"/api/contracts/{contract_id}")
    assert r_still.status_code == 200

    r_all = auth_client.get("/api/contracts")
    assert r_all.status_code == 200
    all_contracts = r_all.json()
    
    typologies = {}
    for c in all_contracts:
        tip = c["tipologia"]
        if tip not in typologies:
            typologies[tip] = {"total": 0, "executed": 0, "value": 0.0}
        typologies[tip]["total"] += 1
        if c["estado"] in ["EN_EJECUCION", "LIQUIDADO", "APROBADO"]:
            typologies[tip]["executed"] += 1
        typologies[tip]["value"] += (c["valor_total"] or 0.0)

    assert "OBRA" in typologies
